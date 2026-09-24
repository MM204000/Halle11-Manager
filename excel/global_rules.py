"""Mappenweite Regeln (Typo-Skala, Ausrichtung, Tabellenstil, Zahlenformate, Farbsemantik, Druck).

early(wb):      vor den Blatt-Layouts – Standards für ALLE Blätter (Blatt-Module dürfen sie danach überschreiben):
                Korrektorat, Zahlenformate, Typo-Skala, keine senkrechten Linien, Seitenkopf, Tabellenstil,
                Zeilenraster, Fuß, Farbsemantik, Blattschutz der Bank-Vorlagen.
final(wb):      nach den Blatt-Layouts – nur Sicherheits-Normalisierungen (kein „An Zellgröße anpassen“,
                Einzug nur mit links/rechts, Mindestschrift 8 pt, Sekundärtext ≤ 9 pt nie in 8A9099,
                nur Calibri, Normzitate in Versaltexten, dxf ohne Schriftnamen).
page_setup(wb): ganz am Ende – A4, Ränder, Kopf-/Fußzeile, Druckbereiche, Maßstab, Umbrüche, Drucktitel,
                Blattreihenfolge, Registerfarben, Ansicht, Objektschutz.

Nur Darstellung: keine Formel der Vorlage, kein Eingabewert und kein Name wird verändert.
"""
import re

from openpyxl.styles import Alignment, Border, Font, Protection
from openpyxl.worksheet.pagebreak import Break, ColBreak, RowBreak
from openpyxl.worksheet.properties import PageSetupProperties

import core as C

# =============================================================================== Blattgruppen
STEP_RE = re.compile(r"S\d\d ")
FORM_SHEETS = ("Eingaben", "Konfiguration")
BANK_SHEETS = ("Bankgespräch", "Haushaltsrechnung", "Vermögensaufstellung")
CALC_SHEETS = ("Projektion", "Steuern", "Finanzierung", "AfA-Vergleich", "Sensitivität")
PRESENTATION = ("Start", "Leitfaden", "Dashboard", "Cockpit") + BANK_SHEETS   # + S01–S12 (ohne Zeilen-/Spaltenköpfe)

ORDER = ["Start", "Leitfaden"] + [None] * 12 + [
    "Dashboard", "Cockpit", "Diagramme", "Eingaben", "Projektion", "Steuern", "AfA-Vergleich", "Finanzierung",
    "Sensitivität", "Bankgespräch", "Haushaltsrechnung", "Vermögensaufstellung", "Hinweise", "Konfiguration"]

TAB_COLOR = {"Start": C.NAVY, "Leitfaden": C.NAVY, "Dashboard": C.NAVY,
             "Cockpit": C.ACCENT, "Diagramme": C.ACCENT, "Eingaben": C.ACCENT, "Projektion": C.ACCENT,
             "Steuern": C.ACCENT, "AfA-Vergleich": C.ACCENT, "Finanzierung": C.ACCENT, "Sensitivität": C.ACCENT,
             "Bankgespräch": "8FA9C9", "Haushaltsrechnung": "8FA9C9", "Vermögensaufstellung": "8FA9C9",
             "Hinweise": "A0A7B1", "Konfiguration": "A0A7B1"}


def tab_color(name):
    return C.BLUE if STEP_RE.match(name) else TAB_COLOR.get(name, C.ACCENT)


# =============================================================================== Seitenkopf-Standard (P2-05)
# Blatt: (Titelspalte, letzte Inhaltsspalte, Eyebrow, H1, Untertitel). Die Blatt-Module überschreiben ihn bei Bedarf.
HEADERS = {
    "Eingaben": ("B", "L", "Eingaben", "Eingaben",
                 "Vollständige Übersicht aller Eingaben und Profi-Felder  ·  Gelb hinterlegt = Eingabe, "
                 "blau = aus dem Leitfaden übernommen (dort ändern)"),
    "Sensitivität": ("C", "M", "Berechnung", "Sensitivität",
                     "Was passiert, wenn Miete, Zins oder Wertentwicklung anders laufen?  ·  Jahr-1-Größen als "
                     "lineare Näherung, die IRR-Matrix rechnet die volle Zahlungsreihe neu"),
    "Haushaltsrechnung": ("B", "E", "Bank  ›  Haushaltsrechnung", "Haushaltsrechnung",
                          "Selbstauskunft für die Bank  ·  Monatswerte in die gelb hinterlegten Felder, "
                          "Jahreswerte rechnen sich selbst"),
    "Vermögensaufstellung": ("B", "E", "Bank  ›  Vermögensaufstellung", "Vermögensaufstellung",
                             "Selbstauskunft für die Bank  ·  aktuelle Werte eingeben und belegen "
                             "(Gutachten, Depot- oder Kontoauszug)"),
    "Hinweise": ("B", "E", "Anhang  ›  Hinweise", "Hinweise",
                 "Steuerliche Regelungen, Rechtsgrundlagen und Modellannahmen  ·  Rechtsstand September 2026 "
                 "(Investitionssofortprogramm 2025, JStG 2024, EStG i. d. F. 2026)  ·  keine Steuerberatung im Einzelfall"),
    "Konfiguration": ("B", "G", "Anhang  ›  Konfiguration", "Konfiguration",
                      "Stammdaten und Rechtsstand September 2026  ·  Werte hier zentral pflegen, "
                      "alle Berechnungen greifen auf diese Zellen zu"),
}

# Tabellen (volle Breite für Kopf-, Summen- und Zeilenlinien, P2-03) und Inhaltsbreite für den Fuß
TABLES = {
    "Eingaben": [("B", "L")],
    "Haushaltsrechnung": [("B", "E")],
    "Vermögensaufstellung": [("B", "E")],
    "Hinweise": [("B", "E")],
    "Konfiguration": [("B", "E"), ("G", "G")],
}
CONTENT = {"Eingaben": ("B", "L"), "Haushaltsrechnung": ("B", "E"), "Vermögensaufstellung": ("B", "E"),
           "Hinweise": ("B", "E"), "Konfiguration": ("B", "G"), "Sensitivität": ("C", "M"),
           "Bankgespräch": ("B", "I")}

# =============================================================================== Zahlenformate (P2-06)
NUMFMT_MAP = {
    '#,##0" €"': C.NUMFMT["eur"],
    '#,##0" €";\\-#,##0" €";\\–': C.NUMFMT["eur"],
    '#,##0" €";-#,##0" €";"–"': C.NUMFMT["eur"],
    '#,##0.00" €";\\-#,##0.00" €";\\–': C.NUMFMT["eur2"],
    '#,##0;\\-#,##0;\\–': C.NUMFMT["num"],
    "0.00%": C.NUMFMT["pct2"], "0.00 %": C.NUMFMT["pct2"],
    "0.0%": C.NUMFMT["pct1"], "0.0 %": C.NUMFMT["pct1"],
    "0.00\\x": C.NUMFMT["mult2"], '0.00"×"': C.NUMFMT["mult2"],
    '0.0"-fach"': C.NUMFMT["mult1"],
    "dd\\.mm\\.yyyy": C.NUMFMT["date"], "dd.mm.yyyy": C.NUMFMT["date"],
}
# Sätze (Zins, Steuer, AfA, GrESt, Soli, KSt) immer mit zwei Nachkommastellen
RATE_CELLS = {
    "Konfiguration": ["C10:C25", "C38:C40", "C45:C53", "C70"],
    "Steuern": ["C18:C19", "D71:AQ71"],
    "Finanzierung": ["D14", "D24", "D39:AQ39"],
}
INT_CELLS = {"Konfiguration": ["C33", "C35"]}

# =============================================================================== Typo-Skala (P2-07)
SCALE = (8, 9, 10, 12.5, 14, 16, 20, 22, 30)
KEEP_SIZE = {("Start", "D23"), ("Start", "G23")}   # Auswahlfeld „Kauf als“ (11 pt) und sein ▾
SYMBOLS = set("✓○●▾›‹→↓↑⚠•✗×")


def _symbol(v):
    return isinstance(v, str) and 0 < len(v.strip()) <= 2 and all(ch in SYMBOLS or ch == " " for ch in v)

VERSAL_FIX = [(re.compile(r"§ 32A\b"), "§ 32a"), (re.compile(r"\bESTG\b"), "EStG"), (re.compile(r"\bGRESTG\b"), "GrEStG"),
              (re.compile(r"\bKSTG\b"), "KStG"), (re.compile(r"\bGEWSTG\b"), "GewStG"), (re.compile(r"\bESTDV\b"), "EStDV"),
              (re.compile(r"\bSOLZG\b"), "SolZG"), (re.compile(r"\bUSTG\b"), "UStG"), (re.compile(r"\bAFA\b"), "AfA"),
              (re.compile(r"\(AFA\)"), "(AfA)"), (re.compile(r"\bABS\. "), "Abs. "), (re.compile(r"\bNR\. "), "Nr. "),
              (re.compile(r"\bI\. D\. F\."), "i. d. F.")]

OLD_FONTS = {None, "Aptos", "Aptos Display", "Aptos Narrow", "Inter", "Fraunces", "IBM Plex Mono", "Arial"}
LINE_COLORS = {C.LINE, C.LINE2, "D6D0C2", "B9B4A6", "E2DCCE", "EBE6DB"}


# =============================================================================== Hilfen
def _rgb(color):
    if color is None or getattr(color, "type", None) != "rgb" or not isinstance(color.rgb, str):
        return None
    return color.rgb[-6:].upper()


def _fill_rgb(c):
    f = c.fill
    if f is None or f.fill_type != "solid":
        return None
    return _rgb(f.fgColor)


def _static(c):
    return isinstance(c.value, str) and c.data_type == "s" and not C.is_formula(c.value)


def _is_input(c):
    return _fill_rgb(c) == C.INPUT_BG or c.protection.locked is False


def _cells(ws, ref):
    if ":" not in ref:
        return [ws[ref]]
    return [c for row in ws[ref] for c in row]


def _merged_anchor_map(ws):
    m = {}
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                m[(r, cc)] = (mr.min_row, mr.min_col, mr.max_col)
    return m


def _numeric_fmt(fmt):
    return fmt not in (None, "General", "@") and any(ch in fmt for ch in "0#")


def _map_size(c):
    """Zwischenstufen auf die Skala 8 · 9 · 10 · 12,5 · 16 · 20 · 22 (Hero 30, Schrittnummer 14)."""
    sz = c.font.sz or 11
    if sz in SCALE:
        return sz
    if sz < 8:
        return 8
    if sz == 8.5:
        return 8 if (c.font.b and isinstance(c.value, str) and c.value.upper() == c.value) else 9
    if sz == 9.5:
        return 10 if (not isinstance(c.value, str) or C.is_formula(c.value) or c.font.b) else 9
    if sz <= 11:
        return 10
    if sz <= 13.5:
        return 12.5
    if sz <= 15:
        return 14
    if sz <= 18:
        return 16
    if sz <= 21:
        return 20
    if sz <= 26:
        return 22
    return 30


def _set_font(c, **kw):
    f = c.font
    base = dict(name=f.name, sz=f.sz, b=f.b, i=f.i, u=f.u, color=f.color, strike=f.strike, vertAlign=f.vertAlign)
    base.update(kw)
    c.font = Font(**base)


def _set_align(c, **kw):
    a = c.alignment
    base = dict(horizontal=a.horizontal, vertical=a.vertical, indent=a.indent, wrap_text=a.wrap_text,
                shrink_to_fit=False, text_rotation=a.text_rotation)
    base.update(kw)
    c.alignment = Alignment(**base)


def _row_has_fill(ws, r, c1, c2, colors):
    return any(_fill_rgb(ws.cell(r, cc)) in colors for cc in range(C.col(c1), C.col(c2) + 1))


# =============================================================================== early
def ungroup_cols(ws):
    """Gruppierte Spaltenbreiten (min..max) in Einzelspalten zerlegen – gleiche Breiten, aber jede Spalte ist
    einzeln adressierbar (core.col_px und Breitenänderungen einzelner Spalten wirken damit korrekt)."""
    from copy import copy
    from openpyxl.utils import get_column_letter
    for key, dim in list(ws.column_dimensions.items()):
        lo, hi = dim.min or 0, dim.max or 0
        if not lo or hi <= lo:
            continue
        dim.max = lo
        for cc in range(lo + 1, hi + 1):
            letter = get_column_letter(cc)
            nd = copy(dim)
            nd.index = letter
            nd.min = nd.max = cc
            ws.column_dimensions[letter] = nd


def cell_rules(ws):
    """Korrektorat, Zahlenformate, Typo-Skala, keine senkrechten Linien (je Zelle)."""
    for c in list(ws._cells.values()):
        if c.number_format in NUMFMT_MAP:
            c.number_format = NUMFMT_MAP[c.number_format]
        if c.value is not None and (ws.title, c.coordinate) not in KEEP_SIZE:
            sz = _map_size(c)
            if sz != c.font.sz:
                _set_font(c, sz=sz)
        b = c.border
        drop_l = b.left is not None and b.left.style in ("thin", "hair", "dotted") and _rgb(b.left.color) in LINE_COLORS
        drop_r = b.right is not None and b.right.style in ("thin", "hair", "dotted") and _rgb(b.right.color) in LINE_COLORS
        if (drop_l or drop_r) and not _is_input(c):
            c.border = Border(left=None if drop_l else b.left, right=None if drop_r else b.right, top=b.top,
                              bottom=b.bottom)
    for ref in RATE_CELLS.get(ws.title, []):
        for c in _cells(ws, ref):
            c.number_format = C.NUMFMT["pct2"]
    for ref in INT_CELLS.get(ws.title, []):
        for c in _cells(ws, ref):
            c.number_format = C.NUMFMT["int"]


def header(ws, spec):
    """Eyebrow (Z. 5) · H1 = Blattname (Z. 6) · Untertitel einzeilig (Z. 7)."""
    c1, c2, eyebrow, title, subtitle = spec
    for r in (5, 6, 7):
        for c in C.iter_cells(ws, 2, r, c2, r):
            if not C.is_formula(c.value):
                c.value = None
            c.fill = C.NOFILL
            c.border = Border()
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row <= 7 and mr.max_row >= 5 and mr.min_col >= 2:
            ws.unmerge_cells(str(mr))
    C.page_header(ws, c1, c2, eyebrow, title, subtitle)
    ws.row_dimensions[4].height = 12
    ws[f"{c1}7"].alignment = C.align("left", "top")


def table_rules(ws, c1, c2, first=8, last=None):
    """Spaltenkopf, Summenzeile und Haarlinie über die volle Tabellenbreite; Köpfe folgen ihren Werten."""
    last = last or ws.max_row
    anchors = _merged_anchor_map(ws)
    band_colors = {C.NAVY, C.TINT} if ws.title not in FORM_SHEETS else {C.NAVY}
    kinds = {}
    for r in range(first, last + 1):
        vals = [ws.cell(r, cc) for cc in range(C.col(c1), C.col(c2) + 1)]
        if not any(v.value is not None for v in vals) and not _row_has_fill(ws, r, c1, c2, band_colors | {C.HEAD}):
            continue
        if any(isinstance(v.value, str) and v.value.startswith("Keine Gewähr") for v in vals):
            break
        fills = {_fill_rgb(v) for v in vals}
        first_cell = ws.cell(r, C.col(c1))
        lb = first_cell.border.left
        if (_fill_rgb(first_cell) == C.NAVY or (_fill_rgb(first_cell) == C.TINT and lb is not None and lb.style == "thick")) \
                and not _is_input(first_cell):
            kinds[r] = "band"
        elif C.HEAD in fills:
            kinds[r] = "head"
        elif (fills & {C.TINT_XL, "F5F8FC", C.TINT}) and any(v.font.b for v in vals if v.value is not None):
            kinds[r] = "sum"
        elif any(v.value is not None for v in vals):
            kinds[r] = "row"
    for r, kind in kinds.items():
        if kind == "band":
            if ws.title in FORM_SHEETS:
                C.band_form(ws, r, c1, c2)
            else:
                C.band_l1b(ws, r, c1, c2)
        elif kind == "head":
            below = next((k for k in range(r + 1, r + 4) if kinds.get(k) in ("row", "sum")), None)
            C.table_head(ws, r, c1, c2)
            for cc in range(C.col(c1), C.col(c2) + 1):
                h = ws.cell(r, cc)
                if h.value is None or below is None:
                    continue
                v = ws.cell(below, cc)
                if (r, cc) in anchors and anchors[(r, cc)][1] != cc:
                    continue
                num = isinstance(v.value, (int, float)) or (C.is_formula(v.value) and _numeric_fmt(v.number_format))
                h.alignment = C.align("right" if num else "left", "center", 1)
                if isinstance(h.value, str) and h.data_type == "s":
                    h.value = h.value.upper()
        elif kind == "sum":
            for c in C.iter_cells(ws, c1, r, c2, r):
                if _is_input(c):
                    continue
                c.fill = C.fill(C.TINT)
                c.border = Border(top=C.side("thin", C.BLUE), bottom=C.side("hair", C.LINE))
                _set_font(c, b=True, color=C.NAVY)
        else:
            for c in C.iter_cells(ws, c1, r, c2, r):
                if _is_input(c):
                    continue
                b = c.border
                c.border = Border(top=b.top if (b.top is not None and b.top.style) else None,
                                  bottom=C.side("hair", C.LINE))
        if kind in ("row", "sum"):
            edges(ws, r, c1, c2, anchors)
    return kinds


def _is_num(c):
    return isinstance(c.value, (int, float)) and not isinstance(c.value, bool) or \
        (C.is_formula(c.value) and _numeric_fmt(c.number_format))


def edges(ws, r, c1, c2, anchors):
    """P1-06: Text links mit Einzug 1, Zahlen rechts mit Einzug 1 (zentrierte Zellen bleiben zentriert).
    Beschriftungen der ersten Spalte brechen um statt abgeschnitten zu werden."""
    for cc in range(C.col(c1), C.col(c2) + 1):
        c = ws.cell(r, cc)
        if c.value is None or ((r, cc) in anchors and anchors[(r, cc)][1] != cc):
            continue
        if c.alignment.horizontal == "center":
            continue
        if _is_num(c):
            _set_align(c, horizontal="right", indent=1)
        else:
            _set_align(c, horizontal="left", indent=1)
        if cc == C.col(c1) and _static(c):
            _set_align(c, wrap_text=True)


NAME_VALUES = {}


def resolve_names(wb):
    """Namen → aktueller Zellwert (für die Zeilenhöhe von Anzeigeformeln wie =Rechtsform)."""
    NAME_VALUES.clear()
    for n, d in wb.defined_names.items():
        try:
            for title, coord in d.destinations:
                if title in wb.sheetnames and ":" not in coord:
                    v = wb[title][coord.replace("$", "")].value
                    if v is not None and not C.is_formula(v):
                        NAME_VALUES[n.lower()] = v
        except Exception:
            continue


def _display_text(c):
    v = c.value
    if C.is_formula(v):
        m = re.fullmatch(r"=([A-Za-z_][A-Za-z0-9_.]*)", v.strip())
        return NAME_VALUES.get(m.group(1).lower()) if m else None
    return v


def _lines(text, width_px, size, bold):
    """Zeilenzahl mit der Calibri-Laufweite aus core.text_px (in der Vorschau gemessen ≈ 3 % zu breit)."""
    import math
    n = 0
    for para in str(text).split("\n"):
        n += max(1, math.ceil(0.97 * C.text_px(para, size, bold) / max(width_px, 20)))
    return n


def fit_row(ws, row, c1, c2, base=C.H_ROW, line_pt=12, pad=6, max_lines=3):
    """Wie core.fit_row_height, zählt aber auch Anzeigeformeln (=Name) mit ihrem aktuellen Text."""
    merged = {}
    for mr in ws.merged_cells.ranges:
        if mr.min_row == row == mr.max_row:
            merged[mr.min_col] = mr.max_col
    need = 1
    for cc in range(C.col(c1), C.col(c2) + 1):
        cell = ws.cell(row, cc)
        text = _display_text(cell)
        if text is None or not (cell.alignment and cell.alignment.wrap_text):
            continue
        w = C.span_px(ws, cc, merged.get(cc, cc)) - 7 * (cell.alignment.indent or 0) - 6
        need = max(need, _lines(text, w, cell.font.sz or C.T_BODY, bool(cell.font.b)))
    need = min(need, max_lines)
    ws.row_dimensions[row].height = base if need == 1 else need * line_pt + pad
    return need


def row_grid(ws, kinds, c1, c2, line_pt=12, pad=6, base=C.H_ROW, keep=None):
    """Zeilenraster: einzeilig 18 pt, zweizeilig 30 pt, dreizeilig 42 pt; einzeilig mittig, mehrzeilig oben.
    keep: Höhen einer bereits gerasterten Nachbartabelle derselben Zeilen (es gilt das Maximum)."""
    for r, kind in kinds.items():
        if kind == "head":
            h = C.H_HEAD
        elif kind == "band":
            h = C.H_BAND_B + 2 if ws.title in FORM_SHEETS else C.H_BAND_B
        else:
            fit_row(ws, r, c1, c2, base=base, line_pt=line_pt, pad=pad)
            h = ws.row_dimensions[r].height
        if keep and r in keep:
            h = max(h, keep[r])
        ws.row_dimensions[r].height = h
        if kind in ("row", "sum"):
            for c in C.iter_cells(ws, c1, r, c2, r):
                if c.value is not None:
                    _set_align(c, vertical="center" if h <= base else "top")


def page_footer(ws, c1, c2, row=None):
    """Fuß: Leerzeile nach dem Inhalt, Oberkante, zwei Zeilen 8 pt grau (core.footer)."""
    if row is None:
        for c in ws._cells.values():
            if isinstance(c.value, str) and c.value.startswith("Keine Gewähr") and c.data_type == "s":
                row = c.row
                break
    if row is None:
        return
    for r in (row, row + 1):
        for mr in list(ws.merged_cells.ranges):
            if mr.min_row <= r <= mr.max_row:
                ws.unmerge_cells(str(mr))
        for c in C.iter_cells(ws, c1, r, c2, r):
            c.value = None
    C.footer(ws, row, c1, c2)
    if row - 1 > 8:
        ws.row_dimensions[row - 1].height = 14


UNIT_TEXT = {"€": None, "%": None, "€/Monat": "pro Monat", "€ p. a.": "pro Jahr", "€/m²": "je m²",
             "€/m² p. a.": "je m² p. a.", "% Darlehen": "vom Darlehen", "% der Miete": "der Miete",
             "% Verkaufspreis": "vom Verkaufspreis", "fach": None}


def forms_pre(ws):
    """Vor dem Tabellenstil: Einheiten, Kopftexte, Jahresspalte (Eingaben / Konfiguration)."""
    if ws.title == "Eingaben":
        for r in range(11, 146):
            e = ws.cell(r, 5)
            if _static(e) and e.value in UNIT_TEXT:
                if e.value == "fach":
                    ws.cell(r, 3).number_format = C.NUMFMT["mult1"]
                new = UNIT_TEXT[e.value]
                e.value = None if new is None else new
                if new is not None:
                    e.data_type = "s"
        for r in range(9, 146):
            c, d = ws.cell(r, 3), ws.cell(r, 4)
            if _fill_rgb(c) == C.HEAD and d.value is None and isinstance(c.value, str):
                C.safe_merge(ws, "C", r, "D", r)
                C.set_text(c, "WERT")
                c.alignment = C.align("right", "center", 1)
        for c in C.iter_cells(ws, "B", 71, "L", 71):  # „Summe Maßnahmen“ ist eine Summenzeile
            if c.value is not None:
                _set_font(c, b=True)
            if not _is_input(c):
                c.fill = C.fill(C.TINT)
    if ws.title == "Konfiguration":
        for r in range(45, 54):
            c = ws.cell(r, 2)
            c.number_format = C.NUMFMT["year"]
            _set_align(c, horizontal="left", indent=1)
        for r in range(10, 81):
            e = ws.cell(r, 5)
            if _static(e):
                _set_align(e, wrap_text=True)


def forms_defaults(ws):
    """Eingaben / Konfiguration: Kopf über Zahlen rechts, Datumsspalten zentriert, Jahre links (P1-06)."""
    if ws.title == "Eingaben":
        if ws["D10"].value is None:
            C.safe_merge(ws, "C", 10, "D", 10)
            C.set_text(ws["C10"], "WERT")
            ws["C10"].alignment = C.align("right", "center", 1)
        for r in range(11, 146):
            e = ws.cell(r, 5)
            if e.value is not None:
                _set_align(e, horizontal="left", indent=1)
    if ws.title == "Konfiguration":
        for r in range(10, 26):
            _set_align(ws.cell(r, 4), horizontal="center", indent=0)
        for coord in ("G27",):
            _set_align(ws[coord], horizontal="left", indent=0)
        g43 = ws["G43"]  # Listenkopf wie G9 / G15
        g43.fill = C.fill(C.HEAD)
        g43.font = C.font(C.T_MICRO, True, C.BLUE)
        g43.border = Border(bottom=C.side("thin", C.ACCENT))
        g43.alignment = C.align("left", "center", 1)
        if isinstance(g43.value, str):
            g43.value = g43.value.upper()
        for coord, text, h in (("C27", "WERT", "right"), ("E27", "FORMEL / ERLÄUTERUNG", "left"),
                               ("C56", "WERT", "right"), ("E56", "ERLÄUTERUNG", "left")):
            c = ws[coord]
            title = ws.cell(c.row, 2)
            if C.text_px(title.value, C.T_BODY, True) + 14 > C.span_px(ws, 2, c.column - 1):
                continue  # Bandtitel braucht den Platz
            if c.value is None:
                C.set_text(c, text)
                c.font = C.font(C.T_MICRO, True, C.SKY)
                c.alignment = C.align(h, "center", 1)
        if all(ws.cell(r, cc).value is None for r in (82, 83) for cc in range(2, 8)):
            page_footer(ws, "B", "G", 82)


NOTE_COLS = {"Eingaben": ("E", 11, 146), "Haushaltsrechnung": ("E", 18, 47), "Vermögensaufstellung": ("E", 10, 34),
             "Konfiguration": ("E", 10, 80)}


def note_cols(ws):
    """Einheiten- und Kommentarspalten: links mit Einzug 1, 9 pt grau (P1-06 (6))."""
    spec = NOTE_COLS.get(ws.title)
    if not spec:
        return
    letter, r1, r2 = spec
    for r in range(r1, r2 + 1):
        c = ws[f"{letter}{r}"]
        if c.value is None or _is_input(c):
            continue
        _set_align(c, horizontal="left", indent=1)
        if not c.font.b:
            _set_font(c, sz=C.T_SMALL, color=C.MUTED)


def bank_defaults(ws):
    """HH / VA: Ergebniszeilen, kritische Puffer rot/grün, Blattschutz mit entsperrten Eingaben (P2-08, P3-02)."""
    from openpyxl.formatting.rule import FormulaRule
    if ws.title == "Haushaltsrechnung":
        res, crit = 46, ("C46:D46", "$C$46")
    elif ws.title == "Vermögensaufstellung":
        res, crit = 32, ("C34", "$C$34")
    else:
        return
    if ws.title == "Haushaltsrechnung" and isinstance(ws["E46"].value, str) and ws["E46"].data_type == "s":
        C.set_text(ws["E46"], "Banken erwarten > 0 nach neuem Kapitaldienst")
    for c in C.iter_cells(ws, "B", res, "E", res):
        c.fill = C.fill(C.TINT)
        c.border = Border(top=C.side("thin", C.NAVY), bottom=C.side("double", C.NAVY))
        if c.value is not None and c.column <= 4:
            _set_font(c, sz=C.T_H3, b=True, color=C.NAVY)
            _set_align(c, vertical="center")
    ws.row_dimensions[res].height = 21
    ref, cell = crit
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{cell}<0"], stopIfTrue=True,
                                                   font=Font(color=C.RED, bold=True), fill=C.fill(C.RED_BG)))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"ISNUMBER({cell})"], stopIfTrue=True,
                                                   font=Font(color=C.GREEN, bold=True)))
    for r in range(9, res + 2):  # Kommentar-/Nachweisspalte neben Eingaben bleibt beschreibbar
        if ws.cell(r, 3).protection.locked is False:
            ws.cell(r, 5).protection = Protection(locked=False)
    ws.protection.sheet = True
    ws.protection.objects = True
    ws.protection.scenarios = True


def calc_colors(ws):
    """Formelergebnisse auf Rechenblättern in 1A1D21 – Blau nur für übernommene Eingaben (P2-08)."""
    for c in ws._cells.values():
        if c.row <= 7 or c.hyperlink is not None or not C.is_formula(c.value):
            continue
        if _rgb(c.font.color) == C.BLUE and not _is_input(c):
            _set_font(c, color=C.INK)


def sensitivity_defaults(ws):
    from openpyxl.formatting.rule import FormulaRule
    for ref in ("D11", "D62"):
        cell = "$" + re.sub(r"(\d+)", r"$\1", ref)
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f"{cell}<0"], stopIfTrue=True,
                                                       font=Font(color=C.RED, bold=True), fill=C.fill(C.RED_BG)))


def early(wb):
    for ws in wb.worksheets:
        ungroup_cols(ws)
        cell_rules(ws)
        ws.sheet_properties.tabColor = tab_color(ws.title)
    for name, spec in HEADERS.items():
        if name in wb.sheetnames:
            header(wb[name], spec)
    resolve_names(wb)
    for name in FORM_SHEETS:
        if name in wb.sheetnames:
            forms_pre(wb[name])
    for name in NOTE_COLS:
        if name in wb.sheetnames:
            note_cols(wb[name])
    for name, spans in TABLES.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        done = {}
        for c1, c2 in spans:
            kinds = table_rules(ws, c1, c2, first=8)
            if name == "Hinweise":
                row_grid(ws, kinds, c1, c2, line_pt=13, pad=10)
            else:
                row_grid(ws, kinds, c1, c2, keep=done)
            for r in kinds:
                done[r] = ws.row_dimensions[r].height or C.H_ROW
    for name, (c1, c2) in CONTENT.items():
        if name in wb.sheetnames and name != "Konfiguration":
            page_footer(wb[name], c1, c2)
    for name in FORM_SHEETS:
        if name in wb.sheetnames:
            forms_defaults(wb[name])
    for name in ("Haushaltsrechnung", "Vermögensaufstellung"):
        if name in wb.sheetnames:
            bank_defaults(wb[name])
    for name in CALC_SHEETS:
        if name in wb.sheetnames:
            calc_colors(wb[name])
    if "Sensitivität" in wb.sheetnames:
        sensitivity_defaults(wb["Sensitivität"])


# =============================================================================== final
def final(wb):
    for ws in wb.worksheets:
        normalise(ws)
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                if rule.dxf is not None and rule.dxf.font is not None:
                    rule.dxf.font.name = None


def _rich_runs(c):
    """Rich-Text: Sekundärtext ≤ 9 pt nie in 8A9099, nur Calibri, Mindestgröße 8 pt."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    v = c.value
    if not isinstance(v, CellRichText):
        return
    for part in v:
        if not isinstance(part, TextBlock) or part.font is None:
            continue
        f = part.font
        col = _rgb(f.color) if f.color is not None else None
        if (f.sz or 11) <= 9 and col == C.MUTED2:
            f.color = C.MUTED
        if f.sz is not None and f.sz < 8:
            f.sz = 8
        if f.rFont in OLD_FONTS - {None}:
            f.rFont = C.SANS


def normalise(ws):
    for c in ws._cells.values():
        _rich_runs(c)
        al = c.alignment
        f = c.font
        # 1) kein „An Zellgröße anpassen“; Einzug nur mit links/rechts
        if al.shrink_to_fit or (al.indent and al.horizontal not in ("left", "right", "distributed")):
            h = al.horizontal
            if al.indent and h not in ("left", "right", "distributed"):
                num = isinstance(c.value, (int, float)) or (C.is_formula(c.value) and _numeric_fmt(c.number_format))
                h = "right" if num else "left"
            _set_align(c, horizontal=h, shrink_to_fit=False)
        if c.value is None:
            if f.name in OLD_FONTS:  # leere Zellen: beim Tippen erscheint sonst die Altschrift
                _set_font(c, name=C.SANS)
            continue
        # 2) Mindestschrift 8 pt; 3) Sekundärtext ≤ 9 pt nie in 8A9099 (außer inaktive Felder)
        kw = {}
        if (f.sz or 11) < 8:
            kw["sz"] = 8
        if (f.sz or 11) <= 9 and _rgb(f.color) == C.MUTED2 and _fill_rgb(c) != C.INACTIVE_BG:
            kw["color"] = C.MUTED
        # 3b) Typo-Skala (P2-07): Zwischenstufen abbilden – Ausnahmen: Auswahlfeld, Symbole, Textergebnisse S-Seiten
        sz0 = f.sz or 11
        if sz0 not in SCALE and (ws.title, c.coordinate) not in KEEP_SIZE and not _symbol(c.value) \
                and not (sz0 == 9.5 and STEP_RE.match(ws.title)):
            kw["sz"] = _map_size(c)
        # 4) nur Calibri
        if f.name in OLD_FONTS:
            kw["name"] = C.SANS
        if kw:
            _set_font(c, **kw)
        # 5) Normzitate in Versaltexten
        if _static(c) and sum(ch.isupper() for ch in c.value) > 0.6 * max(1, sum(ch.isalpha() for ch in c.value)):
            v = c.value
            for pat, rep in VERSAL_FIX:
                v = pat.sub(rep, v)
            if v != c.value:
                C.set_text(c, v)


# =============================================================================== Druck (P1-23) und Register (P3-02)
MARGINS = dict(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.3, footer=0.3)


def _footer_row(ws, c1=1, c2=60):
    rows = [c.row for c in ws._cells.values()
            if isinstance(c.value, str) and c.value in (C.FOOTER_2,) and C.col(c1) <= c.column <= C.col(c2)]
    if rows:
        return max(rows)
    rows = [c.row for c in ws._cells.values() if isinstance(c.value, str) and c.value.startswith("MM HOLDING GMBH ·")]
    return max(rows) if rows else None


def _print(ws, area, orient="landscape", w=1, h=0, scale=None, rows=None, cols=None, breaks=(), col_breaks=(),
           over_then_down=False):
    ws.print_area = area
    ws.page_setup.orientation = orient
    if scale:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
        ws.page_setup.scale = scale
        ws.page_setup.fitToWidth = None
        ws.page_setup.fitToHeight = None
    else:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = w
        ws.page_setup.fitToHeight = h
        ws.page_setup.scale = None
    if rows:
        ws.print_title_rows = rows
    if cols:
        ws.print_title_cols = cols
    ws.row_breaks = RowBreak()
    for r in breaks:
        ws.row_breaks.append(Break(id=r - 1))
    ws.col_breaks = ColBreak()
    for cc in col_breaks:
        ws.col_breaks.append(Break(id=C.col(cc)))
    if over_then_down:
        ws.page_setup.pageOrder = "overThenDown"


def _chart_right_col(ws):
    cols = []
    for ch in getattr(ws, "_charts", []):
        to = getattr(ch.anchor, "to", None)
        if to is not None:
            cols.append(to.col + (1 if to.colOff else 0))
    return max(cols) if cols else 0


def print_setup(ws):
    t = ws.title
    fr = _footer_row(ws)
    if STEP_RE.match(t):
        _print(ws, f"A1:I{_footer_row(ws, 1, 9) or fr or 60}")
    elif t == "Start":
        _print(ws, f"A1:H{max(_footer_row(ws, 1, 8) or 0, 65)}")
        if ws.freeze_panes is None:
            ws.freeze_panes = "A4"
    elif t == "Dashboard":
        last = _footer_row(ws, 1, 18) or ws.max_row
        _print(ws, f"A1:R{last}", "portrait", 1, 1)
    elif t == "Leitfaden":
        _print(ws, f"A1:I{_footer_row(ws, 1, 9) or 53}")
    elif t == "Cockpit":
        _print(ws, f"A1:L{_footer_row(ws, 1, 12) or 76}")
    elif t == "Eingaben":
        _print(ws, f"A1:L{_footer_row(ws, 1, 12) or 148}", rows="9:10", breaks=(48, 75, 95, 119))
    elif t == "Diagramme":
        _print(ws, "A1:P136", breaks=(51, 92))
    elif t in ("Steuern", "Projektion", "Finanzierung"):
        last = _footer_row(ws) or ws.max_row
        rows = "40:42" if t == "Steuern" else "8:10"
        _print(ws, f"A1:AQ{last}", scale=70, rows=rows, cols="A:C", breaks=(40,) if t == "Steuern" else (),
               col_breaks=("M", "W", "AG"), over_then_down=True)
    elif t == "AfA-Vergleich":
        _print(ws, f"A1:Q{_footer_row(ws, 1, 17) or 75}", breaks=(52,))
    elif t == "Sensitivität":
        _print(ws, "$B$5:$M$63", breaks=(43,))
    elif t == "Bankgespräch":
        _print(ws, f"$B$5:$I${max(_footer_row(ws, 2, 9) or 0, 67)}", "portrait", 1, 1)
    elif t in ("Haushaltsrechnung", "Vermögensaufstellung"):
        last = _footer_row(ws, 2, 5) or (50 if t == "Haushaltsrechnung" else 37)
        right = _chart_right_col(ws)
        if right > 5:  # Diagramm rechts neben der Tabelle: quer, sonst hoch (B:E)
            from openpyxl.utils import get_column_letter
            _print(ws, f"$B$5:${get_column_letter(right)}${last}", "landscape", 1, 1)
        else:
            _print(ws, f"$B$5:$E${last}", "portrait", 1, 1)
    elif t == "Hinweise":
        _print(ws, f"A1:E{_footer_row(ws, 1, 5) or 39}", rows="8:9")
    elif t == "Konfiguration":
        _print(ws, f"A1:G{_footer_row(ws, 1, 7) or 83}", rows="8:9")


def page_setup(wb):
    if "Dashboard" in wb.sheetnames:  # entsteht erst nach final(): dieselben Sicherheitsregeln
        ws = wb["Dashboard"]
        normalise(ws)
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                if rule.dxf is not None and rule.dxf.font is not None:
                    rule.dxf.font.name = None
    for ws in wb.worksheets:
        ps = ws.page_setup
        ps.paperSize = 9
        for k, v in MARGINS.items():
            setattr(ws.page_margins, k, v)
        ws.print_options.horizontalCentered = True
        ws.oddHeader.right.text = None if ws.title in BANK_SHEETS else "&A"
        ws.oddHeader.right.size = 8
        ws.oddHeader.right.color = C.MUTED
        ws.oddHeader.right.font = "Calibri,Regular"
        ws.oddFooter.left.text = "MM Holding GmbH · Immobilien-Kalkulation · &A"
        ws.oddFooter.right.text = "Seite &P von &N"
        for part in (ws.oddFooter.left, ws.oddFooter.right):
            part.size = 8
            part.font = "Calibri,Regular"
            part.color = C.MUTED
        try:
            print_setup(ws)
        except Exception as exc:  # ein Blatt darf den Druck-Setup der übrigen nicht verhindern
            print(f"WARNUNG Druck {ws.title}: {exc!r}")
        ws.sheet_properties.tabColor = tab_color(ws.title)
        if ws.title in PRESENTATION or STEP_RE.match(ws.title):
            ws.sheet_view.showRowColHeaders = False
        if ws.protection.sheet:
            ws.protection.objects = True
    reorder(wb)


def reorder(wb):
    """Blattreihenfolge wie der Ablauf (P3-02); Druckbereiche sind blattlokal und wandern mit."""
    steps = [n for n in wb.sheetnames if STEP_RE.match(n)]
    want = []
    for n in ORDER:
        if n is None:
            continue
        want.append(n)
        if n == "Leitfaden":
            want.extend(sorted(steps))
    order = [wb[n] for n in want if n in wb.sheetnames]
    order += [ws for ws in wb.worksheets if ws not in order]
    wb._sheets = order
    wb.active = 0
