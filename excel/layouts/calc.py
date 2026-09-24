"""Rechenblätter Steuern, Projektion, Finanzierung, AfA-Vergleich (Agent D) – Runde 4: EINE Komponentensprache.

Runde 4 (letzte Runde): Kopfschablone mit Rücksprung/Objekt/„Erstellt für“ rechts (P2-03), Werkzeugzeile unter dem
Jahreskopf („Eingaben dazu …“ links, Legende rechts), Verkaufsjahr C8D7EB im Kopf (P1-11), Fünfjahreslinie als linker
Rahmen der Folgespalte, Nullwerte grau (P1-11), Summen-/Kennzahlstufen (P2-06), Abschnittslinks „↑ nach oben“ an der
rechten Kante des Erstbildschirms (P3-08), Fußzeile Zurück · Übersicht · Weiter in Reiterfolge (P3-08),
AfA-Kacheln 18/30/20 (P1-02), Feinschliff Steuern/AfA (P3-06, P3-07).

Alle Bausteine kommen aus core.py (CORE_API.md):
  Seitenkopf     C.page_header (Überzeile „REITER  ›  BLATT“, H1 22 pt, Untertitel 10 pt) + Legende mit Marken ● ■
  Abschnitte     C.section(level=1) – E7EEF7, Akzentkante, 12,5 pt Titelschreibung, Meta rechts 8 pt
                 C.section(level=2) – Tabellenköpfe (EEF3FA, 8 pt Versalien)
  Summenstufen   C.sum_row('deduct' | 'sub' | 'final' | 'kpi' | 'kpi_band' | 'plain' | 'memo')  (P2-06)
  Kacheln        C.tile (calm, 18/30/20) · Einordnung C.callout_box                          (P1-02)
  Links          C.text_link „↑ nach oben“ · C.btn_row Zurück · Übersicht · Weiter           (P3-08)
  Formate        C.NUMFMT (pct1/pct2/years_n/status_dot …)                                   (P2-12)
  CF-Abschluss   C.cf_close                                                                  (LibreOffice-Einzug)
Blattspezifisch: Jahreskopf (Navy/Blau), 5-Jahres-Raster, Verkaufsjahr-Spalte, Rot-Regel für Ergebniszeilen
(P2-01), „entfällt“ = „–“ in 98A2B3 kursiv inkl. Beschriftung (P2-02), Exit-Zeile Jahr 0 (P2-19),
Gliederung auf Steuern (P3-04), gleichmäßiger Zeilenrhythmus 16 pt (P3-05).

Nur Darstellung: keine Formel, auf die etwas verweist, wird verändert; neue Formeln sind reine Anzeigeformeln
in leeren Zellen (bzw. Beschriftungen, auf die nichts verweist).
"""
from copy import copy

from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Border, Font
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.properties import Outline

import core as C
from core import (ACCENT, AMBER, BLUE, GREEN, INK, INK2, LINE, LINE2, MUTED, NAVY, NOFILL, NUMFMT, RED,
                  T_BODY, T_MICRO, T_SMALL, TINT, WHITE, align, fill, font, side)

# ------------------------------------------------------------------------------------------------ Konstanten
FIRST, LAST = 4, 43                 # Jahresspalten D … AQ (Jahr 1 … 40)
GRID = (9, 14, 19, 24, 29, 34, 39)  # I, N, S, X, AC, AH, AM: linke Kante VOR Jahr 6, 11, … 36 (P1-11: Luft an den Zahlen)
FADED = "98A2B3"                    # „entfällt“ / Nullwerte (P1-11, P2-02)
H_DATA, H_YEAR, H_GAP = 16, 20, C.H_GAP
H_TOOL = 20                         # Werkzeugzeile unter dem Jahreskopf („Eingaben dazu …“ · Legende)
YEAR_FMT = '"Jahr "0'
EUR, NUM, PCT1, PCT2 = C.EUR, NUMFMT["num"], C.PCT1, C.PCT2
YESNO = '[=1]"Ja";[=0]"Nein";0'     # P3-06: „Ja/Nein“ großgeschrieben wie die Auswahllisten
DASH_FMT = '"–";"–";"–";"–"'        # bedingtes Zahlenformat „entfällt“
EDGE = "M"                          # rechte Kante des Erstbildschirms (Jahr 10) auf den 40-Jahres-Blättern
UP = "↑ nach oben"                  # P3-08: Blätter ohne Sprungleiste
# Reiterfolge (P3-08): Zurück / Weiter der Fußzeile
CHAIN = {"Projektion": ("Eingaben", "Eingaben", "Steuern", "Steuer-Tabelle"),
         "Steuern": ("Projektion", "Projektion", "AfA-Vergleich", "AfA-Vergleich"),
         "AfA-Vergleich": ("Steuern", "Steuer-Tabelle", "Finanzierung", "Tilgungsplan"),
         "Finanzierung": ("AfA-Vergleich", "AfA-Vergleich", "Sensitivität", "Sensitivität")}
CONTEXT = ("=Obj_Name", '=IFERROR(IF(Erstellt_fuer="","","Erstellt für "&Erstellt_fuer),"")')
ZERO_FONT = Font(color=FADED, bold=False, italic=False)


# ------------------------------------------------------------------------------------------------ Hilfen
def _col(c):
    return C.col(c)


def _ungroup_cols(ws):
    """Gruppierte Spaltenbreiten (<col min max>) in Einzelspalten auflösen, damit Breiten je Spalte setzbar sind."""
    dims = ws.column_dimensions
    for key, d in list(dims.items()):
        lo, hi = d.min, d.max
        if lo and hi and hi > lo:
            for ci in range(lo + 1, hi + 1):
                letter = get_column_letter(ci)
                nd = ColumnDimension(ws, index=letter, width=d.width, hidden=d.hidden, customWidth=d.customWidth,
                                     outlineLevel=d.outlineLevel, collapsed=d.collapsed, min=ci, max=ci)
                nd._style = copy(d._style)
                dims[letter] = nd
            d.max = lo


def _width(ws, letter, w):
    d = ws.column_dimensions[letter]
    ci = _col(letter)
    d.min, d.max = ci, ci
    d.width = w


def _rule(ws, ref, formula, fnt=None, fil=None, numfmt=None, border=None):
    """Bedingte Formatierung (Ausdruck) mit optionalem bedingtem Zahlenformat (dxf numFmt)."""
    nf = NumberFormat(numFmtId=300, formatCode=numfmt) if numfmt else None
    dxf = DifferentialStyle(font=fnt, fill=fil, numFmt=nf, border=border)
    r = Rule(type="expression", dxf=dxf)
    r.formula = [formula]
    ws.conditional_formatting.add(ref, r)


def cf_rules(ws, ref, rules, base, year_cond=None, year_fill=TINT):
    """Regeln [(bedingung, Font, numfmt|None[, fill])] für einen Bereich, danach eine Auffangregel mit der Grundfarbe.

    - Mit year_cond steht vor jeder Regel eine Kombination (Regel + Verkaufsjahr-Fläche), damit die
      Spaltenmarkierung auch dort durchläuft, wo nur die erste zutreffende Regel greift (LibreOffice).
    - Die Auffangregel (immer wahr, nur Grundfarbe) ändert in Excel nichts; LibreOffice verliert sonst in
      Zellen mit nicht zutreffender Regel den Einzug (Vorschau zeigt versetzte Zahlen)."""
    for spec in rules:
        cond, fnt = spec[0], spec[1]
        nf = spec[2] if len(spec) > 2 else None
        fl = spec[3] if len(spec) > 3 else None
        if year_cond:
            _rule(ws, ref, f"AND({cond},{year_cond})", fnt, fill(year_fill), nf)
        _rule(ws, ref, cond, fnt, fl, nf)
    if year_cond:
        _rule(ws, ref, year_cond, base, fill(year_fill))
    _rule(ws, ref, "TRUE", base)


def zero(row, first="D"):
    """P1-11: Nullwerte („–“) grau 98A2B3, nicht fett – mit Vorrang vor Summen-/Negativregeln."""
    return (f"AND(ISNUMBER({first}{row}),{first}{row}=0)", ZERO_FONT)


def row_rules(ws, spec, year_cond):
    """Je Datenzeile EIN Regelsatz: [entfällt …] + Nullwert + [Negativ-Rot …] + Grundfarbe, jeweils mit Verkaufsjahr.
    spec: {zeile: dict(base=Farbe, pre=[…], rules=[…], fill=TINT|MIST)}."""
    for r, sp in spec.items():
        rules = list(sp.get("pre", [])) + [zero(r)] + list(sp.get("rules", []))
        cf_rules(ws, f"D{r}:AQ{r}", rules, Font(color=sp.get("base", INK)), year_cond,
                 year_fill=sp.get("fill", TINT))


def neg_rule(r, bold=None):
    return (f"D{r}<0", Font(color=RED, bold=bold))


def _label(ws, row, text, col=2):
    c = ws.cell(row, col)
    C.set_text(c, text)
    return c


def _clear_row(ws, row, c1, c2):
    for c in C.iter_cells(ws, c1, row, c2, row):
        c.fill = NOFILL
        c.border = Border()


def _calibri(ws):
    """Sicherheitsnetz: Vorlagen-Schriften (Inter/Fraunces) in Band-, Lücken- und Leerzellen auf Calibri."""
    for c in list(ws._cells.values()):
        f = c.font
        if f is not None and f.name not in (C.SANS, None):
            nf = copy(f)
            nf.name = C.SANS
            c.font = nf


def _band_link(ws, row, col=EDGE, text=UP, target_sheet=None, target_cell=None, tooltip=None):
    """P3-08: Abschnittslink rechtsbündig an der rechten Kante des Erstbildschirms (nie über Spalte C),
    8,5 pt 1D4F8A regulär – „↑ nach oben“ (Blätter ohne Sprungleiste) bzw. „Eingaben dazu …  ›“."""
    cell = ws[f"{col}{row}"]
    C.text_link(cell, text, target_sheet or ws.title, target_cell, size=C.T_LABEL, bold=False,
                tooltip=tooltip or ("Zum Seitenanfang" if text == UP else None))
    cell.font = font(C.T_LABEL, False, BLUE)
    cell.alignment = align("right", "center", 1)
    return cell


# ------------------------------------------------------------------------------------------------ Seitenkopf
def header(ws, eyebrow, title, subtitle, edge=EDGE, back=None):
    """Kopfschablone (C.page_header, P2-03): Z. 4–7 feste Höhen 12/18/30/21,75 pt. Rechts bündig an `edge`
    (Erstbildschirm M bzw. Inhaltskante Q): Z. 5 Rücksprung „‹  Dashboard“ (ohne Unterreiter) · Z. 6 Objekt
    („Beispiel: …“, fett Navy) · Z. 7 „Erstellt für …“. „Eingaben dazu …“ und die Legende stehen in der Werkzeugzeile."""
    last = max(LAST, _col(edge))
    for r in (5, 6, 7):
        for mr in list(ws.merged_cells.ranges):
            if mr.min_row == r and mr.min_col >= 4:
                ws.unmerge_cells(str(mr))
        for c in C.iter_cells(ws, 3, r, last, r):
            c.border = Border()
            c.fill = NOFILL
            c.hyperlink = None
            if c.column >= 4 and not C.is_formula(c.value):
                c.value = None
    C.page_header(ws, "B", edge, eyebrow, title, subtitle, context=CONTEXT, context_col=edge, back=back,
                  back_col=edge)


def legend(*items, size=T_SMALL):
    """Legende als Rich-Text: items = (text, marke|None, markenfarbe) – Marke ● / ■ in Farbe, Text 5B6068."""
    parts = []
    for i, (text, mark, color) in enumerate(items):
        if i:
            parts.append(("   ·   ", size, False, MUTED))
        if mark:
            parts.append((mark + " ", size, True, color))
        parts.append((text, size, False, MUTED))
    return C.rich(parts)


def year_legend(*extra):
    """Legende der Jahrestabellen (8,5 pt): Verkaufsjahr-Marke wie im Kopf (C8D7EB), Rot-Regel, Zusätze."""
    return legend(("Verkaufsjahr", "■", C.MIST), ("negatives Ergebnis", "●", RED), *extra, size=C.T_LABEL)


def toolbar(ws, row, link_text, target, legend_rich):
    """Werkzeugzeile direkt unter dem Jahreskopf: links „Eingaben dazu: Sxx …  ›“ (9 pt fett 1D4F8A, in den fixierten
    Spalten), rechts bündig an der Kante des Erstbildschirms die Legende (8,5 pt)."""
    _clear_row(ws, row, 2, LAST)
    for c in C.iter_cells(ws, 2, row, LAST, row):
        if not C.is_formula(c.value):
            c.value = None
    b = ws.cell(row, 2)
    C.text_link(b, link_text, target, size=T_SMALL, bold=True, tooltip=f"Eingaben zu diesem Blatt: {C.sheet_label(target)}")
    b.alignment = align("left", "center", 1)
    lg = ws[f"{EDGE}{row}"]
    lg.value = legend_rich
    lg.font = font(C.T_LABEL, False, MUTED)
    lg.alignment = align("right", "center")
    ws.row_dimensions[row].height = H_TOOL


# ------------------------------------------------------------------------------------------------ Jahrestabelle
def year_head(ws, r_year, r_dup, r_cal):
    """Zweizeiliger Kopf: r_year „Jahr n“ (Navy, 9 pt fett weiß), r_dup ausgeblendet, r_cal Kalenderjahr (Blau)."""
    for c in C.iter_cells(ws, 2, r_year, LAST, r_year):
        c.fill = fill(NAVY)
        c.border = Border()
        if c.column >= FIRST:
            c.number_format = YEAR_FMT
            c.font = font(T_SMALL, True, WHITE)
            c.alignment = align("right", "center", 1)
    C.set_text(ws.cell(r_year, 2), "Jahr der Kalkulation")
    ws.cell(r_year, 2).font = font(T_SMALL, True, WHITE)
    ws.cell(r_year, 2).alignment = align("left", "center", 1)
    ws.cell(r_year, 3).value = None
    ws.row_dimensions[r_year].height = H_YEAR
    C.hide_rows(ws, r_dup, r_dup)
    for c in C.iter_cells(ws, 2, r_cal, LAST, r_cal):
        c.fill = fill(BLUE)
        c.border = Border(bottom=side("medium", ACCENT))
        if c.column >= FIRST:
            c.number_format = NUMFMT["year"]
            c.font = font(T_SMALL, False, WHITE)
            c.alignment = align("right", "center", 1)
    C.set_text(ws.cell(r_cal, 2), "Kalenderjahr")
    ws.cell(r_cal, 2).font = font(T_SMALL, False, WHITE)
    ws.cell(r_cal, 2).alignment = align("left", "center", 1)
    ws.row_dimensions[r_cal].height = H_YEAR


def band(ws, row, title, c2=LAST, years_from=None, meta=None, extra=None, up=None):
    """Abschnittskopf Ebene 1 (C.section) über die Tabellenbreite. years_from: Kalenderjahre (=D$42) in D:AQ als
    echter Spaltenkopf 9 pt fett Navy (P3-07). up: Spalte des „↑ nach oben“-Links (P3-08)."""
    formula = isinstance(title, str) and title.startswith("=")
    C.section(ws, row, "B", c2, None if formula else title, level=1, meta=meta)
    if formula:
        ws.cell(row, 2).value = title
    elif extra:
        # Zusatz inline hinter dem Titel (rechts stehen Sprung-/Eingabelinks bzw. Jahreszahlen)
        ws.cell(row, 2).value = C.rich([(title, C.T_H3, True, NAVY), (f"  {extra}", T_SMALL, False, MUTED)])
    if years_from:
        for cc in range(FIRST, LAST + 1):
            c = ws.cell(row, cc)
            c.value = f"={get_column_letter(cc)}${years_from}"
            c.number_format = NUMFMT["year"]
            c.font = font(T_SMALL, True, NAVY)
            c.alignment = align("right", "center", 1)
    if up:
        _band_link(ws, row, up)


STAGE = {"sub": "sub", "result": "final", "final": "final", "deduct": "deduct", "kpi": "kpi", "plain": "plain"}


def data_row(ws, row, kind="data", fmt=NUM, label=None, unit=True):
    """kind: data · davon · memo · info · deduct · sub · final · kpi · plain (Summenstufen über C.sum_row, P2-06)."""
    if label is not None:
        _label(ws, row, label)
    b, cu = ws.cell(row, 2), ws.cell(row, 3)
    for c in C.iter_cells(ws, 2, row, LAST, row):
        c.fill = NOFILL
        c.border = Border(bottom=side("hair", LINE))
        if c.column >= FIRST:
            c.font = font(T_BODY, False, INK)
            c.alignment = align("right", "center", 1)
            c.number_format = fmt
    b.font = font(T_BODY, False, INK)
    b.alignment = align("left", "center", 1)
    if kind == "davon":
        b.font = font(T_SMALL, False, MUTED)
        b.alignment = align("left", "center", 2)
    elif kind == "memo":
        C.memo(ws, row, 2, LAST)
        b.alignment = align("left", "center", 1)
    elif kind == "info":          # P2-06: Nebeninfo 8,5 pt 98A2B3, ohne Rot-Regel
        for c in C.iter_cells(ws, 2, row, LAST, row):
            c.font = font(C.T_LABEL, False, FADED)
    elif kind == "memo_plain":    # nachrichtlich, aber lesbar (Break-even): 9 pt aufrecht 1A1D21
        for c in C.iter_cells(ws, 2, row, LAST, row):
            c.font = font(T_SMALL, False, INK)
    elif kind in STAGE:
        # Negativ-Rot setzt row_rules() zusammen mit der Verkaufsjahr-Markierung (LibreOffice wertet nur die erste Regel)
        C.sum_row(ws, row, "B", get_column_letter(LAST), STAGE[kind], neg=False)
    if unit:
        cu.font = font(T_SMALL, False, MUTED, italic=(kind == "memo"))
        cu.alignment = align("center", "center")
    else:
        if not C.is_formula(cu.value):
            cu.value = None
    ws.row_dimensions[row].height = C.H_ROW if STAGE.get(kind) == "final" else H_DATA


def gap(ws, row, h=H_GAP):
    _clear_row(ws, row, 2, LAST)
    ws.row_dimensions[row].height = h


def grid(ws, rows):
    """5-Jahres-Raster (P1-11): linke Kante thin D5D9DE an I, N, S, … – die Zahlen davor behalten ihren Einzug."""
    for r in rows:
        for cc in GRID:
            c = ws.cell(r, cc)
            b = c.border
            c.border = Border(right=b.right, top=b.top, bottom=b.bottom, left=side("thin", LINE2))


def footer(ws, row, merge_to="M", line_to=LAST, old=None):
    """Fuß auf dem Tabellenraster: Oberkante über die Tabellenbreite, Text über B:merge_to verbunden.
    old: bisherige Fußzeile der Vorlage (wird geleert, wenn der Fuß für die Buttonzeile nach unten wandert)."""
    rows = {row, row + 1} | ({old, old + 1} if old else set())
    for r in rows:
        for mr in list(ws.merged_cells.ranges):
            if mr.min_row == r:
                ws.unmerge_cells(str(mr))
        for c in C.iter_cells(ws, 2, r, line_to, r):
            if not C.is_formula(c.value):
                c.value = None
            c.border = Border()
    C.footer(ws, row, 2, line_to)
    C.safe_merge(ws, "B", row, merge_to, row)
    C.safe_merge(ws, "B", row + 1, merge_to, row + 1)
    ws.cell(row, 2).alignment = align("left", "center")
    ws.cell(row + 1, 2).alignment = align("left", "center")


def nav_row(ws, row, cols):
    """P3-08: Fußzeile Zurück · Übersicht · Weiter in Reiterfolge (C.btn_row: secondary · tertiary · primary)."""
    prev, prev_lab, nxt, nxt_lab = CHAIN[ws.title]
    (b1, b2), (m1, m2), (n1, n2) = cols
    _clear_row(ws, row, 2, LAST)
    C.btn_row(ws, row, [
        dict(c1=b1, c2=b2, text=f"‹  Zurück: {prev_lab}", target=prev, kind="secondary"),
        dict(c1=m1, c2=m2, text="Übersicht: Dashboard", target="Dashboard", kind="tertiary"),
        dict(c1=n1, c2=n2, text=f"Weiter: {nxt_lab}  ›", target=nxt, kind="primary")])


def year_marks(ws, head_ref, cond):
    """Verkaufsjahr (P1-11): Kopf C8D7EB mit Schrift 0B2A4A fett (≈ 10:1); Tabellenspalte E7EEF7 (row_rules)."""
    _rule(ws, head_ref, cond, Font(color=NAVY, bold=True), fill(C.MIST))
    return cond


def render(ws, rows, fmts, unit=True, up=None):
    for r, spec in rows.items():
        kind = spec[0]
        if kind == "gap":
            gap(ws, r)
        elif kind == "band":
            band(ws, r, spec[1], years_from=spec[2] if len(spec) > 2 else None,
                 extra=spec[3] if len(spec) > 3 else None, up=up)
        else:
            data_row(ws, r, kind, fmts.get(r, NUM), spec[1], unit=unit)


def base_colors(rows):
    """Grundfarbe je Datenzeile für die Auffangregel (entspricht der statischen Schrift aus data_row)."""
    out = {}
    for r, spec in rows.items():
        k = spec[0]
        if k in ("gap", "band"):
            continue
        out[r] = {"davon": MUTED, "memo": MUTED, "info": FADED, "sub": NAVY, "result": NAVY, "final": NAVY,
                  "kpi": NAVY}.get(k, INK)
    return out


# ------------------------------------------------------------------------------------------------ Projektion
def projektion(ws):
    _ungroup_cols(ws)
    _width(ws, "C", 12)
    header(ws, ("Berechnung", "Projektion"), "Projektion",
           "Miete, Kosten, Cashflow und Vermögen über 40 Jahre  ·  Beträge in € pro Jahr", back="Dashboard")
    year_head(ws, 8, 9, 10)
    # P1-12: Spalte C ist ausschließlich „Jahr 0 / Kauf“ (Eigenkapitaleinsatz im Exit-Cashflow C49)
    for r, text in ((8, "Jahr 0"), (10, "Kauf")):
        c = ws.cell(r, 3)
        C.set_text(c, text)
        c.font = font(T_SMALL, r == 8, WHITE)
        c.alignment = align("right", "center", 1)
    toolbar(ws, 11, "Eingaben dazu: S11 Prognose & Exit  ›", "S11 Prognose & Exit",
            year_legend(("Steuer: + Zahlung / − Erstattung", None, None),
                        ("Jahr 1 = erste 12 Monate ab Kauf", None, None)))
    rows = {
        12: ("band", "Mieteinnahmen"),
        13: ("data", None), 14: ("data", None), 15: ("sub", None),
        16: ("memo", "Umlagefähige Betriebskosten (Durchlaufposten)"), 17: ("memo", "Warmmiete"),
        18: ("gap",),
        19: ("band", "Bewirtschaftungskosten", None, "· nicht umlagefähig, Zahlungsabfluss"),
        20: ("data", None), 21: ("data", "Zuführung Erhaltungsrücklage WEG"), 22: ("data", None),
        23: ("data", None), 24: ("data", None), 25: ("data", None), 26: ("sub", None),
        27: ("gap",),
        28: ("band", "Cashflow"),
        29: ("sub", "= Einnahmenüberschuss vor Finanzierung (NOI)"),
        30: ("davon", "Zinsen"), 31: ("davon", "Tilgung (inkl. Sondertilgung)"),
        32: ("sub", "– Kapitaldienst (Zinsen + Tilgung inkl. Sondertilgung)"), 33: ("sub", None),
        34: ("info", "↳ Info: steuerliches Ergebnis (Basis der Steuer)"),
        35: ("data", "– Steuer (+ Zahlung / − Erstattung)"), 36: ("final", None),
        37: ("memo", None), 38: ("memo_plain", None),
        39: ("gap",),
        40: ("band", "Vermögensentwicklung"),
        41: ("data", None), 42: ("data", "– Restschuld Jahresende"), 43: ("final", "= Nettovermögen"),
        44: ("memo", "Vermögenszuwachs im Jahr"),
        45: ("kpi", "Eigenkapitalrendite in % (Vermögenszuwachs / Eigenkapital)"), 46: ("memo", None),
        47: ("gap",),
        48: ("band", "Exit-Cashflow", None, "· Basis für die Eigenkapitalrendite (IRR)"),
        49: ("final", "= Exit-Cashflow (Jahr 0 = Eigenkapitaleinsatz)"),
        50: ("gap",),
    }
    render(ws, rows, {45: PCT1}, unit=False, up=EDGE)
    # Z. 49 (P1-12): C49 wie die Jahreszellen (rechtsbündig, Einzug 1, fett Navy)
    ws["C49"].number_format = NUM
    ws["C49"].font = font(T_BODY, True, NAVY)
    ws["C49"].alignment = align("right", "center", 1)
    footer(ws, 53, old=51)
    nav_row(ws, 51, (("B", "B"), ("D", "G"), ("I", EDGE)))
    gap(ws, 52, H_GAP)
    grid(ws, [8, 10] + [r for r, s in rows.items() if s[0] != "gap"])
    ws.freeze_panes = "D11"
    # Bedingte Formatierung (nur Anzeige) – je Zeile: Nullwert grau, Negativ-Rot, Verkaufsjahr-Spalte
    ycond = year_marks(ws, "D8:AQ10", "D$9=Haltedauer")
    spec = {r: {"base": b} for r, b in base_colors(rows).items()}
    for r in (15, 29, 33, 36, 37, 38, 43, 44, 49):
        spec[r]["rules"] = [neg_rule(r, True if r in (36, 43, 49) else None)]
    for r in (36, 43, 49):
        spec[r]["fill"] = C.MIST
    row_rules(ws, spec, ycond)
    C.neg_red(ws, "C49")
    _rule(ws, "D11:AQ49", ycond, None, fill(TINT))


# ------------------------------------------------------------------------------------------------ Finanzierung
KD_LABEL = "Kapitaldienst (Zinsen + Tilgung inkl. Sondertilgung)"   # mappenweit gleiche Beschriftung (P12)


def finanzierung(ws):
    _ungroup_cols(ws)
    _width(ws, "C", 12)
    header(ws, ("Berechnung", "Tilgungsplan"), "Tilgungsplan",
           "Tilgungsverlauf Darlehen I und II  ·  Annuitätendarlehen mit konstanter Rate", back="Dashboard")
    year_head(ws, 8, 9, 10)
    toolbar(ws, 11, "Eingaben dazu: S07 Finanzierung  ›", "S07 Finanzierung",
            legend(("Verkaufsjahr", "■", C.MIST), ("Anschlusszins nach der Zinsbindung", "●", BLUE),
                   ("– entfällt (getilgt bzw. nicht genutzt)", None, None), size=C.T_LABEL))
    rows = {
        12: ("band", "Darlehen I"),
        13: ("data", None), 14: ("data", None), 15: ("data", None), 16: ("data", None),
        17: ("data", "– Reguläre Tilgung"), 18: ("data", "– Sondertilgung"),
        19: ("final", "= Restschuld Jahresende"), 20: ("plain", KD_LABEL),
        21: ("gap",),
        22: ("band", '="Darlehen II"&IF(Darlehen_II=0,"  ·  nicht genutzt","")'),
        23: ("data", None), 24: ("data", None), 25: ("data", None), 26: ("data", None),
        27: ("data", "– Reguläre Tilgung"), 28: ("data", "– Sondertilgung"),
        29: ("data", "– Tilgungszuschuss (KfW, kein Zahlungsabfluss)"),
        30: ("final", "= Restschuld Jahresende"), 31: ("plain", KD_LABEL),
        32: ("gap",),
        33: ("band", "Summe Darlehen"),
        34: ("data", None), 35: ("data", None), 36: ("data", None),
        37: ("sub", "= Kapitaldienst gesamt (Zinsen + Tilgung inkl. Sondertilgung)"),
        38: ("final", "= Restschuld Jahresende"),
        39: ("memo", None), 40: ("memo", None),
        41: ("gap",),
    }
    render(ws, rows, {14: PCT2, 24: PCT2, 39: PCT2}, up=EDGE)
    # P43: Darlehen II ungenutzt → Zeilen 23–31 als Gliederungsgruppe, beim Build eingeklappt; Hinweis im Band
    d2 = _input_value(ws.parent, "Darlehen_II")
    for r in range(23, 32):
        ws.row_dimensions[r].outlineLevel = 1
        ws.row_dimensions[r].hidden = d2 == 0
    ws.row_dimensions[32].collapsed = d2 == 0
    hint = ws.cell(22, _col("K"))
    hint.value = ('=IF(Darlehen_II=0,"▸ ausgeblendet  ·  über das „+“ der Gliederung links einblenden",'
                  '"▸ über die Gliederung links ein-/ausklappbar")')
    hint.font = font(C.T_LABEL, False, MUTED)
    hint.alignment = align("right", "center", 1)
    ws.sheet_format.outlineLevelRow = max(ws.sheet_format.outlineLevelRow or 0, 1)
    footer(ws, 44, old=42)
    nav_row(ws, 42, (("B", "B"), ("D", "G"), ("I", EDGE)))
    gap(ws, 43, H_GAP)
    grid(ws, [8, 10] + [r for r, s in rows.items() if s[0] != "gap"])
    ws.freeze_panes = "D11"
    ycond = year_marks(ws, "D8:AQ10", "D$9=Haltedauer")
    faded = Font(color=FADED, bold=False, italic=True)
    switch = Font(bold=True, color=BLUE)          # Zinswechsel nach der Zinsbindung
    # P2-02/P43: entfallene Werte „–“ in 98A2B3 kursiv, nicht fett – Darlehen I nach Volltilgung (Z. 13–20),
    # Summenblock nach Volltilgung (Z. 34–40), Darlehen II ungenutzt (Z. 23–31); sonst Nullwerte grau (P1-11)
    base = base_colors(rows)
    spec = {}
    for r in range(13, 21):
        spec[r] = {"base": base[r], "pre": [("D$13=0", faded, DASH_FMT)],
                   "rules": [("AND(ISNUMBER(C14),D14<>C14)", switch)] if r == 14 else [],
                   "fill": C.MIST if r == 19 else TINT}
    for r in range(23, 32):
        pre = [("Darlehen_II=0", faded, DASH_FMT, fill(WHITE) if r == 30 else None)]
        spec[r] = {"base": base[r], "pre": pre,
                   "rules": [("AND(ISNUMBER(C24),D24<>C24)", switch)] if r == 24 else [],
                   "fill": C.MIST if r == 30 else TINT}
    for r in range(34, 41):
        spec[r] = {"base": base[r], "pre": [("D$34=0", faded, DASH_FMT)], "fill": C.MIST if r == 38 else TINT}
    row_rules(ws, spec, ycond)
    # Beschriftung B:C der inaktiven Zeilen ebenso zurücknehmen (Summenzeilen dann weiß, nicht fett)
    _rule(ws, "B23:C29", "Darlehen_II=0", Font(color=FADED, bold=False, italic=True))
    _rule(ws, "B30:C31", "Darlehen_II=0", Font(color=FADED, bold=False, italic=True), fill(WHITE))
    _rule(ws, "D11:AQ40", ycond, None, fill(TINT))


def _input_value(wb, name):
    """Wert einer Eingabe hinter einem Namen beim Build (nur Konstanten; Formeln → None)."""
    try:
        dn = wb.defined_names.get(name)
        sheet, ref = next(iter(dn.destinations))
        v = wb[sheet][ref.replace("$", "")].value
        return None if C.is_formula(v) else v
    except Exception:
        return None


# ------------------------------------------------------------------------------------------------ Steuern
STEUERN_LABELS = {
    17: "Wohngebäude", 21: "Sofortige Verlustverrechnung", 23: "Disagio sofort abziehbar",
    27: "Typisierte Nutzungsdauer",
    28: "Degressive AfA aktiv", 30: "AfA-Bemessungsgrundlage Gebäude (ohne Sanierungsanteil)",
    32: "Sonder-AfA § 7b aktiv", 35: "Renovierung als Herstellungskosten aktivieren",
    78: "Haltedauer",
    85: "Veräußerungsgewinn (Preis – Verkaufskosten – Buchwert)",
    86: "Veräußerungsgewinn steuerpflichtig",
    91: "= Nettoerlös nach Steuern und Darlehensablösung",
    94: "= Gesamtertrag nach Steuern",
    96: "IRR n. St.  ·  Eigenkapitalrendite p. a.",
    98: "Summe Gebäude-AfA – Standard nach Baujahr",
    99: "Steuereffekt der gewählten Variante",
}
AFA_BLOCK_TITLE = "AfA-Vergleich  ·  Jahre 1–10"
AFA_BLOCK_97 = "Summe Gebäude-AfA – gewählte Variante"
STEUERN_NOTES = {
    24: "Fallback auf „Automatisch“, wenn Voraussetzungen fehlen (siehe Prüfhinweise im Cockpit)",
    88: "nur bei Option „Verlustvortrag“ bzw. GmbH; Mindestbesteuerung (§ 10d Abs. 2 EStG) nicht modelliert",
    94: "Cashflows + Nettoerlös – Eigenkapital",
    99: "gegenüber Standard-AfA · positiv = Steuerersparnis (Liquiditätsvorteil, keine endgültige Ersparnis)",
}
STEUERN_FMT = {
    17: YESNO, 18: PCT2, 19: PCT2, 20: PCT2, 21: YESNO, 22: NUMFMT["years_n"], 23: YESNO,
    25: PCT2, 26: PCT2, 27: NUMFMT["years_n"], 28: YESNO, 31: C.typo_minus('#,##0" €/m²";-#,##0" €/m²";"–"'),
    32: YESNO, 35: YESNO, 78: NUMFMT["years_n"], 86: YESNO, 87: PCT2, 95: C.MULT2, 96: PCT1,
}
KV_LAST = 13   # M
PARAM_GROUPS = (25, 32, 35)   # P3-06: AfA · Sonder-AfA § 7b · Maßnahmen


def kv_rows(ws, rows, last=KV_LAST):
    """Parameter-/Exit-Block: Beschriftung B, Wert C rechts, Hinweis E:M grau; Haarlinie, keine Seitenrahmen."""
    for r in rows:
        if r in STEUERN_LABELS:
            _label(ws, r, STEUERN_LABELS[r])
        if r in STEUERN_NOTES:
            C.set_text(ws.cell(r, 5), STEUERN_NOTES[r])
        for c in C.iter_cells(ws, 2, r, last, r):
            c.fill = NOFILL
            c.border = Border(bottom=side("hair", LINE))
        b, v, n = ws.cell(r, 2), ws.cell(r, 3), ws.cell(r, 5)
        b.font = font(T_BODY, False, INK)
        b.alignment = align("left", "center", 1)
        v.font = font(T_BODY, False, INK)
        v.alignment = align("right", "center", 1)
        v.number_format = STEUERN_FMT.get(r, EUR)
        n.font = font(T_SMALL, False, MUTED)
        n.alignment = align("left", "center", 1)
        ws.row_dimensions[r].height = H_DATA


def kv_sum(ws, row, stage, top=True):
    """Summen-/Kennzahlstufe im Parameter-/Exit-Block (B:M); Hinweis in E bleibt 9 pt grau."""
    C.sum_row(ws, row, "B", get_column_letter(KV_LAST), stage, top=top)
    n = ws.cell(row, 5)
    n.font = font(T_SMALL, False, MUTED)
    ws.row_dimensions[row].height = C.H_ROW


def _outline(ws, r1, r2, level=1):
    for r in range(r1, r2 + 1):
        ws.row_dimensions[r].outlineLevel = max(ws.row_dimensions[r].outlineLevel or 0, level)


def steuern(ws):
    _ungroup_cols(ws)
    _width(ws, "C", 12)
    header(ws, ("Berechnung", "Steuer-Tabelle"), "Steuer-Tabelle",
           "Abgeleitete Parameter, AfA-Verlauf und steuerliches Ergebnis über 40 Jahre")
    # ---- Parameterblock (P3-04: per Gliederung einklappbar; technische Indizes Ebene 2, bleiben verborgen)
    band(ws, 8, "Abgeleitete steuerliche Parameter", c2=KV_LAST, extra="· über die Gliederung links einklappbar")
    # Sprunglinks im Band rechtsbündig (P3-08)
    for c1, c2, text, row in (("I", "K", "↓ Jahrestabelle", 40), ("L", "M", "↓ Verkaufsszenario", 77)):
        C.safe_merge(ws, c1, 8, c2, 8)
        c = ws[f"{c1}8"]
        C.text_link(c, text, ws.title, C.link_row(ws.title, row), size=C.T_LABEL, bold=False,
                    tooltip=f"Zum Abschnitt „{text[2:]}“")
        c.font = font(C.T_LABEL, False, BLUE)
        c.alignment = align("right", "center", 1)
    C.hide_rows(ws, 9, 16)
    C.hide_rows(ws, 24, 24)
    kv_rows(ws, [r for r in range(17, 38) if r != 24])
    # P3-06: Gruppentrennung (AfA · Sonder-AfA § 7b · Maßnahmen) als Oberlinie 4A86C8 – ohne Zeileneinfügung
    for r in PARAM_GROUPS:
        prev = r - 2 if r == 25 else r - 1
        for cc in range(2, KV_LAST + 1):
            c = ws.cell(r, cc)
            c.border = Border(top=side("thin", ACCENT), bottom=c.border.bottom)
            p = ws.cell(prev, cc)
            p.border = Border(top=p.border.top, bottom=side("thin", ACCENT))
    gap(ws, 38, H_GAP)
    C.hide_rows(ws, 39, 39)
    _outline(ws, 9, 37, 1)
    _outline(ws, 9, 16, 2)
    _outline(ws, 24, 24, 2)
    # ---- Jahrestabelle
    year_head(ws, 40, 41, 42)
    toolbar(ws, 43, "Eingaben dazu: S09 Steuern  ›", "S09 Steuern",
            year_legend(("Steuer: + Zahlung / − Erstattung (Verrechnung mit anderen Einkünften)", None, None)))
    rows = {
        44: ("band", "Abschreibung (AfA)", 42),
        45: ("data", None), 46: ("data", None), 47: ("data", None), 48: ("data", None), 49: ("data", None),
        50: ("data", None), 51: ("sub", None), 52: ("memo", None), 53: ("memo", None), 54: ("data", None),
        55: ("data", None), 56: ("sub", None), 57: ("memo", None),
        58: ("gap",),
        59: ("band", "Steuerliches Ergebnis aus dem Objekt", 42),
        60: ("data", None),
        61: ("davon", "Bewirtschaftungskosten (steuerlich)"), 62: ("davon", "Schuldzinsen"),
        63: ("davon", "Finanzierungsnebenkosten und Disagio"),
        64: ("davon", "Erhaltungsaufwand (ggf. verteilt, § 82b EStDV)"),
        65: ("davon", "Abschreibungen"),
        66: ("sub", "– Werbungskosten / Betriebsausgaben gesamt"), 67: ("sub", "= Steuerliches Ergebnis"),
        68: ("memo", None), 69: ("data", None), 70: ("memo", None),
        71: ("data", "Körperschaftsteuersatz im Kalenderjahr (nur GmbH)"),
        72: ("data", "Angewendeter Steuersatz (Grenzsatz bzw. KSt + Soli + GewSt)"),
        73: ("final", "= Steuer (+ Zahlung / − Erstattung)"), 74: ("memo", None),
    }
    render(ws, rows, {71: PCT2, 72: PCT2})
    grid(ws, [40, 42] + [r for r, s in rows.items() if s[0] != "gap"])
    gap(ws, 75, H_GAP)
    C.hide_rows(ws, 76, 76)
    # ---- Exit-Block
    band(ws, 77, "Verkaufsszenario (Exit)", c2=KV_LAST, extra="· Steuer und Gesamtrendite")
    _band_link(ws, 77)
    kv_rows(ws, range(78, 100))
    # P2-06: Zwischensumme · Endsumme · Kennzahlenpaar im Band (Doppellinie nur unter dem IRR)
    kv_sum(ws, 91, "sub")
    kv_sum(ws, 94, "final")
    kv_sum(ws, 95, "kpi_band")
    kv_sum(ws, 96, "kpi_band")
    for cc in range(2, KV_LAST + 1):
        c = ws.cell(96, cc)
        c.border = Border(bottom=side("double", NAVY))
    # P3-06: AfA-Vergleich (Z. 97–99) als Unterabschnitt mit Minititel statt dreifachem Präfix
    for r in (97, 98, 99):
        C.sum_row(ws, r, "B", get_column_letter(KV_LAST), "plain")
        ws.cell(r, 3).alignment = align("right", "center", 1)
        ws.cell(r, 5).font = font(T_SMALL, False, MUTED)
        ws.cell(r, 5).alignment = align("left", "center", 1)
        ws.row_dimensions[r].height = H_DATA
    b97 = ws["B97"]
    b97.value = C.rich([(AFA_BLOCK_TITLE + "\n", T_SMALL, True, NAVY), (AFA_BLOCK_97, T_BODY, False, INK)])
    b97.alignment = align("left", "bottom", 1, wrap=True)
    ws.row_dimensions[97].height = C.H_ROW3        # Luft über dem Minititel (Abstand zur Doppellinie)
    for cc in range(3, KV_LAST + 1):
        c = ws.cell(97, cc)
        c.alignment = align(c.alignment.horizontal or "left", "bottom", c.alignment.indent)
    lk = ws["E97"]
    C.text_link(lk, "Zum AfA-Vergleich  ›", "AfA-Vergleich", size=T_SMALL, bold=True,
                tooltip="Alle Abschreibungsvarianten im Vergleich")
    lk.alignment = align("left", "bottom", 1)
    _outline(ws, 78, 99, 1)
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=True, applyStyles=False,
                                            showOutlineSymbols=True)
    ws.sheet_format.outlineLevelRow = 2
    gap(ws, 100, H_GAP)
    footer(ws, 103, line_to=KV_LAST, old=101)
    nav_row(ws, 101, (("B", "B"), ("D", "G"), ("I", EDGE)))
    gap(ws, 102, H_GAP)
    # ---- Bedingte Formatierung
    ycond = "D$41=Haltedauer"
    year_marks(ws, "D40:AQ42", ycond)
    faded = Font(color=FADED, bold=False, italic=True)
    spec = {r: {"base": b} for r, b in base_colors(rows).items()}
    spec[71]["pre"] = [("Rechtsform_Idx=1", faded, DASH_FMT)]
    spec[67]["rules"] = [neg_rule(67)]
    spec[69]["rules"] = [neg_rule(69)]
    # Z. 73: Steuer-Sicht laut Beschriftung – negative Werte (Erstattung) bewusst NICHT rot (P07)
    spec[73]["fill"] = C.MIST
    row_rules(ws, spec, ycond)
    _rule(ws, "B71:C71", "Rechtsform_Idx=1", faded)
    _rule(ws, "D43:AQ74", ycond, None, fill(TINT))
    # Exit-Block: negative Ergebnisse rot (P2-01, C92 kumulierter Cashflow)
    C.neg_red(ws, "C92")
    # Nullwerte im Parameter-/Exit-Block grau (P1-11), außer Ja/Nein-Feldern
    for r in list(range(17, 38)) + list(range(78, 100)):
        if STEUERN_FMT.get(r) != YESNO and r not in (91, 94, 95, 96):
            _rule(ws, f"C{r}", f"AND(ISNUMBER(C{r}),C{r}=0)", ZERO_FONT)


# ------------------------------------------------------------------------------------------------ AfA-Vergleich
AFA_VARIANTS = {
    19: ("Linear 2,0 %", "§ 7 Abs. 4 S. 1 Nr. 2b EStG · Fertigstellung 1925–2022"),
    20: ("Linear 2,5 %", "§ 7 Abs. 4 S. 1 Nr. 2c EStG · Fertigstellung vor 1925"),
    21: ("Linear 3,0 %", "§ 7 Abs. 4 S. 1 Nr. 2a EStG · Fertigstellung ab 2023"),
    22: ("Gutachten (RND)", "§ 7 Abs. 4 S. 2 EStG · tatsächliche Restnutzungsdauer"),
    23: ("Degressiv 5 %", "§ 7 Abs. 5a EStG · Baubeginn 10/2023–9/2029"),
    24: ("Linear 3 % + § 7b", "5 % p. a. in Jahren 1–4, danach Restwert-AfA"),
    25: ("Degressiv + § 7b", "kombiniert; die Sonder-AfA mindert den Restwert"),
    26: ("Im Modell angewendet", "lt. Steuer-Tabelle · inkl. Sonder-AfA, ohne § 7h/7i"),
}
AFA_BASICS = {
    9: "AfA-Basis Gebäude (ohne § 7h/7i-Anteil)",
    10: "Sonder-AfA § 7b – Basis (max. 4.000 €/m²)",
    11: "Sonder-AfA § 7b p. a. (5 %)",
    12: "Steuersatz Jahr 1 (angewendet)",
    13: "Kalkulationszins für den Barwert",
    14: "Restnutzungsdauer lt. Gutachten",
    15: "Baujahr · Objektart",
}
AFA_HINTS = {
    47: ("Liquiditätseffekt: ", "Eine höhere AfA in den ersten Jahren ist ein Liquiditäts- und Zinsvorteil "
         "(Steuerstundung), keine endgültige Ersparnis – bei Verkauf innerhalb der 10-Jahres-Frist (Privat) bzw. "
         "immer (GmbH) mindert die kumulierte AfA den Buchwert und erhöht den steuerpflichtigen Veräußerungsgewinn."),
    48: ("Prüfumfang: ", "Die Spalte „Anwendbar“ prüft nur die im Tool erfassten Kriterien (Baujahr, Objektart, "
         "Baukosten je m²). Nicht geprüft werden: Baubeginn im Zeitfenster (§ 7 Abs. 5a), Bauantrag/Bauanzeige, "
         "Effizienzhaus-40-Standard mit QNG-Siegel (§ 7b), Anschaffung bis Ende des Fertigstellungsjahres (§ 7b)."),
    49: ("Kombination: ", "Degressive AfA und § 7b sind kombinierbar (§ 7b Abs. 1 S. 1 i. V. m. § 7 Abs. 5a EStG); "
         "im Tool wird die Sonder-AfA vereinfachend vom Restwert abgezogen (konservativ). Ein Wechsel von degressiv "
         "zu linear (Restwert / Restnutzungsdauer) erfolgt automatisch, sobald vorteilhaft – innerhalb der ersten "
         "10 Jahre bei 3 % typisierter Nutzungsdauer nicht."),
    50: ("§ 7h/7i: ", "Erhöhte Absetzungen nach § 7h / § 7i EStG (9 % / 7 %) betreffen nur den bescheinigten "
         "Sanierungsanteil und sind hier nicht als Variante dargestellt; sie werden in der Steuer-Tabelle zusätzlich "
         "zur regulären AfA der Altsubstanz berechnet."),
}
AFA_LAST = 17  # Q
AFA_SIDE_LAST = 13  # M: Nebenrechnungen nur Jahre 1–10 (P3-06)
EMU_PT = 12700
# Kurzform der angewendeten Methode für die Kachel (reine Anzeigeformel, P1-11)
AFA_SHORT = ('=IF(Is_Degressiv=1,"Degressiv "&FIXED(Degressiv_Satz*100,1)&" %",'
             'IF(Methode_eff=6,"Gutachten ","Linear ")&FIXED(AfA_Satz_linear*100,1)&" %")'
             '&IF(SonderAfA_aktiv=1," + § 7b","")')
AFA_NOTE = ("Die Methode folgt automatisch aus Baujahr und Objektart (Eingaben in S10). Die Tabelle zeigt alle "
            "Varianten zum Vergleich – zulässig sind nur die als „anwendbar“ markierten.")


def afa(ws):
    _ungroup_cols(ws)
    # Spaltenraster (P10/P28): B:C Beschriftung (C breit genug für „Automatisch nach Baujahr (§ 7 Abs. 4 EStG)“),
    # D:O einheitlich 74 px (Jahre, Σ-Spalten, Kachelrinnen I und N gleich breit), P:Q je 111 px → alle drei Kacheln 296 px
    _width(ws, "B", 37)
    _width(ws, "C", 34)
    for cc in range(4, 16):
        _width(ws, get_column_letter(cc), 10.57)
    _width(ws, "P", 15.86)
    _width(ws, "Q", 15.86)
    header(ws, ("Steuer-Tabelle", "AfA-Vergleich"), "AfA-Vergleich",
           "Abschreibungsvarianten für dieses Objekt in den Jahren 1–10  ·  Beträge in €", edge="Q")
    # ---- Grundlagen + Kacheln
    band(ws, 8, "Grundlagen", c2=AFA_LAST)
    _band_link(ws, 8, "Q", "Eingaben dazu: S10 Abschreibung  ›", "S10 Abschreibung",
               tooltip="Eingaben zu diesem Blatt: Schritt 10 · Abschreibung")
    fmts = {9: EUR, 10: EUR, 11: EUR, 12: PCT2, 13: PCT2, 14: '0" Jahre";;"– (kein Gutachten)"', 15: "@"}
    for r in range(9, 16):
        _label(ws, r, AFA_BASICS[r])
        _clear_row(ws, r, 2, AFA_LAST)
        for c in C.iter_cells(ws, 2, r, 3, r):
            c.border = Border(bottom=side("hair", LINE))
        b, v = ws.cell(r, 2), ws.cell(r, 3)
        b.font = font(T_BODY, False, INK)
        b.alignment = align("left", "center", 1)
        v.font = font(T_BODY, False, INK)
        v.alignment = align("right", "center", 1)
        v.number_format = fmts[r]
        ws.row_dimensions[r].height = C.H_ROW
    ws["C15"].value = '=Baujahr&" · "&Objektart'
    tiles(ws)
    gap(ws, 16, H_GAP)
    # ---- Variantentabelle
    band(ws, 17, "AfA-Beträge je Jahr", c2=AFA_LAST,
         meta="Steuerersparnis = AfA × Steuersatz Jahr 1 · Barwert mit dem Kalkulationszins")
    # Versalien nur für die Textköpfe – „Jahr n“ bleibt in Satzschreibung (Kategorien des Liniendiagramms)
    heads = {"B": ("VARIANTE", "left"), "C": ("ANWENDBAR FÜR DIESES OBJEKT?", "left"),
             "N": ("Σ JAHRE\n1–4", "right"), "O": ("Σ JAHRE\n1–10", "right"),
             "P": ("STEUERERSPARNIS\n1–10 NOMINAL", "right"), "Q": ("BARWERT\nSTEUERERSPARNIS", "right")}
    for k in range(10):
        heads[get_column_letter(4 + k)] = (f"Jahr {k + 1}", "right")
    C.section(ws, 18, "B", get_column_letter(AFA_LAST), level=2, labels=heads, height=C.H_ROW2, caps=False)
    for cc in range(2, AFA_LAST + 1):
        c = ws.cell(18, cc)
        c.alignment = align(c.alignment.horizontal, "center", 1, wrap=cc >= 14)
    for r in range(19, 27):
        name, basis = AFA_VARIANTS[r]
        key = r == 26
        b = ws.cell(r, 2)
        b.value = C.rich([(name, T_BODY, True, NAVY if key else INK), ("\n" + basis, T_MICRO, False, MUTED)])
        for c in C.iter_cells(ws, 2, r, AFA_LAST, r):
            c.fill = NOFILL
            c.border = Border(bottom=side("hair", LINE))
            if c.column >= 4:
                c.font = font(T_BODY, c.column in (15, 17), INK)
                c.alignment = align("right", "center", 1)
                c.number_format = NUM
        b.font = font(T_BODY, True, INK)
        b.alignment = align("left", "center", 1, wrap=True)
        cst = ws.cell(r, 3)
        cst.font = font(T_SMALL, True, MUTED)
        cst.alignment = align("left", "center", 1, wrap=True)
        cst.number_format = NUMFMT["status_dot"]
        ws.row_dimensions[r].height = C.H_ROW2
    C.sum_row(ws, 26, "B", get_column_letter(AFA_LAST), "final")
    ws.row_dimensions[26].height = C.H_ROW2
    ws["C26"].font = font(T_SMALL, True, NAVY)
    ws["C26"].alignment = align("left", "center", 1, wrap=True)
    ws["C26"].number_format = "General"
    # Legende der Spalte „Anwendbar“ als Fußnote direkt unter der Tabelle (P1-18)
    _clear_row(ws, 27, 2, AFA_LAST)
    ws["B27"].value = legend(("anwendbar", "●", GREEN), ("nur mit Gutachten", "●", AMBER),
                             ("nicht anwendbar (Beträge nur zum Vergleich)", "●", FADED), size=T_MICRO)
    ws["B27"].font = font(T_MICRO, False, MUTED)
    ws["B27"].alignment = align("left", "top", 1)
    ws.row_dimensions[27].height = 22
    # ---- Nebenrechnungen (P3-06: Kopfbänder nur bis Jahr 10 = Spalte M)
    side_last = get_column_letter(AFA_SIDE_LAST)
    for head_row, title, first in ((28, "NEBENRECHNUNG: RESTWERT GEBÄUDE AM JAHRESENDE (€)", 29),
                                   (37, "NEBENRECHNUNG: KUMULIERTE AFA JE VARIANTE (€)", 38)):
        _clear_row(ws, head_row, 2, AFA_LAST)
        heads = {"B": (title, "left")}
        for k in range(10):
            heads[get_column_letter(4 + k)] = (f"Jahr {k + 1}", "right")
        C.section(ws, head_row, "B", side_last, level=2, labels=heads, caps=False)
        for i in range(7):
            r = first + i
            if head_row == 28:
                _label(ws, r, AFA_VARIANTS[19 + i][0])
            else:
                ws.cell(r, 2).value = f"=$S{19 + i}"
            _clear_row(ws, r, 2, AFA_LAST)
            for c in C.iter_cells(ws, 2, r, AFA_SIDE_LAST, r):
                c.border = Border(bottom=side("hair", LINE))
                c.font = font(T_SMALL, False, MUTED)
                if c.column >= 4:
                    c.alignment = align("right", "center", 1)
                    c.number_format = NUM
            ws.cell(r, 2).alignment = align("left", "center", 1)
            ws.row_dimensions[r].height = H_DATA
    gap(ws, 36, H_GAP)
    gap(ws, 45, H_GAP)
    # ---- Hinweise: zwei Spalten (B:C | E:Q), je Absatz ein fettes Stichwort (P3-10)
    band(ws, 46, "Hinweise zur Interpretation", c2=AFA_LAST)
    for r in AFA_HINTS:
        for mr in list(ws.merged_cells.ranges):
            if mr.min_row == r:
                ws.unmerge_cells(str(mr))
        _clear_row(ws, r, 2, AFA_LAST)
        for c in C.iter_cells(ws, 2, r, AFA_LAST, r):
            c.value = None
    # P33: zwei gleich breite Spaltengruppen mit einer Rinnenspalte; Zeilenhöhe aus dem längeren Block (P27-Raster)
    split = min(range(5, AFA_LAST - 3), key=lambda k: abs(C.span_px(ws, 2, k) - C.span_px(ws, k + 2, AFA_LAST)))
    layout = ((47, 2, split, 47), (47, split + 2, AFA_LAST, 49), (48, 2, split, 48), (48, split + 2, AFA_LAST, 50))
    need = {}
    for row, c1, c2, key in layout:
        kw, text = AFA_HINTS[key]
        C.safe_merge(ws, c1, row, c2, row)
        c = ws.cell(row, c1)
        c.value = C.rich([(kw, T_SMALL, True, NAVY), (text, T_SMALL, False, INK2)])
        c.font = font(T_SMALL, False, INK2)
        c.alignment = align("left", "top", 1, wrap=True)
        n = C.lines_needed_metric(kw + text, C.span_px(ws, c1, c2), T_SMALL, False, 1)
        need[row] = max(need.get(row, 1), n)
    for row, n in need.items():
        ws.row_dimensions[row].height = C.text_row_height(n)
    gap(ws, 49, H_GAP)
    C.hide_rows(ws, 50, 51)
    # ---- Diagramme
    band(ws, 52, "Diagramme", c2=AFA_LAST, extra="· Jahre 1–10")
    for r in range(53, 70):
        ws.row_dimensions[r].height = 15
    place_charts(ws)
    gap(ws, 70, C.H_ROW)
    # P19/P3-04: leere Variante als Fußnote unter den Diagrammen (reine Anzeigeformel) – 8 pt aufrecht 5B6068
    fn = ws["B70"]
    fn.value = ('=IF(N(C14)=0,"Gutachten (RND): kein Restnutzungsdauer-Gutachten erfasst – Variante ohne Balken '
                'und Linie","")')
    fn.font = font(T_MICRO, False, MUTED)
    # vertikal zentriert: eigene Zellvorlage – mit B27 (Rich-Text, „●“ grün fett) teilt LibreOffice sonst die Schrift
    fn.alignment = align("left", "center", 1)
    for r in (71, 72):
        ws.row_dimensions[r].hidden = False
    footer(ws, 74, merge_to="Q", line_to=AFA_LAST, old=73)
    gap(ws, 71, H_GAP)
    nav_row(ws, 72, (("B", "C"), ("F", "I"), ("L", "Q")))
    gap(ws, 73, H_GAP)
    C.hide_cols(ws, "R", "S")
    # ---- Bedingte Formatierung: Anwendbarkeit als Statuswort (● in Statusfarbe, keine Flächen); Nullwerte grau
    cf_rules(ws, "C19:C25", [('LEFT($C19,2)="Ja"', Font(color=GREEN, bold=True)),
                             ('LEFT($C19,4)="Nur "', Font(color=AMBER, bold=True)),
                             ('LEFT($C19,4)="Nein"', Font(color=FADED, bold=False))], Font(color=MUTED))
    cf_rules(ws, "D19:Q25", [("AND(ISNUMBER(D19),D19=0)", ZERO_FONT),
                             ('LEFT($C19,4)="Nein"', Font(color=FADED, bold=False))], Font(color=INK))


def tiles(ws):
    """Kacheln rechts neben den Grundlagen (C.tile, calm): drei gleich breite Kacheln (je 296 px, Rinnen I und N),
    Anatomie Kopf 18 · Wert · Fuß 20 pt (P1-02). Die Grundlagenliste läuft im 18-pt-Raster daneben weiter; deshalb
    liegt der Wert über zwei Listenzeilen (Z. 10–11) und die Fußzeile in Z. 12. Darunter – bündig mit dem Listenende –
    eine Einordnung (C.callout_box) statt eines Lochs in der rechten Spalte."""
    specs = (("E", "H", "Im Modell angewendet", AFA_SHORT, "General", "=C26"),
             ("J", "M", "AfA Jahre 1–10 (angewendet)", "=O26", EUR, "Summe der Gebäude-AfA lt. Steuer-Tabelle"),
             ("O", "Q", "Barwert der Steuerersparnis", "=Q26", EUR, "Jahre 1–10, abgezinst mit dem Kalkulationszins"))
    for r in range(9, 16):
        ws.row_dimensions[r].height = C.H_ROW
    for c1, c2, lab, val, fmt, sub in specs:
        C.tile(ws, c1, c2, 9, 10, 12, label=lab, value=val, sub=sub, fmt=fmt, value_rows=2,
               gap_right=False, gap=False, status=None, neg=False)
    for r in (10, 11):
        ws.row_dimensions[r].height = C.H_ROW
    # Einordnung E14:Q15 (Kopf 18 pt im Listenraster, Körper einzeilig)
    C.callout_box(ws, "E", 14, "Q", 15, 15, title="Einordnung", text=AFA_NOTE, pill=False, fit=None)
    for r in (13, 14, 15):
        ws.row_dimensions[r].height = C.H_ROW


def place_charts(ws):
    """chart33 (Balken) exakt B53 bis Ende C (Fixierlinie frei), chart34 (Linien) E53 bis Ende Q;
    Höhe 3.059.280 EMU wie auf „Diagramme“ (16 Zeilen à 15 pt + Rest)."""
    rest = 3059280 - 16 * 15 * EMU_PT
    for ch in ws._charts:
        kind = type(ch).__name__
        if kind.startswith("Bar"):
            c1, c2 = 1, 3
        else:
            c1, c2 = 4, 17
        ch.anchor = TwoCellAnchor(editAs="twoCell", _from=AnchorMarker(col=c1, colOff=0, row=52, rowOff=0),
                                  to=AnchorMarker(col=c2, colOff=0, row=68, rowOff=rest))


# ------------------------------------------------------------------------------------------------ Einstieg
def apply(wb):
    for name, fn in (("Projektion", projektion), ("Finanzierung", finanzierung), ("Steuern", steuern),
                     ("AfA-Vergleich", afa)):
        if name in wb.sheetnames:
            fn(wb[name])
            _calibri(wb[name])
            C.cf_close(wb[name])
