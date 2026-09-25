"""Blatt „Start“: Hero, Entscheidungskarte „① Kauf als“, vier Kennzahl-Kacheln, Ablauf ①–⑥, Legende mit
Tipp-Kasten, Blätterverzeichnis, Fuß – ausschließlich mit den zentralen Bausteinen aus core.py.

Raster (px, Runde 4 / P1-08): C 224 | D 112 | E 336 | F 336 | G 336  →  C:D = E = F = G (vier gleich breite Spalten);
  Hero B:H von 31 bis 1 417 px = linke/rechte Inhaltskante der Schrittseiten und der festen Reiterleiste (NAV_END).
  Hero:  Logo-Spalte C:D (Logo zentriert), Inhalt E:G; Kennzahlen und Buttons exakt in E | F | G.
  Kacheln Z. 27–29: C:D | E | F | G – C.tile: Kopfstreifen · Wert in Statusfarbe · Fußzeile Kontext links / Status rechts.
Namensziele bleiben an ihrem Platz: Rechtsform = D23, ST_BMR = E28, ST_CF = C31, ST_DSCR = E31, ST_IRR = F31.
Die zweite Kachelreihe (Z. 30–32, Namensziele ST_CF/ST_DSCR/ST_IRR) wird ausgeblendet (P3-12: Start zeigt vier
Kacheln); ihre Formeln bleiben unverändert. Geändert werden nur Darstellung, statische Beschriftungen und nicht
referenzierte Anzeigeformeln.
"""
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Border, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

import core as C
from core import BLUE, GOLD, INK, INK2, LINE2, MIST, MUTED, NAVY, NAVY_2, RED, SKY, WHITE, align, fill, font, side

ui = C                                   # Altname (leitfaden.py importiert Helfer aus diesem Modul)
SHEET = "Start"
HERO_ROWS = range(5, 20)                 # B5:H19 Navy-Fläche, Z. 20 Weißraum
HERO_COLS = "BCDEFGH"
HERO_RULE = C.mix(NAVY, SKY, 0.3)       # zarte Trennlinie AUF Navy (Kennzahlenband) – aus Tokens gemischt, kein Blau-Akzent
# Runde 4 (P1-08): rechte Kante des Heros = NAV_END der festen Reiterleiste (1 417 px = Kante der Schrittseiten)
WIDTHS = {"C": 32, "D": 16, "E": 48, "F": 48, "G": 48}   # C:D = E = F = G = 336 px, H 21 px → Hero bis 1 417 px
EMU = 9525                               # EMU je Pixel
KAUF_ALS = 'CHOOSE(Rechtsform_Idx,"Privatperson","vv-GmbH","GmbH / Holding")'

# Ablauf „So gehen Sie vor“: (Titel, Zielblatt, Zielzelle, Beschreibung ≤ 95 Zeichen, einzeilig)
FLOW = [
    ("Kauf als wählen", "Start", "D23",
     "Privatperson oder Gesellschaft – die Auswahl oben steuert Steuern, Exit und Bankunterlagen."),
    ("Leitfaden: Schritte 01–12", "Leitfaden", None,
     "Vom Objekt bis zum Ergebnis – jede Seite mit Eingaben, Zwischenergebnis und Einordnung."),
    ("Dashboard", "Dashboard", None,
     "Die Gesamtbewertung auf einer Seite: Kennzahlen mit Status, Cashflow und Vermögen."),
    ("Cockpit & Diagramme", "Cockpit", None,
     "Detailkennzahlen und Prüfhinweise im Cockpit, alle Auswertungen als Diagramm."),
    ("Sensitivität", "Sensitivität", None,
     "Break-even-Miete und -Zins, Miete × Zins und IRR-Matrix – wie robust ist die Rechnung?"),
    ("Bankunterlagen", "Bankgespräch", None,
     "Bankgespräch als A4-Übersicht, dazu Haushaltsrechnung und Vermögensaufstellung."),
]
ETAPPE = "ABCDEF"                        # Etappen A–F („Schritt“ bleibt S01–S12 vorbehalten, P37)

# Blätterverzeichnis in Reiterreihenfolge; „{n}“ = Anzahl Diagramme
SHEETS = [
    ("Leitfaden", "Zwölf Schritte vom Objekt zum Ergebnis"),
    ("Dashboard", "Gesamtbewertung auf einer Seite"),
    ("Cockpit", "Detailkennzahlen und Prüfhinweise"),
    ("Diagramme", "{n} Diagramme zur Kalkulation"),
    ("Eingaben", "Alle Eingaben und Profi-Felder"),
    ("Projektion", "Miete, Kosten, Vermögen – 40 Jahre"),
    ("Steuern", "Steuerberechnung je Jahr  (Eingaben: Schritt 09)"),
    ("AfA-Vergleich", "Abschreibungsvarianten im Vergleich"),
    ("Finanzierung", "Tilgungspläne I und II  (Eingaben: Schritt 07)"),
    ("Sensitivität", "Break-even, Miete × Zins, IRR-Matrix"),
    ("Bankgespräch", "Investitionsübersicht für die Bank"),
    ("Haushaltsrechnung", "Selbstauskunft Einnahmen/Ausgaben"),
    ("Vermögensaufstellung", "Selbstauskunft Vermögen"),
    ("Hinweise", "Rechtsgrundlagen und Modellannahmen"),
    ("Konfiguration", "Stammdaten, Tarif 2026, Ampel-Schwellen"),
]

# Anzeigename im Verzeichnis = Reiterbeschriftung (P2-05: ein Name pro Ziel); das Blatt selbst heißt unverändert
SHEET_LABEL = {"Steuern": "Steuer-Tabelle", "Finanzierung": "Tilgungsplan"}


# ============================================================================ lokale Helfer (auch für leitfaden.py)
def rich(*parts):
    """Rich-Text aus (Text, Größe, fett, Farbe)-Teilen – dünne Hülle um core.rich."""
    return C.rich([tuple(p[:4]) for p in parts])


def unmerge_in(ws, c1, r1, c2, r2):
    c1, c2 = C.col(c1), C.col(c2)
    for mr in list(ws.merged_cells.ranges):
        if not (mr.max_row < r1 or mr.min_row > r2 or mr.max_col < c1 or mr.min_col > c2):
            ws.unmerge_cells(str(mr))


def drop_cf(ws, coords):
    """Bedingte Formate entfernen, deren Bereich vollständig in den genannten Zellen liegt."""
    targets = set(coords)
    new = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        cells = {C.L(c) + str(r) for cr in cf.sqref.ranges for (r, c) in cr.cells}
        if cells and cells <= targets:
            continue
        for rule in cf.rules:
            new.add(str(cf.sqref), rule)
    ws.conditional_formatting = new


def cells_of(c1, r1, c2, r2):
    return [f"{C.L(c)}{r}" for r in range(r1, r2 + 1) for c in range(C.col(c1), C.col(c2) + 1)]


def wipe(ws, c1, r1, c2, r2, formulas=False):
    """Bereich leeren: Verbünde lösen, statische Werte, Links und Stile entfernen.
    Formeln bleiben, außer formulas=True (nur für Bereiche ohne referenzierte Formeln)."""
    unmerge_in(ws, c1, r1, c2, r2)
    for c in C.iter_cells(ws, c1, r1, c2, r2):
        if isinstance(c, MergedCell):
            continue
        c.hyperlink = None
        if formulas or not C.is_formula(c.value):
            c.value = None
        c.fill = C.NOFILL
        c.border = Border()
        c.font = font()
        c.alignment = align("general", "center")
        c.number_format = "General"


def link(cell, text, sheet, target=None, size=C.T_BODY, bold=True, color=BLUE, tooltip=None, h="left", indent=1):
    """Textlink (nur in Listen/Fließtext – nie in einer Button-Reihe)."""
    C.text_link(cell, text, sheet, target_cell=target, size=size, bold=bold,
                tooltip=tooltip or f"Zum Blatt „{sheet}“")
    cell.hyperlink.location = C.link_loc(sheet, target)
    cell.font = font(size, bold, color)
    cell.alignment = align(h, "center", indent)


def set_widths(ws, widths):
    """Spaltenbreiten setzen; Bereichs-Dimensionen der Vorlage (z. B. D:G in einem <col>) vorher auftrennen."""
    from openpyxl.worksheet.dimensions import ColumnDimension
    for key, d in list(ws.column_dimensions.items()):
        if d.min and d.max and d.max > d.min:
            for idx in range(d.min + 1, min(d.max, 60) + 1):
                letter = C.L(idx)
                if letter not in ws.column_dimensions or ws.column_dimensions[letter].min != idx:
                    ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=d.width,
                                                                   customWidth=d.customWidth, hidden=d.hidden)
            d.max = d.min
    for k, w in widths.items():
        dim = ws.column_dimensions[k]
        dim.width = w
        dim.min = dim.max = C.col(k)


def heights(ws, spec):
    for r, h in spec.items():
        C.set_height(ws, r, h)


fixed_m = C.fixed_m                      # FIXED mit typografischem Minus (zentral in core, Runde 4)


# Kanonische Kachel-Fußzeilen (P11/P20) – wortgleich mit dem Dashboard; Status setzt C.tile selbst (rechts).
TILE_SUB = {
    "GI": '="Kaufpreis "&FIXED(Kaufpreis,0)&" €  ·  NK "&FIXED(NK_Quote*100,1)&" %"',
    "EK": '="Quote "&FIXED(EK_Quote*100,1)&" %  ·  Beleihung "&FIXED(Beleihung*100,1)&" %"',
    "CF": f'="Jahr 2: "&{fixed_m(C.CF_YEAR2)}&" €"',
    "BMR": "=" + C.threshold_text("BMR"),
    "DSCR": '="Jahr 1  ·  "&' + C.threshold_text("DSCR"),
    "IRR": "=" + C.threshold_text("IRR"),
    "RATE": '="Volltilgung "&Volltilgung_Txt',
}


def clear_rows(ws, r1, r2, c1="B", c2="H"):
    """Nur Stile (keine Formeln) in ausgeblendeten Zeilen neutralisieren."""
    unmerge_in(ws, c1, r1, c2, r2)
    for c in C.iter_cells(ws, c1, r1, c2, r2):
        if isinstance(c, MergedCell):
            continue
        c.fill = C.NOFILL
        c.border = Border()


# ============================================================================ „Kauf als“ (Logik bleibt)
def purchase_selector(wb):
    """Direktauswahl Privat / Kapitalgesellschaft auf der Startseite (steuert den Namen „Rechtsform“)."""
    ws, s09 = wb[SHEET], wb["S09 Steuern"]
    current = s09["D12"].value
    unmerge_in(ws, "D", 23, "H", 23)
    for c in "EFGH":                                   # D23:F23 wird das Feld, G23 der Auswahl-Knopf (selector)
        ws[f"{c}23"].value = None
        ws[f"{c}23"].hyperlink = None
    ws.merge_cells("D23:F23")
    sel = ws["D23"]
    sel.value = current
    sel.protection = Protection(locked=False)
    dv = DataValidation(type="list", formula1="=L_Rechtsform", allow_blank=False, showDropDown=False,
                        showErrorMessage=True, errorStyle="stop", showInputMessage=True)
    dv.promptTitle = "Kauf als"
    dv.prompt = "Privatperson oder Kapitalgesellschaft (vermögensverwaltende bzw. gewerbliche GmbH) aus der Liste wählen."
    dv.errorTitle = "Bitte aus der Liste wählen"
    dv.error = "Privatperson, vermögensverwaltende GmbH oder gewerbliche GmbH/Holding über den Pfeil am Feldrand wählen."
    ws.add_data_validation(dv)
    dv.add("D23")
    wb.defined_names["Rechtsform"] = DefinedName("Rechtsform", attr_text="Start!$D$23")
    # Schritt 09 zeigt die Auswahl nur noch an (verknüpfter Wert mit Link zur Startseite)
    for dvs in list(s09.data_validations.dataValidation):
        if "D12" in str(dvs.sqref):
            s09.data_validations.dataValidation.remove(dvs)
    d12 = s09["D12"]
    d12.value = "=Rechtsform"
    C.input_style(d12, "linked")
    d12.protection = Protection(locked=True)
    d12.hyperlink = Hyperlink(ref="D12", location=C.link_loc("Start", "D23"), display="Kauf als – Startseite",
                              tooltip="Kauf als auf der Startseite ändern")
    s09["F12"].value = "Kauf als auf der Startseite ändern  ›"
    s09["F12"].data_type = "s"


# ============================================================================ Icons (Runde 6)
def icon(ws, cell, name, color=NAVY, px=18, **kw):
    """Icon aus excel/icons.py setzen – ohne harte Abhängigkeit (fehlt Pillow/icons, bleibt das Blatt ohne Icon)."""
    try:
        import icons
        return icons.place_icon(ws, cell, name, color, px=px, **kw)
    except Exception as exc:                             # pragma: no cover
        print(f"   Hinweis start.icon: {name} nicht gesetzt ({exc})")
        return None


def section_icon(ws, row, c1, name, px=18):
    """Abschnittskopf Ebene 1 mit Linien-Icon vor dem Titel (Runde 6): Icon Nachtblau 18 px, 12 px hinter der
    Goldkante, Titel mit Einzug 4 (≈ 36 px) – EIN Stil für alle Abschnittsköpfe von Start und Leitfaden."""
    t = ws.cell(row, C.col(c1))
    t.alignment = align("left", "center", 4)
    icon(ws, f"{c1}{row}", name, NAVY, px, dx=12, valign="middle")


def tile_icon(ws, c1, row, name, px=16):
    """Goldenes Linien-Icon rechts im Nachtblau-Kachelkopf (Runde 6) – Kennzeichen der Kennzahl."""
    icon(ws, f"{c1}{row}", name, GOLD, px, dx=-12, align="right", valign="middle")


# ============================================================================ Cover (Runde 6)
# Titelblatt: Nachtblau über die volle Inhaltsbreite B:H, links der Satzspiegel (Monogramm + Marke · Eyebrow ·
# 30-pt-Titel · Goldlinie · Untertitel · Objektzeile · Aktionen), rechts die goldene Linien-Illustration (icons.cover_art),
# unten die Kennzahlen-Leiste auf NAVY_2 mit Goldlinie als Abschluss.
# Bild-Regel: Bilder liegen über den Zellen und fangen Klicks ab → die Illustration beginnt erst in F8 (unter „Erstellt
# für“ G7) und endet über der Kennzahlen-Leiste; Buttons (C:D, E) und das Eingabefeld G7 bleiben frei.
COVER_H = {5: 14, 6: 16, 7: 20, 8: 30, 9: 14, 10: 45.75, 11: 12, 12: 24, 13: 16, 14: 21, 15: C.H_BTN, 16: 30,
           17: 21, 18: 36, 19: 14}                   # 17/18 Kennzahlen-Leiste, 19 Luft + Goldlinie; Z. 20 Weißraum
ART_ROWS = (8, 16)                                   # Illustration F8 … Ende Z. 16
BAND = [  # Kennzahlen-Leiste im Kachelraster C:D | E | F | G (Icon, Label, Formel, Format)
    ("C", "muenzen", "GESAMTINVESTITION", "=Gesamtinvestition", "eur"),
    ("E", "haus", "KAUFPREIS", "=Kaufpreis", "eur"),
    ("F", "schluessel", "EIGENKAPITAL INKL. RESERVE", "=EK_Bedarf_gesamt", "eur"),
    ("G", "kalender", "HALTEDAUER", "=Haltedauer", "years_n"),
]
BAND_RULE = C.mix(NAVY_2, SKY, 0.28)                  # Trennlinie AUF NAVY_2 (Kennzahlen-Leiste)


def _drop_template_logo(ws):
    """Das Vorlagen-Logo (Bild links oben im Hero) weicht dem Monogramm aus icons.py."""
    keep = []
    for img in ws._images:
        a = getattr(img.anchor, "_from", None)
        if a is not None and a.col <= 3 and 3 <= a.row <= 20:
            continue
        keep.append(img)
    ws._images = keep


def _row_px(ws, r):
    return round((ws.row_dimensions[r].height or 15) / 0.75)


def cover(ws):
    unmerge_in(ws, "B", 5, "H", 21)
    for r in range(5, 20):
        for c in HERO_COLS:
            cell = ws[f"{c}{r}"]
            if cell.coordinate != "G7":                    # Eingabe „Erstellt für“ bleibt (Name Erstellt_fuer)
                cell.value = None
            cell.fill = fill(NAVY_2 if r >= 17 else NAVY)
            cell.border = Border()
            cell.hyperlink = None
    heights(ws, {4: C.H_HDR[4], **COVER_H})
    for r in (20, 21, 22):                                 # Eckmarken der Vorlage (└ ┘) und Restflächen neben dem Hero
        for c in "BH":
            cell = ws[f"{c}{r}"]
            if not C.is_formula(cell.value):
                cell.value = None
            cell.fill = C.NOFILL
            cell.border = Border()

    def put(cell, value, size, bold, color, h="left", v="center", merge_to=None, fmt=None, indent=0):
        if merge_to:
            C.safe_merge(ws, cell[0], int(cell[1:]), merge_to, int(cell[1:]))
        c = ws[cell]
        if isinstance(value, str) and not value.startswith("="):
            C.set_text(c, value)
        else:
            c.value = value
        c.font = font(size, bold, color)
        c.alignment = align(h, v, indent)
        if fmt:
            c.number_format = fmt
        return c

    # Kopfzeile des Titelblatts: Monogramm + Marke links, „Erstellt für“ rechts (G6/G7)
    put("C6", "MM HOLDING GMBH", C.T_MICRO, True, WHITE, v="bottom", merge_to="E", indent=9)
    put("C7", "Unternehmerische Beteiligungen  ·  Weingarten", C.T_MICRO, False, SKY, v="top", merge_to="E", indent=9)
    created_for(ws)

    # Satzspiegel links (C:E): Eyebrow · Titel · Goldlinie · Untertitel · Objektzeile
    put("C9", "INVESTMENT-KALKULATION   ·   WOHNIMMOBILIEN", C.T_MICRO, True, SKY, merge_to="E")
    put("C10", "Immobilien-Kalkulation", C.T_HERO, True, WHITE, merge_to="E")
    ws["C11"].border = Border(bottom=side("medium", GOLD))      # kurze Goldlinie unter dem Titel (Hero-Linie)
    put("C12", "Kauf  ·  Finanzierung  ·  Cashflow  ·  Steuern  ·  Exit", C.T_H3, False, MIST, merge_to="E")
    put("C13", '=Obj_Name&"   ·   Stand "&TEXT(DAY(TODAY()),"00")&"."&TEXT(MONTH(TODAY()),"00")&"."&YEAR(TODAY())'
               f'&"   ·   Kauf als: "&{KAUF_ALS}', C.T_MICRO, False, SKY, merge_to="E")

    # Aktionen (C:D primär, E ghost) – exakt im Kachelraster; 3-px-Fuge in Heldenfarbe
    C.btn(ws, "C", 15, "D", "Leitfaden starten  ›", "Leitfaden", "primary",
          tooltip="Zur Übersicht der zwölf Schritte – Schritt 01 beginnt mit dem Objekt")
    C.btn(ws, "E", 15, "E", "Dashboard  ›", "Dashboard", "ghost", tooltip="Gesamtbewertung auf einer Seite")
    gap = side("thick", NAVY)
    g2 = side("medium", GOLD)
    ws["C15"].border = Border(top=g2, bottom=g2, left=g2)
    ws["D15"].border = Border(top=g2, bottom=g2, right=g2)
    b = ws["E15"].border
    ws["E15"].border = Border(top=b.top, bottom=b.bottom, left=gap, right=b.right)
    heights(ws, {15: C.H_BTN})

    # Kennzahlen-Leiste (NAVY_2, Z. 17–19): Icon + Label · Wert 22 pt · senkrechte Trennlinien im Kachelraster
    rule = side("thin", BAND_RULE)
    for c, ic, lab, val, fmt in BAND:
        c2 = "D" if c == "C" else c
        if c2 != c:
            C.safe_merge(ws, c, 17, c2, 17)
            C.safe_merge(ws, c, 18, c2, 18)
        put(f"{c}17", lab, C.T_MICRO, True, SKY, v="bottom", indent=4)
        put(f"{c}18", val, C.T_H1, True, WHITE, fmt=C.NUMFMT[fmt], indent=1)
        if c != "C":
            for r in (17, 18, 19):
                ws[f"{c}{r}"].border = Border(left=rule)
    for c in HERO_COLS:                                                 # Abschluss: Goldlinie (Hero-Linie)
        cur = ws[f"{c}19"].border
        ws[f"{c}19"].border = Border(left=cur.left, bottom=side("thick", GOLD))
    ws["G19"].border = Border(left=rule, bottom=side("thick", GOLD))

    # Bilder (nach allen Maßen): Monogramm, Illustration, Icons der Kennzahlen-Leiste
    _drop_template_logo(ws)
    try:
        import icons
        mono = 60
        top = _row_px(ws, 5) + (_row_px(ws, 6) + _row_px(ws, 7) - mono) // 2
        icons.place_image(ws, "C5", icons.monogram_path(mono), mono, mono, dx=0, dy=top)
        w = C.span_px(ws, "F", "H")
        h = sum(_row_px(ws, r) for r in range(ART_ROWS[0], ART_ROWS[1] + 1))
        icons.place_image(ws, f"F{ART_ROWS[0]}", icons.cover_art_path(w, h, background=False, focus="center"), w, h)
    except Exception as exc:                                            # pragma: no cover
        print(f"   Hinweis start.cover: Illustration nicht gesetzt ({exc})")
    for c, ic, *_ in BAND:
        icon(ws, f"{c}17", ic, GOLD, 16, dx=10, valign="bottom", dy=-1)


# ============================================================================ Erstellt für (Deckblatt-Feld)
def created_for(ws):
    """Eingabefeld „Erstellt für“ oben rechts im Hero + Name Erstellt_fuer (Start!$G$7).
    Gedämpft (P2-15): FFF5D6, Rahmen 1 pt E6CB77, 10 pt, Zeilenhöhe 18."""
    wb = ws.parent
    lab, inp = ws["G6"], ws["G7"]
    C.set_text(lab, "ERSTELLT FÜR")
    lab.font = font(C.T_MICRO, True, SKY)
    lab.alignment = align("left", "center", 1)
    if inp.value is None:
        C.set_text(inp, "Max Mustermann")
    C.input_style(inp, "required")
    inp.font = font(C.T_BODY, True, C.INPUT_FG)
    inp.alignment = align("left", "center", 1)
    dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="60", allow_blank=True,
                        showErrorMessage=True, errorStyle="stop", showInputMessage=True)
    dv.promptTitle = "Erstellt für"
    dv.prompt = "Name oder Firma des Investors – erscheint auf Dashboard, Bankgespräch und den Schrittseiten."
    dv.errorTitle = "Zu lang"
    dv.error = "Bitte höchstens 60 Zeichen eingeben."
    ws.add_data_validation(dv)
    dv.add("G7")
    wb.defined_names["Erstellt_fuer"] = DefinedName("Erstellt_fuer", attr_text="Start!$G$7")


# ============================================================================ Kauf als – Entscheidungskarte
SEL_LINE = GOLD                        # Edel-Gold C9A14A: Rahmen der Pflichtauswahl (Eingabefamilie, P06, Runde 5)


def selector(ws):
    """Pflichtauswahl als erkennbare Auswahlliste (P06, P3-01): Beschriftung „KAUF ALS / Pflichtfeld“ (ohne ▾),
    Feld D23:F23 (gelb, Datenüberprüfung Liste) + Hinweis G23 „Liste öffnen: Alt + ↓“ (weiß, linke Trennlinie E6CB77,
    Sprung auf D23) – beide in EINEM Rahmen der Eingabefamilie (medium C9A94A, links 3 px). Das ▾ steht nur noch
    einmal: in der Erklärzeile D24 unter dem Feld."""
    unmerge_in(ws, "C", 21, "H", 22)
    C.section(ws, 21, "C", "G", "Kauf als – Privatperson oder Gesellschaft", meta="Pflichtauswahl")
    section_icon(ws, 21, "C", "waage")
    lab = ws["C23"]
    lab.value = rich(("KAUF ALS", C.T_BODY, True, NAVY), ("\nPflichtfeld", C.T_MICRO, False, MUTED))
    lab.font = font(C.T_BODY, True, NAVY)
    lab.alignment = align("left", "center", 1, wrap=True)
    lab.border = Border()
    lab.fill = C.NOFILL
    lab.hyperlink = None
    unmerge_in(ws, "D", 23, "G", 23)
    for c in "EFG":
        ws[f"{c}23"].value = None
        ws[f"{c}23"].hyperlink = None
    C.safe_merge(ws, "D", 23, "F", 23)
    frame = side("medium", SEL_LINE)
    for c in "DEF":
        cell = ws[f"{c}23"]
        cell.fill = fill(C.INPUT_BG)
        cell.border = Border(top=frame, bottom=frame, left=side("thick", SEL_LINE) if c == "D" else None)
    sel = ws["D23"]
    sel.font = font(C.T_BODY, True, C.INPUT_FG)
    sel.alignment = align("left", "center", 1)
    # Hinweis im Feld (kein Knopf, keine Fläche): Tastenkürzel für die Liste; Klick markiert das Feld D23
    hint = ws["G23"]
    C.set_text(hint, "Liste öffnen:  Alt + ↓")
    hint.font = font(C.T_LABEL, False, BLUE)
    hint.alignment = align("center", "center")
    hint.fill = fill(WHITE)
    hint.border = Border(top=frame, bottom=frame, right=frame, left=side("thin", C.INPUT_LINE))
    hint.hyperlink = Hyperlink(ref="G23", location=C.link_loc(SHEET, "D23"), display="Kauf als",
                               tooltip="Feld „Kauf als“ markieren – dann Alt + ↓ oder den Pfeil am Feldrand")
    ws["H23"].fill = C.NOFILL
    ws["H23"].border = Border()

    wipe(ws, "C", 24, "H", 24)
    C.safe_merge(ws, "D", 24, "G", 24)
    C.set_text(ws["D24"], "▾  Aus der Liste wählen: Privatperson · vv-GmbH · gewerbliche GmbH/Holding"
                          "   ·   wirkt auf S09 Steuern, S11 Exit, Dashboard und Bankgespräch")
    ws["D24"].font = font(C.T_MICRO, False, MUTED)
    ws["D24"].alignment = align("left", "center", 1)
    for c in HERO_COLS:
        for r in (20, 22):
            ws[f"{c}{r}"].fill = C.NOFILL
            ws[f"{c}{r}"].border = Border()
    heights(ws, {20: 12, 22: 6, 23: 36, 24: 18, 25: C.H_GAP})


# ============================================================================ Kacheln
TILE_ICON = {"CF": "muenzen", "BMR": "prozent", "DSCR": "schild", "IRR": "trend", "GI": "haus",
             "EK": "schluessel", "RATE": "bank"}   # Kachelkopf-Icons (auch Leitfaden)
TILES = [  # (c1, c2, KPI, Wertformel) – Teilmenge der kanonischen Reihenfolge C.KPI_ORDER (P20)
    ("C", "D", "CF", "=CF_nSt_Monat_J1"),
    ("E", "E", "BMR", "=Bruttomietrendite"),
    ("F", "F", "DSCR", "=DSCR_J1"),
    ("G", "G", "IRR", "=EK_IRR"),
]


def tiles(ws):
    """Vier Kacheln mit C.tile (dark) – EINE Anatomie wie Dashboard/Leitfaden (P10/P11/P20):
    Kopfstreifen mit kanonischem Label · Wert in Statusfarbe · Fußzeile Kontext links, Status rechts
    (2-spaltig als Chip, 1-spaltig inline), Rinnen 3 px horizontal = vertikal (gap='auto')."""
    unmerge_in(ws, "C", 26, "G", 26)
    C.section(ws, 26, "C", "G", "Aktuelle Kalkulation im Überblick", meta="Details im Dashboard  ›")
    section_icon(ws, 26, "C", "diagramm")
    m = ws["G26"]
    m.hyperlink = Hyperlink(ref="G26", location=C.link_loc("Dashboard"), display="Dashboard",
                            tooltip="Gesamtbewertung auf einer Seite")
    for r in range(27, 33):
        for c in "BCDEFGH":
            ws[f"{c}{r}"].border = Border()
            ws[f"{c}{r}"].fill = C.NOFILL
    drop_cf(ws, cells_of("C", 27, "G", 32))
    unmerge_in(ws, "C", 27, "G", 32)
    for c in "CDEFG":                                   # Kopf- und Fußzeile (Z. 27/29) sind in der Vorlage reine Texte
        ws[f"{c}29"].value = None
    ws["D27"].value = ws["D28"].value = None
    for c1, c2, key, val in TILES:
        vref = val[1:]
        C.tile(ws, c1, c2, 27, 28, 29, kpi=key, value=val, value_ref=vref, sub=TILE_SUB[key],
               label=C.kpi_label(key, caps=True, formula=True), gap_right=c2 != "G")
        tile_icon(ws, c1, 27, TILE_ICON[key])
    # zweite Reihe der Vorlage (Namensziele ST_CF/ST_DSCR/ST_IRR) ausblenden – Formeln bleiben unverändert
    clear_rows(ws, 30, 32)
    for c, fmt in (("C", C.NUMFMT["eur"]), ("E", C.NUMFMT["dscr"]), ("F", C.NUMFMT["pct1"])):
        ws[f"{c}31"].number_format = fmt
    C.hide_rows(ws, 30, 32)
    heights(ws, {33: C.H_GAP})


# ============================================================================ Ihr Weg durch das Tool (Etappen-Karten)
FLOW_ICON = ("waage", "dokument", "diagramm", "ziel", "lupe", "bank")      # je Etappe A–F
CARD_ROWS = (35, 38, 41)                                                  # Titelzeile je Kartenreihe (+1 = Text)
CARD_COLS = (("C", "E"), ("F", "G"))                                      # 2 Karten je Reihe, Fuge = Kachelrinne
BADGE = 36


def flow(ws):
    """Sechs Etappen-Karten A–F (2 × 3) statt Linkliste: Navy-Badge mit Gold-Icon · „A  Titel  ›“ (12,5 pt) ·
    Beschreibung (9 pt) – die ganze Karte ist ein Link. Karten FBFAF7 mit feinem Rahmen, Rinne 3 px weiß."""
    wipe(ws, "C", 34, "H", 45)
    C.section(ws, 34, "C", "G", "Ihr Weg durch das Tool", meta="6 Etappen  ·  A–F")
    section_icon(ws, 34, "C", "pfeil")
    heights(ws, {r: 22 for r in CARD_ROWS})
    heights(ws, {r + 1: 20 for r in CARD_ROWS})
    heights(ws, {37: 6, 40: 6, 43: C.H_GAP, 44: 4, 45: 4})              # A44 = Sprungziel „≡ Alle Bereiche“
    edge = side("thin", LINE2)
    gutter = side("thick", WHITE)
    for i, (title, sheet, target, desc) in enumerate(FLOW):
        r = CARD_ROWS[i // 2]
        c1, c2 = CARD_COLS[i % 2]
        tip = "Zur Auswahl „Kauf als“" if sheet == SHEET else f"Zum Blatt „{sheet}“"
        for rr in (r, r + 1):
            C.safe_merge(ws, c1, rr, c2, rr)
            ws[f"{c1}{rr}"].hyperlink = None
            cols = range(C.col(c1), C.col(c2) + 1)
            for k, cc in enumerate(cols):
                cell = ws.cell(rr, cc)
                cell.fill = fill(C.TINT_XL)
                cell.border = Border(top=edge if rr == r else None, bottom=edge if rr == r + 1 else None,
                                     left=edge if (k == 0 and c1 == "C") else None,
                                     right=(gutter if c2 == "E" else edge) if cc == cols[-1] else None)
        t, d = ws[f"{c1}{r}"], ws[f"{c1}{r + 1}"]
        t.hyperlink = Hyperlink(ref=t.coordinate, location=C.link_loc(sheet, target), display=title, tooltip=tip)
        t.value = rich((ETAPPE[i], C.T_H3, True, C.GOLD_INK), ("   " + title, C.T_H3, True, NAVY),
                       ("   ›", C.T_H3, True, BLUE))
        t.font = font(C.T_H3, True, NAVY)
        t.alignment = align("left", "bottom", 7)
        C.set_text(d, desc)
        d.font = font(C.T_SMALL, False, MUTED)
        d.alignment = align("left", "top", 7)
        h = _row_px(ws, r) + _row_px(ws, r + 1)
        icon(ws, f"{c1}{r}", FLOW_ICON[i], GOLD, BADGE, dx=14, dy=(h - BADGE) / 2, bg=NAVY)


# ============================================================================ Legende + Tipp | Blätter
def legend_and_sheets(ws):
    """Links Legende (C Muster | D:E Erklärung) + Kasten „Tipp & Support“, rechts Blätterverzeichnis
    (F Link | G Beschreibung). Beide Spalten gleich hoch (Z. 48–62).
    Legende (P3-20): die vier Eingabezustände über C.input_legend (C.INPUT_STATES), danach Verknüpfung, Berechnung,
    Ergebnis, Negativ-Rot, Status und Prüfhinweis. Blattschutz und Navigation stehen im Tipp-Kasten."""
    wipe(ws, "C", 46, "H", 63)
    C.section(ws, 46, "C", "E", "Legende")
    C.section(ws, 46, "F", "G", "Alle Blätter", meta="Schritte S01–S12: siehe Leitfaden")
    section_icon(ws, 46, "C", "info")
    section_icon(ws, 46, "F", "dokument")
    ws["E46"].border = Border(bottom=side("thin", LINE2), right=side("thick", WHITE))   # Unterlinie wie section()
    C.set_height(ws, 47, 6)
    gutter = side("thick", WHITE)                 # Rinne Legende | Alle Blätter (P26) – wie die Kachelrinnen
    top = 48

    # Eingabezustände (Muster in C mit Kurzwort, Erklärung in D:E)
    states = {  # Kurzwort im Muster · Erklärung (Satzbau wie die übrigen Legendenzeilen: „Farbe – Bedeutung“)
        "required": ("Eingabe", "Gelb – hier eingeben, Beispielwerte überschreiben"),
        "optional": ("optional", "Gestrichelt – darf leer bleiben"),
        "override": ("überschreibbar", "Aus der Kalkulation vorgeschlagen – bei Bedarf überschreiben"),
        "inactive": ("inaktiv", "Grau – für die gewählte Methode ohne Wirkung"),
    }
    n_states = len(C.INPUT_STATES)
    for i in range(n_states):
        C.safe_merge(ws, "D", top + i, "E", top + i)
    C.input_legend(ws, top, [(f"C{top + i}", f"D{top + i}") for i in range(n_states)])
    for i, (state, word) in enumerate(C.INPUT_STATES):
        r = top + i
        sw, tx = ws[f"C{r}"], ws[f"D{r}"]
        short, expl = states.get(state, (word, word))
        C.set_text(sw, short)
        sw.alignment = align("left", "center", 1)
        C.set_text(tx, expl)
        tx.font = font(C.T_SMALL, False, MUTED)
        tx.alignment = align("left", "center", 1)
        C.hairline(ws, r, "D", "E")
        C.set_height(ws, r, C.H_ROW)

    legend = [
        ("Verknüpfung", "Blau – übernommen aus einer Eingabe an anderer Stelle"),
        ("Berechnung", "Schwarz – berechnet, bitte nicht ändern"),
        ("Ergebnis", "Hervorgehoben – Summe bzw. Blockergebnis"),
        (f"{C.MINUS}1.234 €", C.NEG_RULE_TEXT),
        ("Status", None),
        ("Prüfhinweis", None),
    ]
    r1 = top + n_states
    for i, (term, text) in enumerate(legend):
        r = r1 + i
        c = ws[f"C{r}"]
        C.set_text(c, term)
        c.font = font(C.T_BODY, False, INK)
        c.alignment = align("left", "center", 1)
        C.safe_merge(ws, "D", r, "E", r)
        d = ws[f"D{r}"]
        if text:
            C.set_text(d, text)
        d.font = font(C.T_SMALL, False, MUTED)
        d.alignment = align("left", "center", 1)
        C.hairline(ws, r, "C", "E")
        C.set_height(ws, r, C.H_ROW)
    ws[f"C{r1}"].font = font(C.T_BODY, False, BLUE)
    C.sum_row(ws, r1 + 2, "C", "C", "final")
    ws[f"C{r1 + 2}"].font = font(C.T_BODY, True, NAVY)
    ws[f"C{r1 + 3}"].font = font(C.T_BODY, True, RED)
    ws[f"D{r1 + 4}"].value = C.status_legend(("erfüllt", "prüfen", "kritisch"), size=C.T_SMALL)
    ws[f"D{r1 + 5}"].value = rich(("▲", C.T_SMALL, True, RED), ("  Warnung      ", C.T_SMALL, False, MUTED),
                                  ("ⓘ", C.T_SMALL, True, BLUE), ("  Information", C.T_SMALL, False, MUTED))
    last_leg = r1 + len(legend) - 1                                # 57

    # Tipp & Support (Einordnungs-Box aus core, neutral) unter der Legende: Kopf 59, Körper 60–62 (3 Zeilen)
    r0 = last_leg + 2
    C.set_height(ws, r0 - 1, C.H_ROW)                             # Luft; Zeile gehört rechts zur Blätterliste (18 pt Raster)
    C.callout_box(ws, "C", r0, "E", r0 + 1, 62, title="Tipp & Support", pill=False, fit=None,
                  text="Hilfe zu jedem Eingabefeld: Zelle markieren – der Eingabehinweis erscheint.\n"
                       "Navigation: Reiter oben, „Weiter  ›“ führt zum nächsten Schritt · Strg + Bild ↓ / Bild ↑.\n"
                       "Blattschutz ohne Passwort · jedes Blatt ist für den Druck auf A4 eingerichtet.")
    for rr in range(r0, 63):
        C.set_height(ws, rr, C.H_ROW)

    n_charts = len(ws.parent["Diagramme"]._charts) if "Diagramme" in ws.parent.sheetnames else 0
    r = top
    for name, desc in SHEETS:
        if name not in ws.parent.sheetnames and name != "Dashboard":      # Dashboard entsteht erst nach den Layouts
            continue
        link(ws[f"F{r}"], f"{SHEET_LABEL.get(name, name)}  ›", name)
        g = ws[f"G{r}"]
        C.set_text(g, desc.format(n=n_charts or "Alle"))
        g.font = font(C.T_SMALL, False, MUTED)
        g.alignment = align("left", "center", 1)
        C.hairline(ws, r, "F", "G")
        C.set_height(ws, r, C.H_ROW)
        r += 1
    # Rinne zwischen den beiden Spalten: weiße 3-px-Kante rechts an E (unterbricht Haarlinien und Flächen)
    from copy import copy
    for rr in range(46, 63):
        cell = ws[f"E{rr}"]
        b = copy(cell.border)
        cell.border = Border(left=b.left, top=b.top, bottom=b.bottom, right=gutter)
        f_ = ws[f"F{rr}"]
        fb = copy(f_.border)
        if rr != 46:                                      # F46 behält die Akzentkante des Abschnittskopfs
            f_.border = Border(left=gutter, top=fb.top, bottom=fb.bottom, right=fb.right)


def foot(ws):
    wipe(ws, "C", 63, "H", 66)
    C.set_height(ws, 63, C.H_GAP)
    C.footer(ws, 64, "C", "G")
    C.set_height(ws, 66, 15)


def layout(ws):
    set_widths(ws, WIDTHS)
    cover(ws)
    selector(ws)
    tiles(ws)
    flow(ws)
    legend_and_sheets(ws)
    foot(ws)
    ws.freeze_panes = "A4"
    ws.sheet_view.showRowColHeaders = False
    C.cf_close(ws)


def apply(wb):
    purchase_selector(wb)
    layout(wb[SHEET])
