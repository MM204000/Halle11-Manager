"""Kopfleiste (Farbfläche Z. 1–3), Monogramm-Entfernung und Zeilen für Schritt- und Sprungleisten.

Die anklickbaren Reiter, die Schritt-Leiste und die Sprungleisten setzt navigation.py nach der Neuberechnung
als Formen darüber. Hier wird nur die Zellfläche vorbereitet (Farben, Zeilenhöhen) – keine Werte, keine Formeln.
"""
import re

from openpyxl.styles import Border, Side

from openpyxl.utils import column_index_from_string, get_column_letter

from core import CONTENT_EDGE, GOLD, H_GAP, NAVY, NOFILL, SKY, col_px, fill, mix

# Runde 4 (P1-08): Die Reiterleiste hat auf allen Blättern dieselbe feste Pixelgeometrie (navigation.NAV_X …
# NAV_END); das Navy-Band reicht bis navigation.BAND_END (rechts derselbe Innenabstand wie links), bei etwas
# breiterem Inhalt bis zur Inhaltskante, auf den Jahrestabellen nie über die Jahresspalten. Hier wird die
# Navy-Fläche nur großzügig vorgelegt (bis BAND_PREFILL_PX bzw. core.CONTENT_EDGE) – navigation.trim_band schneidet
# sie nach der Neuberechnung exakt zu, navigation.band_extension ergänzt den Rest als (nicht gedruckte) Fläche.
JUMP_ROW_SHEETS = ("Eingaben", "Diagramme")  # Zeile 8 trägt die Sprungleiste (navigation.jump_bar)
BAND_TO = CONTENT_EDGE                       # Kompatibilität für Altaufrufer
BAND_PREFILL_PX = 1800                       # Vorlage der Navy-Fläche; Zuschnitt auf die Inhaltskante: navigation.py
# Deckblatt (Runde 6): auf „Start“ gehen Kopfleiste und Hero ineinander über – keine Goldlinie in Z. 3, stattdessen
# eine feine Trennlinie auf Nachtblau (dieselbe Mischung wie die Hero-Haarlinien) und Z. 4 als navy Fuge bis zum Hero.
# navigation.py schneidet Z. 1–4 danach exakt auf die Hero-Kanten zu (eine durchgehende Fläche, kein Absatz).
COVER = "Start"
COVER_RULE = mix(NAVY, SKY, 0.3)


def band_end_col(ws):
    """Letzte Spalte der vorgelegten Kopfleisten-Fläche (Zellen).

    Großzügig: alle Spalten bis BAND_PREFILL_PX, mindestens bis core.CONTENT_EDGE bzw. zum breitesten Diagramm –
    spaltenbasiert, weil die Blatt-Module die Spaltenbreiten erst nach diesem Modul setzen. Den exakten Zuschnitt
    auf die rechte Inhaltskante (= rechte Kante der Reiterleiste) macht navigation.trim_band am fertigen Blatt."""
    last = column_index_from_string(CONTENT_EDGE[ws.title]) if ws.title in CONTENT_EDGE else 1
    for ch in getattr(ws, "_charts", []):
        to = getattr(ch.anchor, "to", None)
        if to is not None:
            last = max(last, to.col + (1 if to.colOff else 0))
    px, c = 0, 1
    while c < 400:
        w = col_px(ws, get_column_letter(c))          # ausgeblendete Spalten = 0 px (core.col_px)
        if px + w > BAND_PREFILL_PX:
            break
        px += w
        c += 1
    return max(last, c - 1)


def masthead(ws):
    """Kopfleiste als Farbfläche: Z. 1 (6 pt) und 2 (33 pt) Navy, Z. 3 (3 pt) Goldlinie (Runde 5). Endet nach dem
    Zuschnitt (navigation.trim_band) an der Bandkante (navigation.Frame.band); dahinter bleibt die Kopfzone weiß."""
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row <= 2 <= mr.max_row:
            ws.unmerge_cells(str(mr))
    for c in ws[2]:
        if not (isinstance(c.value, str) and c.value.startswith("=")):
            c.value = None
        c.hyperlink = None
    last_col = band_end_col(ws)
    cover = ws.title == COVER
    for r, h in ((1, 6), (2, 33), (3, 3)):
        ws.row_dimensions[r].height = h
        for cc in range(1, max(last_col, ws.max_column) + 1):
            c = ws.cell(r, cc)
            c.fill = fill(GOLD if r == 3 and not cover else NAVY) if cc <= last_col else NOFILL
            c.border = Border(bottom=Side("thin", color=COVER_RULE)) if cover and r == 3 and cc <= last_col \
                else Border()
    if cover:                                  # Fuge Z. 4 navy (Höhe setzt die Kopfschablone), nur leere Zellen
        for cc in range(1, last_col + 1):
            c = ws.cell(4, cc)
            if c.value is None:
                c.fill = fill(NAVY)


def remove_monogram(ws):
    """P3-01: das kleine Monogramm-Bild in A1/A2 entfällt – die Marke „MM HOLDING“ (Form) genügt."""
    keep = []
    for img in ws._images:
        a = getattr(img.anchor, "_from", None)
        small = (getattr(img, "width", 0) or 0) <= 64 and (getattr(img, "height", 0) or 0) <= 64
        if a is not None and a.col == 0 and a.row <= 1 and small:
            continue
        keep.append(img)
    ws._images = keep


def stepper_row(ws):
    """Leitfaden-Seiten: Punktreihe (●○○) entfernen, Zeile 8 als Schritt-Leiste (30 pt), Zeile 9 = eine
    Standard-Fuge (core.H_GAP) bis zum ersten Abschnittskopf."""
    if not re.match(r"S\d\d ", ws.title):
        return
    for c in ws[8]:
        if isinstance(c.value, str) and set(c.value) <= set("●○ "):
            c.value = None
    ws.row_dimensions[8].height = 30
    ws.row_dimensions[9].height = H_GAP


def jump_row(ws):
    """Eingaben/Diagramme: Zeile 8 als ruhige Zeile für die Sprungleiste (22 px Reiter, mittig)."""
    if ws.title not in JUMP_ROW_SHEETS:
        return
    if all(c.value is None for c in ws[8]):
        ws.row_dimensions[8].height = 30


def apply(wb):
    for ws in wb.worksheets:
        remove_monogram(ws)
        if ws.title == "Dashboard":
            continue
        masthead(ws)
        stepper_row(ws)
        jump_row(ws)
