"""Blatt „Cockpit“: Seitenkopf mit Gesamtbewertung, sechs KPI-Kacheln (drei Paare „Betrag | Rendite“ im Raster
B | C:D · F | G:H · J | K:L), zwei Kartenreihen, Prüfhinweise, Diagrammzeile und Fuß.

Runde 2 – durchgängig auf die zentralen Bausteine aus core.py umgestellt (CORE_API.md):
  C.page_header (Seitenkopf P1-18) · C.tile (Kachel P1-11/P1-10, dreizeilig, Wert neutral, Status-Chip)
  C.section (Abschnittskopf Ebene 1/2, P1-12) · C.sum_row (Summenstufen P1-14) · C.status_cf / C.neg_red (P1-10)
  C.NUMFMT (Zahlenformat-Katalog P2-12/P3-08) · C.link_loc / C.link_row (Sprungziele P1-08) · C.cf_close.
Weitere Planpunkte: P2-14 (Prüfhinweise ohne Emoji, Semantik über Farbe/Kante), P3-02 (Unterzeilen 9 pt grau,
Abstände, Drill-down-Links je Blockkopf), P3-03 („↑ Übersicht“), P3-08 (Steuerwirkungs-Konvention).

Runde 3: P07 (Steuerwirkung C42:D42 in Cash-Sicht per Format), P11/P20 (C.tile ohne eigene Statustexte, kanonischer
KPI-Satz GI · EK · CF | BMR · DSCR · IRR, IRR-Label mit Haltedauer), P15 (Negativ-Rot über sum_row + neg_red),
P21 (DSCR-Schwelle „Ziel ≥ 1,20×“ im Prüfhinweis), P32 (Zähler, Abschlusslinie, Z. 56 = 6 pt), P37 (Dachzeile
„Ergebnis · Übersicht auf einer Seite“), P41 (Gesamtbewertung und Hinweise ohne Signalfläche), P16 („@“ → Standard),
Weißraum: die Lücken F28:L30 und B45:D47 tragen jetzt Kennzahlen (Anzeigeformeln in leeren Zellen).

Rechenlogik bleibt unberührt: Es ändern sich nur Stile, Zahlenformate, Zeilenhöhen, Verbünde, statische
Beschriftungen und nicht referenzierte Anzeigeformeln (B5/B6, B42, B49:B55, C9/G9/K9, K15, G26, G27, G35 …).
Die benannten Zellen (u. a. J9 = CF_nSt_Monat_J1) bleiben, wo sie sind. Die Spalten N:X (alter Zeitverlauf,
Kacheln EKR_Tile/IRR_Tile) bleiben erhalten und sind nur ausgeblendet.
"""
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Border, Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.hyperlink import Hyperlink

import core as C

try:                                   # Runde 6: Icon-System (M) und Sparklines (K) – fehlt ein Modul, bleibt das Blatt ohne
    import icons as ICN
except Exception:  # noqa: BLE001
    ICN = None
try:
    import sparklines as SP
except Exception:  # noqa: BLE001
    SP = None
from core import (ACCENT, BLUE, INK, INK2, LINE, LINE2, MUTED, NAVY, NOFILL, NUMFMT, RECHTSFORM_SHORT, RED,
                  RED_BG, RED_LINE, T_BODY, T_MICRO, T_SMALL, TINT_XL, WHITE, align, col, fill, font, is_formula,
                  iter_cells, safe_merge, set_height, set_text, side)

SHEET = "Cockpit"
BLOCKS = (("B", "C", "D"), ("F", "G", "H"), ("J", "K", "L"))   # (Beschriftung, Wert, Randspalte)
FIRST, LAST = "B", "L"
NBSP = " "

# --------------------------------------------------------------------------------------------- Inhalte
# Zeilentypen: "row" Datenzeile · "sub" Unterposition (Label und Wert 9 pt grau) · "sum" Zwischensumme (Stufe 1) ·
# "result" Blockergebnis (Stufe 2) · "memo" nachrichtlich · "l2" Unterabschnitt.
# fmt: Schlüssel aus core.NUMFMT oder ein Formatstring; label None = Beschriftung unverändert lassen.
# Kaufpreis je m² ohne Nachkommastellen, Einheit im Format wie Miete je m² (P2-07: € je m² immer „€/m²“)
EUR_M2_0 = '#,##0" €/m²";"' + C.MINUS + '"#,##0" €/m²";"–"'
UPPER = {
    "B": [
        (14, "row", "Objektart", None),
        (15, "row", "Wohn-/Nutzfläche", "m2"),
        (16, "row", "Baujahr", "year"),
        (17, "row", "Kaufpreis", "eur"),
        (18, "sub", "Kaufpreis je m² (Immobilie)", EUR_M2_0),
        (19, "row", "+ Kaufnebenkosten", "eur"),
        (20, "sub", "Kaufnebenkostenquote", "pct1"),
        (21, "row", "+ Finanzierungsnebenkosten", "eur"),
        (22, "row", "+ Maßnahmen Jahr 1", "eur"),
        (23, "result", "= Gesamtinvestition", "eur"),
        (25, "l2", "Kaufpreisaufteilung (AfA-Bemessungsgrundlage)", None),
        (26, "row", "Bodenanteil am Kaufpreis", "pct1"),
        (27, "row", "AK Grund und Boden (inkl. Nebenkosten)", "eur"),
        (28, "sum", "= AfA-Basis Gebäude (inkl. Nebenkosten)", "eur"),
        (29, "row", "Bewegliche Gegenstände (eigene AfA)", "eur"),
        (30, "row", "Erhaltungsrücklage (nicht abschreibbar)", "eur"),
    ],
    "F": [
        (14, "row", "Nettokaltmiete pro Monat (Soll)", "eur"),
        (15, "sub", "Miete je m² und Monat", "eur_qm"),
        (16, "row", "Jahresnettokaltmiete (Soll)", "eur"),
        (17, "row", "Nettokaltmiete Ist Jahr 1 (nach Ausfall)", "eur"),
        (18, "row", "– Bewirtschaftung Jahr 1 (inkl. Rücklagen)", "eur"),
        (19, "sub", "in % der Nettokaltmiete", "pct1"),
        (20, "result", "= Einnahmenüberschuss (NOI) Jahr 1", "eur"),
        (22, "l2", "Rendite, Annahmen und Break-even", None),
        (23, "row", "Bruttomietrendite (Jahresmiete / Kaufpreis)", "pct1"),
        (24, "row", "Nettomietrendite (NOI / Gesamtinvestition)", "pct1"),
        (25, "row", "Kaufpreisfaktor (Kaufpreis / Jahresmiete)", "mult1"),
        (26, "row", "Mietsteigerung · Kosten · Wert p. a.", None),
        (27, "row", "Mietausfallwagnis · Leerstand Jahr 1", None),
        (28, "row", "Eigenkapitalrendite Jahr 1", "pct1"),
        (29, "row", "Break-even-Miete v. St. / Monat", "eur"),
        (30, "row", "Break-even-Miete n. St. / Monat", "eur"),
    ],
    "J": [
        (14, "row", "Darlehen gesamt", "eur"),
        (15, "sub", "Darlehen I · Darlehen II", None),
        (16, "row", "Eigenkapital (ohne Reserve)", "eur"),
        (17, "row", "Eigenkapitalquote", "pct1"),
        (18, "row", "Beleihungsauslauf (Darlehen / Kaufpreis)", "pct1"),
        (19, "row", "Gewichteter Sollzins Jahr 1", "pct2"),
        (21, "l2", "Kapitaldienst und Zinsrisiko", None),
        (22, "row", "Kapitaldienst pro Monat (Jahr 1)", "eur"),
        (23, "row", "Kapitaldienst pro Jahr (Jahr 1)", "eur"),
        (24, "row", "Kapitaldienstdeckung (DSCR)", "dscr"),
        (25, "row", "Restschuld nach Zinsbindung (Darlehen I)", "eur"),
        (26, "row", "Volltilgung (Restschuld = 0)", None),
        (27, "row", f"Anschlusszins +1{NBSP}%-Pkt. kostet p. a.", "eur"),
        (28, "row", "Zinsbindung Darlehen I", "years_n"),
        (29, "row", "Anschlusszins nach Zinsbindung (Annahme)", "pct2"),
        (30, "row", "Anfängliche Tilgung Darlehen I", "pct2"),
    ],
}
# Steuerzeile im Cashflow (P07, Cash-Sicht wie S12/Dashboard): + = Geld fließt zu (Erstattung), − = Zahlung.
# Die Formeln C42:D42 liefern die Steuer-Sicht (Erstattung < 0) – das Vorzeichen wird nur per Format umgedreht.
TAX_CASH = '"−"#,##0" €";"+"#,##0" €";"–"'
TAX_LABEL = "± Steuerwirkung (Cash-Sicht)"
LOWER = {
    "B": [
        (36, "row", "Nettokaltmiete Ist", "eur"),
        (37, "row", "– Bewirtschaftungskosten (inkl. Rücklagen)", "eur"),
        (38, "sum", "= Einnahmenüberschuss (NOI)", "eur"),
        (39, "row", "– Zinsen", "eur"),
        (40, "row", "– Tilgung", "eur"),
        (41, "sum", "= Cashflow vor Steuern", "eur"),
        (42, "row", TAX_LABEL, TAX_CASH),
        (43, "result", "= Cashflow nach Steuern", "eur"),
        (44, "memo", "nachrichtlich: Warmmiete", "eur"),
        (45, "l2", "Ausblick und Verlauf", None),
        (46, "row", "Cashflow vor Steuern Jahr 2", "eur"),
        (47, "row", "Cashflow nach Steuern Jahr 2", "eur"),
    ],
    "F": [
        (35, "row", "Rechtsform", None),
        (36, "row", "Steuersatz angewendet (Jahr 1)", "pct2"),
        (37, "row", "AfA-Methode angewendet", None),
        (38, "row", "Reguläre Gebäude-AfA Jahr 1", "eur"),
        (39, "row", f"+ Sonder-AfA §{NBSP}7b Jahr 1", "eur"),
        (40, "row", f"+ Erhöhte Absetzungen §{NBSP}7h/7i Jahr 1", "eur"),
        (41, "row", "+ AfA bewegliche Gegenstände", "eur"),
        (42, "sum", "= Abschreibungen gesamt Jahr 1", "eur"),
        (43, "row", "Erhaltungsaufwand sofort abziehbar Jahr 1", "eur"),
        (44, "row", "Aktivierte Herstellungskosten Jahr 1", "eur"),
        (45, "sum", "= Steuerliches Ergebnis Jahr 1", "eur"),
        (46, "result", "= Steuerwirkung Jahr 1", "tax_effect"),
        (47, "memo", f"Steuereffekt gewählte AfA vs. Standard (10{NBSP}J.)", "eur"),
    ],
    "J": [
        (35, "row", "Verkauf nach", "years_n"),
        (36, "row", "Verkaufspreis (Prognose bzw. Eingabe)", "eur"),
        (37, "row", "– Verkaufskosten", "eur"),
        (38, "row", "– Restschuld (Ablösung)", "eur"),
        (39, "sub", "Veräußerungsgewinn (steuerlich)", "eur"),
        (40, "sub", "Steuerpflichtig?", None),
        (41, "row", "– Steuer auf Veräußerungsgewinn", "eur"),
        (42, "sum", "= Nettoerlös nach Steuern und Ablösung", "eur"),
        (43, "row", "+ Kumulierter Cashflow n. St. bis Verkauf", "eur"),
        (44, "row", "– Eingesetztes Eigenkapital", "eur"),
        (45, "sum", "= Gesamtertrag nach Steuern", "eur"),
        (46, "row", C.KPI_LABELS.get("MULT", "Eigenkapital-Multiplikator"), "mult2"),
        (47, "result", "= Eigenkapitalrendite p. a. (IRR n. St.)", "pct1"),
    ],
}
# Anzeigeformeln (nicht referenziert): Trennzeichen „ · “ statt „|“, Rechtsform kurz (P1-19, P2-06)
DISPLAY_FORMULAS = {
    # Textverkettungen mit Zahlen: typografisches Minus über C.minus_text (Runde 4, P2-07)
    # Runde 6: kompakt für die Wertspalte neben der Verlaufsspalte („250.000 € · –“ bzw. „200.000 · 50.000 €“)
    "K15": C.minus_text('=IF(Darlehen_II=0,FIXED(Darlehen_I,0)&" € · –",'
                        'FIXED(Darlehen_I,0)&" · "&FIXED(Darlehen_II,0)&" €")'),
    "G26": C.minus_text('=FIXED(Mietsteigerung*100,1)&" · "&FIXED(Kostensteigerung*100,1)&" · "'
                        '&FIXED(Wertsteigerung*100,1)&" %"'),
    "G27": C.minus_text('=FIXED(Mietausfall_Pct*100,1)&" % · "&Leerstand_Monate&IF(Leerstand_Monate=1," Monat"," Monate")'),
    "G35": "=" + RECHTSFORM_SHORT,
    # Lückenfüller der oberen Karten F/J und der Cashflow-Karte (leere Zellen der Vorlage, reine Anzeige)
    "G28": "=EKR_Tile",
    "G29": "=BE_Miete_vSt",
    "G30": "=BE_Miete_nSt",
    "K28": "=Zinsbindung_I",
    "K29": "=Anschlusszins_I",
    "K30": "=Tilgung_I",
    "C46": "=INDEX(Projektion!$D$33:$AQ$33,2)/12",
    "C47": "=INDEX(Projektion!$D$36:$AQ$36,2)/12",
}


# Kacheln (P11/P20): drei Paare „Betrag | Kennzahl“ auf dem Spaltenraster der Karten darunter. Die großen Kacheln
# (B/F/J) tragen die Formeln der Vorlage in Z. 9 (J9 = Name CF_nSt_Monat_J1), die kleinen Kacheln (C:D, G:H, K:L)
# Anzeigeformeln. Kanonischer Satz und Reihenfolge (C.KPI_ORDER): Beträge GI · EK · CF, Kennzahlen BMR · DSCR · IRR.
# Fußzeile = reiner Kontext (Ziel über C.threshold_text), den Status setzt C.tile selbst (chip bzw. inline).
# (c1, c2, kpi, value, Kontextzeile)
CF2 = 'SUBSTITUTE(FIXED(' + C.CF_YEAR2 + ',0),"-","' + C.MINUS + '")'
TILES = [
    ("B", "B", "GI", None, '="Kaufpreis "&FIXED(Kaufpreis,0)&" € · Nebenkosten "&FIXED(NK_Quote*100,1)&" %"'),
    ("C", "D", "BMR", "=Bruttomietrendite", "=" + C.threshold_text("BMR")),
    ("F", "F", "EK", None,
     '=IF(Reserve_Einmalig=0,"ohne Liquiditätsreserve","davon Liquiditätsreserve "&FIXED(Reserve_Einmalig,0)&" €")'),
    ("G", "H", "DSCR", "=DSCR_J1", "=" + C.threshold_text("DSCR")),
    ("J", "J", "CF", None, '="ab Jahr 2: "&' + CF2 + '&" € / Monat"'),
    ("K", "L", "IRR", "=IRR_Tile", "=" + C.threshold_text("IRR")),
]
AMPEL = {"G23": "BMR", "G24": "NMR", "G28": "EKR", "K24": "DSCR", "K47": "IRR"}   # Tabellenwerte 10 pt: Statusfarbe
NEGATIVE_RED = ("K43", "C46:C47")      # Kumul- und Cashflow-Werte außerhalb der Summenzeilen (P15)

HINT_ROWS = range(50, 56)
CHART_TOP, CHART_BOTTOM = 58, 73          # Diagramme füllen B58:D73 | F58:H73 | J58:L73
CHART_ROW = 16                            # Rasterhöhe der Diagrammzeilen (core.ROW_RASTER)
# Soll-Anker (0-basiert: from col/row → to col/row, Offsets 0) in Diagramm-Reihenfolge.
ANCHORS = {SHEET: [(1, CHART_TOP - 1, 4, CHART_BOTTOM), (5, CHART_TOP - 1, 8, CHART_BOTTOM),
                   (9, CHART_TOP - 1, 12, CHART_BOTTOM)]}
NAV_ROW = 75                              # Zurück / Weiter (Reihenfolge der Reiterleiste: Dashboard → Cockpit → Diagramme)
FOOTER_ROW = 77
ROW_AUSBLICK = 24
GAP_SECTION = C.H_ROW                     # Weißraum zwischen den großen Abschnitten (Karten, Hinweise, Diagramme)

# --------------------------------------------------------------------------------------------- Runde 6 „Wow-Paket“
# Hero: Z. 5–6 nachtblaues Band B:L (wie Deckblatt/Dashboard-Hero) mit Titel links, feiner Gold-Stadtsilhouette in
# der Mitte und dem Gesamturteil rechts (Label · „n von 6 erfüllt“ · Urteil 20 pt weiß, Statuskante links in
# leuchtender Statusfarbe); Z. 7 weiß: Objektzeile und „Erstellt für …“ – darunter die Kacheln wie bisher.
H_HERO = {4: 12, 5: 30, 6: 52, 7: 24}
HERO_PAD = 2                              # Einzug der Hero-Texte (≈ 18 px Innenabstand)
HERO_ICON, HERO_ICON_PX, HERO_ICON_DX = "rechner", 30, 16
HERO_TITLE_INDENT = 7                     # Titel/Eyebrow hinter dem Icon (16 + 30 + ≈ 17 px Luft)
# Statuskante des Urteils auf Navy: leuchtende Semantikfarben (die Text-Statusfarben sind auf Navy zu dunkel)
HERO_STATUS = {"red": C.C_NEG, "amber": C.C_ORANGE, "green": C.C_POS}
# Verlaufsspalte (Sparklines): die Wertspalten der Karten F und J werden 118 | 64 px statt 91 | 91 px aufgeteilt –
# das Paar bleibt 182 px breit (Kacheln G:H/K:L und Diagrammanker unverändert). Werte stehen in G/K, H/L trägt die
# Sparklines (nur Anzeige, Daten aus Projektion, Jahre 1–30 wie die Diagramme darunter).
W_VALUE, W_LANE = 16.8, 9.1               # 118 px · 64 px
LANE_CARDS = {("F", "upper"), ("J", "upper"), ("J", "lower")}
TREND_YEARS = 30
PJ = "Projektion"
TRENDS = [   # (Zielzelle, Projektion-Zeile, Sparkline-Vorlage, Zusatzargumente) – Farben nach CHART_SEMANTIC
    ("H16", 13, "rent", {"min_zero": True}),                                           # Jahresnettokaltmiete → Miete Blau
    ("H18", 26, "costs", {"min_zero": True}),                                          # Bewirtschaftung → Orange
    ("H20", 29, "trend", {"color": "C_BLUE", "markers": True}),        # Einnahmenüberschuss (NOI) → Blau
    ("H28", 45, "value", {}),                                          # Eigenkapitalrendite je Jahr → Aqua
    ("L14", 42, "debt", {}),                                           # Darlehen → Restschuld Jahresende (grau)
    ("L23", 32, "costs", {"min_zero": True}),                                          # Kapitaldienst pro Jahr → Orange
    ("L36", 41, "value", {}),                                          # Verkaufspreis → Immobilienwert (Aqua)
    ("L38", 42, "debt", {}),                                           # Restschuld (Ablösung)
    ("L43", 38, "cashflow", {"color": "C_BLUE"}),                      # kumulierter Cashflow n. St. (Blau/Rot)
    ("D46", 33, "cashflow", {}),                                       # Ausblick: Cashflow v. St. (Aqua/Rot)
    ("D47", 36, "cashflow", {}),                                       # Ausblick: Cashflow n. St.
]
LANE_CAPTION = f"Verlauf Jahre 1–{TREND_YEARS}"
LANE_CAPTIONS = ("H31", "L31", "L48")      # unter der Verlaufsspalte, direkt unter der Kartenunterkante
# Icons der Abschnittsköpfe (ein Motiv je Block, Nachtblau 20 px wie auf den Schrittseiten, steps.ICON_*)
BLOCK_ICON = {(13, "B"): "haus", (13, "F"): "prozent", (13, "J"): "bank", (32, "B"): "muenzen",
              (32, "F"): "paragraf", (32, "J"): "schluessel", (49, "B"): "schild", (57, "B"): "diagramm"}
ICON_PX, ICON_DX, ICON_INDENT = 20, 10, 4


# --------------------------------------------------------------------------------------------- Helfer
def _nf(key):
    return NUMFMT.get(key, key) if key else None


def _hide_cols(ws, first, last):
    """Spalten ausblenden, ohne überlappende <col>-Einträge zu erzeugen (Gruppen wie O:X bleiben eine Gruppe)."""
    lo, hi = col(first), col(last)
    covered = set()
    for key, dim in list(ws.column_dimensions.items()):
        mn, mx = dim.min or col(key), dim.max or col(key)
        if mn >= lo and mx <= hi:
            dim.hidden = True
            covered.update(range(mn, mx + 1))
    for c in range(lo, hi + 1):
        if c not in covered:
            d = ws.column_dimensions[get_column_letter(c)]
            d.min = d.max = c
            d.hidden = True


def _clear_cf(ws, predicate):
    """Bedingte Formate entfernen, deren Bereich predicate(sqref-String) erfüllt."""
    new = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        if predicate(str(cf.sqref)):
            continue
        for rule in cf.rules:
            new.add(str(cf.sqref), rule)
    ws.conditional_formatting = new


def _in_region(sqref):
    for part in sqref.split():
        c1, r1, c2, r2 = range_boundaries(part)
        if r2 >= 5 and r1 <= 80 and c1 <= col("X") and c2 >= col("B"):
            return True
    return False


def _reset(ws, r1, r2, c1=FIRST, c2="M"):
    """Neutraler Grundzustand für den Seitenbereich (entfernt Stilreste der Vorlage wie Inter/F4F0E8)."""
    for c in iter_cells(ws, c1, r1, c2, r2):
        c.font = font(T_BODY, False, INK)
        c.fill = NOFILL
        c.border = Border()
        c.alignment = align("general", "center")
        if c.value is None:
            c.number_format = "General"


def _unmerge_rows(ws, r1, r2, lo="B", hi="L"):
    """Verbünde lösen, die im Seitenbereich B:L beginnen (die ausgeblendeten Kacheln N:X bleiben unberührt)."""
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= r1 and mr.max_row <= r2 and col(lo) <= mr.min_col <= col(hi):
            ws.unmerge_cells(str(mr))


def _link(cell, text, target_sheet, target_cell=None, tooltip=None, size=C.T_LABEL, h="right", formula=False):
    """Zell-Link im Stil der Abschnitts-Meta (8,5 pt 1D4F8A, rechtsbündig) – Ziel über core.link_loc."""
    if formula:
        cell.value = text
    else:
        set_text(cell, text)
    cell.font = font(size, False, BLUE)
    cell.alignment = align(h, "center", 1)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=C.link_loc(target_sheet, target_cell), tooltip=tooltip,
                               display=None if formula else text)


def _up_link(cell):
    _link(cell, "↑ Übersicht", SHEET, C.link_target(SHEET), tooltip="Zum Seitenanfang (Kennzahlen)")


def _verdict_ref():
    """Zelle der Gesamtbewertung auf dem Dashboard (dashboard.py: Label in Z. R_VERDICT, Status darunter in B)."""
    try:
        import dashboard
        r = int(getattr(dashboard, "R_VERDICT"))
        return f"Dashboard!$B${r + 1}"
    except Exception:
        return "Dashboard!$B$12"


def _dashboard_attr(name, default):
    try:
        import dashboard
        return int(getattr(dashboard, name))
    except Exception:
        return default


def _check_range():
    """Statusspalte des Kennzahlen-Checks auf dem Dashboard (6 Zeilen, Wörter erfüllt/prüfen/kritisch)."""
    r = _dashboard_attr("R_CHECK", 20)
    return f"Dashboard!$H${r}:$H${r + 5}"


def _icon(ws, cell, name, color=NAVY, px=ICON_PX, dx=ICON_DX, dy=0, valign="middle"):
    """Icon setzen (icons.place_icon, OneCellAnchor) – Spaltenbreiten/Zeilenhöhen müssen vorher stehen.
    Fehler → Blatt ohne Icon (nie ohne Blatt)."""
    if ICN is None or not name:
        return None
    try:
        return ICN.place_icon(ws, cell, name, color, px=px, dx=dx, dy=dy, valign=valign)
    except Exception as exc:  # noqa: BLE001
        print(f"  cockpit: Icon „{name}“ in {cell} übersprungen ({exc})")
        return None


def grid(ws):
    """Runde 6: Wertspalten G/K 118 px, Verlaufsspalten H/L 64 px (Summe je Paar unverändert 182 px)."""
    for c, w in (("G", W_VALUE), ("H", W_LANE), ("K", W_VALUE), ("L", W_LANE)):
        d = ws.column_dimensions[c]
        d.min = d.max = col(c)
        d.width = w
        d.hidden = False


# --------------------------------------------------------------------------------------------- Seitenkopf
def header(ws):
    """Runde 6 – Hero im Stil des Deckblatts/Dashboard-Heros: Z. 5–6 nachtblaues Band B:L mit feiner Goldlinie unten.
      links   Eyebrow „COCKPIT“ (helles Gold 8,5 pt) · Titel „Cockpit“ 30 pt weiß · Gold-Icon „rechner“
      Mitte   feine goldene Stadtsilhouette (icons.cover_art_path, transparent, zurückhaltend)
      rechts  „GESAMTBEWERTUNG“ · „n von 6 Kennzahlen erfüllt“ · Urteil 20 pt weiß; linke 3-px-Kante in leuchtender
              Statusfarbe (bedingte Formatierung, EINE Regel je Status); Link auf die Prüfhinweise.
    Z. 7 (weiß): Objektzeile links, „Erstellt für …“ rechts. Die Kacheln (Z. 8–10) bleiben, wo sie sind."""
    _unmerge_rows(ws, 4, 7)
    for c in iter_cells(ws, "C", 4, "L", 7):
        c.value = None
        c.hyperlink = None
    date = 'TEXT(DAY(Kaufdatum),"00")&"."&TEXT(MONTH(Kaufdatum),"00")&"."&YEAR(Kaufdatum)'
    subtitle = (f'=Obj_Name&"  ·  "&Obj_Adresse&"  ·  Kauf am "&{date}&"  ·  Haltedauer "&Haltedauer&" Jahre  ·  "'
                f'&{RECHTSFORM_SHORT}')
    C.page_header(ws, FIRST, LAST, "Cockpit", "Cockpit", subtitle=subtitle)
    for r, h in H_HERO.items():
        set_height(ws, r, h)
    gold_rule = side("medium", C.GOLD)
    for c in iter_cells(ws, "B", 5, "L", 6):
        c.fill = fill(NAVY)
        c.border = Border(bottom=gold_rule if c.row == 6 else None)
    eyebrow, title = ws["B5"], ws["B6"]
    eyebrow.value = '="COCKPIT"'     # Anzeigeformel: global_rules.breadcrumb_links färbt nur statische Krumen (BLUE)
    eyebrow.font = font(C.T_LABEL, True, C.GOLD_LINE)
    eyebrow.alignment = align("left", "bottom", HERO_TITLE_INDENT)
    title.font = font(C.T_HERO, True, WHITE, C.DISPLAY)
    title.alignment = align("left", "center", HERO_TITLE_INDENT)
    # Objektzeile und „Erstellt für …“ unter dem Band (weiß)
    safe_merge(ws, "B", 7, "I", 7)
    sub = ws["B7"]
    sub.font = font(T_BODY, False, MUTED)
    sub.alignment = align("left", "center")
    safe_merge(ws, "J", 7, "L", 7)
    meta = ws["J7"]
    meta.value = '=IFERROR(IF(Erstellt_fuer="","","Erstellt für "&Erstellt_fuer),"")'
    meta.font = font(T_SMALL, False, MUTED)
    meta.alignment = align("right", "center")
    # rechts: Gesamtbewertung (Label · Zähler · Urteil)
    cap = ws["J5"]
    set_text(cap, "GESAMTBEWERTUNG")
    cap.font = font(C.T_LABEL, True, C.SKY)
    cap.alignment = align("left", "bottom", HERO_PAD)
    safe_merge(ws, "K", 5, "L", 5)
    cnt = ws["K5"]
    cnt.value = f'=IFERROR(COUNTIF({_check_range()},"erfüllt")&" von 6 Kennzahlen erfüllt","")'
    cnt.font = font(C.T_LABEL, False, C.SKY)
    cnt.alignment = align("right", "bottom", HERO_PAD)
    safe_merge(ws, "J", 6, "L", 6)
    verdict = ws["J6"]
    verdict.value = f"={_verdict_ref()}"
    verdict.number_format = "General"
    verdict.font = font(C.T_KPI, True, WHITE, C.DISPLAY)
    verdict.alignment = align("left", "center", HERO_PAD)
    verdict.hyperlink = Hyperlink(ref="J6", location=C.link_loc(SHEET, C.link_row(SHEET, 49)),
                                  tooltip="Zu den Prüfhinweisen und der steuerlichen Einordnung")
    v = "$J$6"
    conds = [(f'ISNUMBER(SEARCH("kritisch",{v}))', "red"), (f'ISNUMBER(SEARCH("Prüfpunkten",{v}))', "amber"),
             (f'ISNUMBER(SEARCH("Solide",{v}))', "green")]
    # statische Kante (vor der Neuberechnung / ohne Treffer): Gold; je Status EINE Regel mit Kante in Statusfarbe
    for r in (5, 6):
        c = ws[f"J{r}"]
        c.border = Border(left=side("thick", C.GOLD), bottom=gold_rule if r == 6 else None)
    for c in iter_cells(ws, "K", 6, "L", 6):
        c.border = Border(bottom=gold_rule)
    for cond, lvl in conds:      # je Zelle EINE Regel (LibreOffice wertet nur die erste zutreffende aus)
        C.cf_rule(ws, "J5", cond, border=Border(left=side("thick", HERO_STATUS[lvl])))
        C.cf_rule(ws, "J6", cond, border=Border(left=side("thick", HERO_STATUS[lvl]), bottom=gold_rule))
    # Bildebene: Gold-Icon vor dem Titel, Stadtsilhouette zwischen Titel und Urteil
    _icon(ws, "B6", HERO_ICON, C.GOLD, px=HERO_ICON_PX, dx=HERO_ICON_DX)
    if ICN is not None:
        try:
            w = C.span_px(ws, "C", "I")
            h = round((H_HERO[5] + H_HERO[6]) * 4 / 3)
            ICN.place_image(ws, "C5", ICN.cover_art_path(w, h, background=False, focus="center", intensity=0.75),
                            w, h)
        except Exception as exc:  # noqa: BLE001
            print(f"  cockpit: Hero-Illustration übersprungen ({exc})")


# --------------------------------------------------------------------------------------------- Kacheln
def tiles(ws):
    """Sechs Kacheln in einer Reihe (P1-11): Label 18 · Wert 30 · Kontext 16 pt, Werte neutral (P1-10),
    Status als Chip (Renditen) bzw. als Kontextzeile in Statusfarbe (Cashflow). Paare „Betrag | Rendite“ sitzen
    exakt auf den Spalten der Karten darunter; zwischen den Paaren liegt die Rinnenspalte E/I."""
    _unmerge_rows(ws, 8, 12)
    keep = {"B9", "F9", "J9"}                        # Formeln der Vorlage (J9 = Name CF_nSt_Monat_J1)
    for c in iter_cells(ws, "B", 8, "L", 12):
        if c.coordinate not in keep:
            c.value = None
            c.hyperlink = None
    for c1, c2, key, value, sub in TILES:
        label = C.kpi_label(key, caps=True, formula=True) if key == "IRR" else None
        C.tile(ws, c1, c2, 8, 9, 10, kpi=key, label=label, value=value, sub=sub, value_ref=f"${c1}$9")
    set_height(ws, 11, 6)       # Fuge Kacheln → Karten: 2 × 6 pt = eine Leerzeile H_GAP
    set_height(ws, 12, 6)


# --------------------------------------------------------------------------------------------- Karten
def _label(cell, text, kind):
    if text is not None:
        if text.startswith("=IF("):          # dynamische Beschriftung (Anzeigeformel)
            cell.value = text
        else:                                # statischer Text, auch „= Summe …“
            set_text(cell, text)
    if kind == "sub":
        cell.font = font(T_SMALL, False, MUTED)
        cell.alignment = align("left", "center", 2)
    else:
        cell.font = font(T_BODY, False, INK)
        cell.alignment = align("left", "center", 1)


def _value(ws, row, vcol, ecol, fmt, kind, merge=True):
    if merge:
        safe_merge(ws, vcol, row, ecol, row)
    v = ws.cell(row, col(vcol))
    v.font = font(T_SMALL, False, MUTED) if kind == "sub" else font(T_BODY, False, INK)
    v.alignment = align("right", "center", 1)
    if fmt:
        v.number_format = _nf(fmt)
    return v


def _block(ws, rows, lab, val, edge, merge_values=True):
    kinds = {}
    for r, kind, text, fmt in rows:
        kinds[r] = kind
        cell = ws[f"{lab}{r}"]
        if kind == "l2":
            C.section(ws, r, lab, edge, text, level=2, variant="line", height=C.H_ROW)
            continue
        _label(cell, text, kind)
        v = _value(ws, r, val, edge, fmt, kind, merge=merge_values and val != edge)
        if kind in ("sum", "result"):
            # Negativ-Rot nur für Beträge; Quoten (IRR) tragen ihren Status als Schriftfarbe (AMPEL)
            C.sum_row(ws, r, lab, edge, "sub" if kind == "sum" else "result",
                      neg=False if fmt in ("pct1", "pct2") else None)
            cell.alignment = align("left", "center", 1)
            for c in iter_cells(ws, val, r, edge, r):
                c.alignment = align("right", "center", 1)
        elif kind == "memo":
            C.memo(ws, r, lab, edge)
            for c in iter_cells(ws, val, r, edge, r):
                c.alignment = align("right", "center", 1)
        else:
            C.hairline(ws, r, lab, edge)
    # Die Linie über einer Summe gehört der Summe: keine Haarlinie darunter doppeln
    for r, kind in kinds.items():
        if kinds.get(r + 1) in ("sum", "result", "l2"):
            for c in iter_cells(ws, lab, r, edge, r):
                c.border = Border(top=c.border.top)
    return kinds


def _card_bottom(ws, row):
    """Gemeinsame Kartenunterkante; die Doppellinie eines Blockergebnisses bleibt erhalten."""
    for lab, _, edge in BLOCKS:
        for c in iter_cells(ws, lab, row, edge, row):
            b = c.border
            if b.bottom is not None and b.bottom.style == "double":
                continue
            c.border = Border(top=b.top, bottom=side("thin", LINE2))


def _block_head(ws, row, lab, edge, title, link):
    """Blockkopf Ebene 1 mit Drill-down-Link rechts (P3-02). Runde 6: Icon links vor dem Titel (Nachtblau 20 px),
    der Link sitzt im Verbund der beiden Wertspalten (die rechte Spalte ist bei Verlaufskarten nur 64 px breit)."""
    C.section(ws, row, lab, edge, title, level=1)
    text, target, tip = link
    val = get_column_letter(col(lab) + 1)
    safe_merge(ws, val, row, edge, row)
    _link(ws[f"{val}{row}"], text, target, tooltip=tip)
    name = BLOCK_ICON.get((row, lab))
    if name and _icon(ws, f"{lab}{row}", name) is not None:
        ws[f"{lab}{row}"].alignment = align("left", "center", ICON_INDENT)


def blocks(ws):
    """Obere und untere Kartenreihe (P1-12 Köpfe, P1-14 Summenstufen, P3-02 Unterzeilen/Abstände/Drill-down)."""
    _unmerge_rows(ws, 13, 47)
    for c in iter_cells(ws, "B", 13, "L", 13):
        c.value = None
    heads_up = (("B", "D", "Objekt & Investition", ("Eingaben ›", "Eingaben", "Alle Annahmen auf einen Blick")),
                ("F", "H", "Miete, Bewirtschaftung & Rendite",
                 ("Projektion ›", "Projektion", "Miete, Kosten und Cashflow über die Jahre")),
                ("J", "L", "Finanzierung", ("Finanzierung ›", "Finanzierung", "Tilgungsplan beider Darlehen")))
    for lab, edge, title, link in heads_up:
        _block_head(ws, 13, lab, edge, title, link)
    for r in range(14, 31):
        set_height(ws, r, C.H_ROW)
    for lab, val, edge in BLOCKS:
        _block(ws, UPPER[lab], lab, val, edge, merge_values=(lab, "upper") not in LANE_CARDS)
    _card_bottom(ws, 30)
    set_height(ws, 31, GAP_SECTION)

    # untere Reihe – Kopf direkt über der ersten Zeile wie oben (Z. 33/34 ausgeblendet, P3-02)
    for c in iter_cells(ws, "B", 32, "L", 34):
        if not is_formula(c.value):
            c.value = None
    heads_low = (("B", "D", "Cashflow Jahr 1", ("Projektion ›", "Projektion", "Cashflow über die Haltedauer")),
                 ("F", "H", "Steuern & AfA Jahr 1", ("Steuern ›", "Steuern", "Steuerliche Ermittlung je Jahr")),
                 ("J", "L", "Exit-Szenario nach Haltedauer",
                  ("Sensitivität ›", "Sensitivität", "Was-wäre-wenn: Verkaufspreis, Zins, Miete")))
    for lab, edge, title, link in heads_low:
        _block_head(ws, 32, lab, edge, title, link)
    for r in (33, 34):
        ws.row_dimensions[r].hidden = True
        set_height(ws, r, 3)
    for r in range(35, 48):
        set_height(ws, r, C.H_ROW)
    C.section(ws, 35, "B", "D", None, level=2, labels={"C": ("Pro Monat", "right"), "D": ("Pro Jahr", "right")},
              height=C.H_ROW)
    ws["B35"].value = None
    _block(ws, LOWER["B"], "B", "C", "D", merge_values=False)
    for r, kind, _, fmt in LOWER["B"]:          # zweite Wertspalte (pro Jahr) der Ertragsrechnung
        d = ws[f"D{r}"]
        if fmt:
            d.number_format = _nf(fmt)
        d.alignment = align("right", "center", 1)
        if kind == "row":
            d.font = font(T_BODY, False, INK)
    _block(ws, LOWER["F"], "F", "G", "H", merge_values=("F", "lower") not in LANE_CARDS)
    _block(ws, LOWER["J"], "J", "K", "L", merge_values=("J", "lower") not in LANE_CARDS)
    _card_bottom(ws, 47)
    set_height(ws, 48, GAP_SECTION)
    # P3-02: „AUSBLICK FOLGEJAHR“ folgt direkt auf die Memo-Zeile – etwas Luft darüber (Titel unten ausgerichtet);
    # die Summenzeilen F45/J45 bleiben vertikal zentriert
    set_height(ws, 45, ROW_AUSBLICK)
    ws["B45"].alignment = align("left", "bottom", 1)
    # Runde 6: Ausblick mit Verlaufsspalte – C = Jahr 2 pro Monat, D = Sparkline Cashflow Jahre 1–30
    for ref, text in (("C45", "Jahr 2 / Monat"), ("D45", f"Jahre 1–{TREND_YEARS}")):
        c = ws[ref]
        set_text(c, text.upper() if not text.startswith("Jahr") else text)
        c.font = font(C.T_LABEL, True, BLUE)
        c.alignment = align("right", "bottom", 1)

    for ref, formula in DISPLAY_FORMULAS.items():
        ws[ref].value = formula


def trends(ws):
    """Runde 6: echte Excel-Sparklines (sparklines.py) in den Verlaufsspalten H/L und im Ausblick (D46:D47).
    Zielzellen bleiben leer; Daten Projektion Jahre 1–30. Unter jeder Verlaufsspalte eine Achsnotiz (8 pt)."""
    last = get_column_letter(col("D") + TREND_YEARS - 1)
    for ref, prow, style, extra in TRENDS:
        cell = ws[ref]
        cell.value = None
        cell.hyperlink = None
        if SP is None:
            continue
        try:
            SP.register(ws, ref, f"{PJ}!D{prow}:{last}{prow}", style=style, **extra)
        except Exception as exc:  # noqa: BLE001
            print(f"  cockpit: Sparkline {ref} übersprungen ({exc})")
    if SP is None:
        return
    for ref in LANE_CAPTIONS:
        c = ws[ref]
        set_text(c, LANE_CAPTION)
        c.font = font(T_MICRO, False, MUTED)
        c.alignment = align("right", "top", 1)


def semantics(ws):
    """Status nur semantisch (P1-10): Zielkennzahlen in den Karten per Schriftfarbe, negative Ergebnisse rot."""
    for ref, key in AMPEL.items():
        C.add_ampel(ws, ref, key)
    for ref in NEGATIVE_RED:
        C.neg_red(ws, ref)


# --------------------------------------------------------------------------------------------- Prüfhinweise
# Anzeigeformeln B50:B55 (nicht referenziert): farbige Emoji der Hinweistexte → monochrome Zeichen wie auf dem
# Dashboard (P2-14, dashboard.WARN_SYM/INFO_SYM): „⚠“ → „▲“, „ℹ“ → „•“ (ⓘ fehlt in Calibri). Typografisches Minus vor
# Beträgen („(−18.397 €)“, „ −“), Bindestriche in Wörtern bleiben. DSCR-Schwelle wie überall „Ziel ≥ 1,20×“ (P21).
WARN_SYM, INFO_SYM = "▲", "•"
_HINT_PICK = 'INDEX($AB$50:$AB$74,SMALL($AC$50:$AC$74,{k}))'
_HINT_SRC = ('SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({pick},"(-","(' + C.MINUS + '")," -"," ' + C.MINUS + '"),'
             '" – Banken erwarten meist ≥ 1,1 bis 1,2","× – "&' + C.threshold_text("DSCR") + ')')
_HINT_TEXT = ('IF(LEFT({src},1)="⚠","' + WARN_SYM + '  "&TRIM(MID({src},2,999)),'
              'IF(LEFT({src},1)="ℹ","' + INFO_SYM + '  "&TRIM(MID({src},2,999)),{src}))')
_N_HINT = "COUNT($AC$50:$AC$74)"


def _hint_text(k):
    src = _HINT_SRC.replace("{pick}", _HINT_PICK.format(k=k))
    return "IFERROR(" + _HINT_TEXT.replace("{src}", src) + ',"")'


def hint_formula(k, n):
    """Hinweisplatz k (1…n): erster Platz meldet „keine Hinweise“, letzter Platz wird bei Überlauf zur Sammelzeile."""
    if k == 1:
        return (f'=IF({_N_HINT}=0,"Keine Prüfhinweise – alle Plausibilitätsprüfungen ohne Befund.",'
                f'{_hint_text(k)})')
    if k == n:
        return (f'=IF({_N_HINT}>{n},"+ "&({_N_HINT}-{n - 1})&" weitere Hinweise – Eingaben im Leitfaden '
                f'(S01–S12) prüfen",{_hint_text(k)})')
    return "=" + _hint_text(k)


def hints(ws):
    """Prüfhinweise B49:L55 (P2-14, P32, P41, P3-02): Zähler rechts im Kopf („1 Warnung“ rot fett · „3 Hinweise“ grau,
    direkt vor „↑ Übersicht“). Die sechs Plätze bilden EINE geschlossene Box (TINT_XL, weiße Fugen, Abschlusslinie
    unter Z. 55) – freie Plätze lesen sich als Teil der Box, nicht als Loch. Gefüllte Zeilen erhalten per bedingter
    Formatierung die linke 3-px-Kante (Warnung rot mit roter Schrift, Information Akzent)."""
    _unmerge_rows(ws, 49, 57)
    for c in iter_cells(ws, "C", 49, "L", 49):
        c.value = None
    C.section(ws, 49, "B", "L", "Prüfhinweise & steuerliche Einordnung", level=1)
    if _icon(ws, "B49", BLOCK_ICON[(49, "B")]) is not None:
        ws["B49"].alignment = align("left", "center", ICON_INDENT)
    cnt_w = 'COUNTIF($AB$50:$AB$74,"⚠*")'
    cnt_i = 'COUNTIF($AB$50:$AB$74,"ℹ*")'
    # Runde 6: Zähler in EINER Zelle J49 (rechtsbündig, grau; mit Warnung rot fett per Regel) – der Rücksprung
    # „↑ Übersicht“ braucht den Verbund K:L, weil L jetzt die schmale Verlaufsspalte ist
    j49 = ws["J49"]
    j49.value = (f'=IF({_N_HINT}=0,"keine Hinweise",IF({cnt_w}=0,"",{cnt_w}&IF({cnt_w}=1," Warnung"," Warnungen"))'
                 f'&IF(AND({cnt_w}>0,{cnt_i}>0),"  ·  ","")'
                 f'&IF({cnt_i}=0,"",{cnt_i}&IF({cnt_i}=1," Hinweis"," Hinweise")))')
    j49.font = font(C.T_LABEL, False, MUTED)
    j49.alignment = align("right", "center", 1)
    C.cf_rule(ws, "J49", f"{cnt_w}>0", font_=Font(color=RED, bold=True))
    safe_merge(ws, "K", 49, "L", 49)
    _up_link(ws["K49"])
    top, bot = HINT_ROWS.start, HINT_ROWS.stop - 1
    white, close = side("thin", WHITE), side("thin", C.LINE_SUB)
    for k, r in enumerate(HINT_ROWS, start=1):
        ws[f"B{r}"].value = hint_formula(k, len(HINT_ROWS))
        safe_merge(ws, "B", r, "L", r)
        c = ws[f"B{r}"]
        c.font = font(T_BODY, False, INK2)
        c.alignment = align("left", "center", 1)
        for cc in iter_cells(ws, "B", r, "L", r):      # statische Box: Fläche + Fuge bzw. Abschlusslinie,
            cc.border = Border(left=side("thick", LINE2) if cc.column == col("B") else None,    # ruhige Kante (warm)
                               bottom=close if r == bot else white)
            cc.fill = fill(TINT_XL)
        set_height(ws, r, C.H_ROW)
    warn, info = f'LEFT($B{top},1)="{WARN_SYM}"', f'$B{top}<>""'
    last = f'ROW($B{top})={bot}'
    # Ankerzelle B (gilt für den ganzen Verbund): linke Kante in Statusfarbe; Unterkante wie die Box
    anchors = f"B{top}:B{bot}"
    for cond, fg, edge in ((warn, RED, RED), (info, INK2, ACCENT)):
        C.cf_rule(ws, anchors, f"AND({cond},{last})", font_=Font(color=fg), fill_=fill(TINT_XL),
                  border=Border(left=side("thick", edge), bottom=close))
        C.cf_rule(ws, anchors, cond, font_=Font(color=fg), fill_=fill(TINT_XL),
                  border=Border(left=side("thick", edge), bottom=white))
    ws["B56"].value = None
    ws["B56"].hyperlink = None
    set_height(ws, 56, GAP_SECTION)  # P3-02: Blockabstand vor „Diagramme“ = Standardabstand (Z. 31/48)


def text_formats(ws):
    """P16: Formelzellen mit Textformat „@“ auf Standard (sonst zeigt Excel nach F2 + Enter die Formel als Text)."""
    for row in ws.iter_rows():
        for c in row:
            if c.number_format == "@" and is_formula(c.value):
                c.number_format = "General"


# --------------------------------------------------------------------------------------------- Diagramme
def charts(ws):
    """Diagrammzeile: drei Diagramme exakt auf B:D | F:H | J:L, das vierte (AfA) nur noch auf „Diagramme“."""
    _unmerge_rows(ws, 57, 57)
    for c in iter_cells(ws, "B", 57, "L", 57):
        c.value = None
    C.section(ws, 57, "B", "L", "Diagramme", level=1)
    if _icon(ws, "B57", BLOCK_ICON[(57, "B")]) is not None:
        ws["B57"].alignment = align("left", "center", ICON_INDENT)
    _link(ws["J57"], "Alle Diagramme ›", "Diagramme", tooltip="Alle Auswertungen auf dem Blatt Diagramme")
    safe_merge(ws, "K", 57, "L", 57)
    _up_link(ws["K57"])
    keep = sorted((ch for ch in ws._charts if ch.anchor._from.col < col("N") - 1),
                  key=lambda ch: ch.anchor._from.col)
    ws._charts = [ch for ch in ws._charts if ch in keep]
    for ch, (c1, r1, c2, r2) in zip(keep, ANCHORS[SHEET]):
        a = ch.anchor
        a._from.col, a._from.colOff, a._from.row, a._from.rowOff = c1, 0, r1, 0
        a.to.col, a.to.colOff, a.to.row, a.to.rowOff = c2, 0, r2, 0
    for r in range(CHART_TOP, CHART_BOTTOM + 1):
        set_height(ws, r, CHART_ROW)


def page_footer(ws):
    """Fuß wie auf den übrigen Auswertungsblättern: Leerzeile · Zurück (sekundär, unter Karte 1) · Weiter (primär,
    unter Karte 3, bündig an der Inhaltskante) · Leerzeile · Seitenfuß."""
    _unmerge_rows(ws, CHART_BOTTOM + 1, FOOTER_ROW + 3)
    for c in iter_cells(ws, "A", CHART_BOTTOM + 1, "M", FOOTER_ROW + 3):
        if not is_formula(c.value):
            c.value = None
        c.hyperlink = None
        c.border = Border()
        c.fill = NOFILL
    set_height(ws, NAV_ROW - 1, GAP_SECTION)      # Luft zwischen Diagrammlegenden und Buttons
    C.btn_row(ws, NAV_ROW, [dict(c1="B", c2="D", text="‹  Zurück: Dashboard", target="Dashboard", kind="secondary",
                                 tooltip="Zurück: Dashboard"),
                            dict(c1="J", c2="L", text="Weiter: Diagramme  ›", target="Diagramme", kind="primary",
                                 tooltip="Weiter: Diagramme")])
    set_height(ws, NAV_ROW + 1, C.H_GAP)
    C.footer(ws, FOOTER_ROW, FIRST, LAST)


def apply(wb):
    if SHEET not in wb.sheetnames:
        return
    ws = wb[SHEET]
    _clear_cf(ws, _in_region)
    _reset(ws, 4, FOOTER_ROW + 1)
    _hide_cols(ws, "N", "X")
    _hide_cols(ws, "Y", "AA")          # leere Randspalten (Hilfsspalten AB:AC sind bereits ausgeblendet)
    grid(ws)
    header(ws)
    tiles(ws)
    blocks(ws)
    trends(ws)
    semantics(ws)
    hints(ws)
    charts(ws)
    page_footer(ws)
    text_formats(ws)
    C.cf_close(ws)
