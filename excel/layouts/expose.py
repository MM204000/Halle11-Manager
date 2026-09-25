"""Investment-Exposé (Runde 6, Agent N) – druckfertige A4-Hochformat-Seite für Investoren und Bank im Magazin-Stil.

Aufbau (ein Raster, zwei gleich breite Hälften B:D | F:H mit Mittelrinne E, Randspalten A/I):
  Kopf (core.page_header) · nachtblaues Titelband mit Gold-Linienzeichnung (icons.cover_art_path), Monogramm,
  Objektname, Adresse, Kurzbeschreibung und „Erstellt für · Stand · Kauf als“ · großer Fotoplatzhalter links neben
  sechs Kennzahl-Kacheln (core.tile, 2 × 3, Gold-Icons im Kachelkopf) · Objektdaten | Finanzierung ·
  Cashflow Jahr 1 als Herleitung (+ Sparkline 40 Jahre) | Finanzierungsstruktur (Donut mit Legende) ·
  Vermögensentwicklung (Linie) · Gesamturteil als Status-Banner · Annahmen · Fuß mit Haftung.

Nur Darstellung: ein NEUES Blatt mit reinen Anzeigeformeln auf bestehende Namen/Zellen (auf nichts davon verweist
eine Formel oder ein Name der Vorlage). Hilfszellen der Diagramme und des Urteils liegen in ausgeblendeten Spalten.
"""
from openpyxl.chart import DoughnutChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
from openpyxl.styles import Border, Font

import core as C

try:
    import icons as I
except Exception:  # pragma: no cover – Icons sind Zierde, nie Voraussetzung
    I = None
try:
    import sparklines as SP
except Exception:  # pragma: no cover
    SP = None

SHEET = "Exposé"
AFTER = "Vermögensaufstellung"

# ------------------------------------------------------------------------------------------------ Raster (px)
COL_W = {"A": 3.71, "B": 27.0, "C": 2.0, "D": 27.0, "E": 2.86, "F": 27.0, "G": 2.0, "H": 27.0, "I": 3.71}
HELP = ("K", "L", "M")                 # ausgeblendete Hilfsspalten (Donut-Daten, Urteil)
LEFT, RIGHT = ("B", "D"), ("F", "H")   # zwei Hälften à 392 px, Mittelrinne E 20 px
C1, C2 = "B", "H"                      # Inhaltskante 804 px

# ------------------------------------------------------------------------------------------------ Zeilen
R_BAND = 9          # Titelband 9 … 16
R_TILES = 18        # Kacheln 18–20 · 22–24 · 26–28, Foto B18:D28
R_CAPTION = 29
R_OBJ = 31          # Objekt | Finanzierung: Kopf 31, Zeilen 32–38
N_OBJ = 7
R_CF = 40           # Cashflow Jahr 1 | Finanzierungsstruktur: Kopf 40, Zeilen 41–48
N_CF = 8
R_WEALTH = 50       # Vermögensentwicklung: Kopf 50, Diagramm 51–58
N_WEALTH = 8
R_VERDICT = 60      # Gesamturteil 60–61
R_NOTE = 63         # Annahmen / Hinweis
R_FOOT = 65         # Fuß 65/66

DATE_TODAY = 'TEXT(DAY(TODAY()),"00")&"."&TEXT(MONTH(TODAY()),"00")&"."&YEAR(TODAY())'
DATE_BUY = 'IFERROR(TEXT(DAY(Kaufdatum),"00")&"."&TEXT(MONTH(Kaufdatum),"00")&"."&YEAR(Kaufdatum),"–")'
RF_WORD = 'CHOOSE(Rechtsform_Idx,"Privatperson","vermögensverwaltende GmbH","GmbH / Holding")'
MADE_FOR = 'IFERROR(IF(Erstellt_fuer="","","Erstellt für "&Erstellt_fuer&"  ·  "),"")'
EUR_QM0 = '#,##0" €/m²";"−"#,##0" €/m²";"–"'
Y1 = "INDEX(Projektion!$D${r}:$AQ${r},1)/12"      # Monatswert Jahr 1 aus der Projektion


# ================================================================================================= Helfer
def _row(ws, r, h):
    ws.row_dimensions[r].height = h


def _cell(ws, ref, value=None, size=C.T_BODY, bold=False, color=C.INK, h="left", v="center", indent=0, wrap=False,
          fmt=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = C.font(size, bold, color, italic=italic)
    c.alignment = C.align(h, v, indent, wrap=wrap)
    if fmt:
        c.number_format = fmt
    return c


def _fill(ws, c1, r1, c2, r2, color):
    for r in range(r1, r2 + 1):
        for c in C.iter_cells(ws, c1, r, c2, r):
            c.fill = C.fill(color)


def _rows_px(ws, r1, r2):
    return sum(round((ws.row_dimensions[r].height or 15) * 96 / 72) for r in range(r1, r2 + 1))


def _icon(ws, ref, name, color, px, **kw):
    if I is None:
        return None
    try:
        return I.place_icon(ws, ref, name, color, px=px, **kw)
    except Exception as exc:  # nie den Build brechen
        print(f"WARNUNG expose Icon {name}: {exc!r}")
        return None


def _image(ws, ref, path_fn, w, h, **kw):
    if I is None:
        return None
    try:
        return I.place_image(ws, ref, path_fn(), w, h, **kw)
    except Exception as exc:
        print(f"WARNUNG expose Bild {ref}: {exc!r}")
        return None


def _section(ws, row, c1, c2, title, icon=None, meta=None):
    """Abschnittskopf Ebene 1 (core.section) mit feinem Linien-Icon vor dem Titel (Titel eingerückt)."""
    C.section(ws, row, c1, c2, title, meta=meta)
    if icon:
        first = ws.cell(row, C.col(c1))
        first.alignment = C.align("left", "center", 4)
        _icon(ws, f"{c1}{row}", icon, C.NAVY, 18, dx=12, valign="middle")


def _anchor(chart, c1, r1, c2, r2, dx1=0, dy1=0, dx2=0, dy2=0):
    """Zwei-Zellen-Anker: from = c1/r1 (+Versatz px), to = Zelle NACH c2/r2 (−Versatz px)."""
    emu = 9525
    chart.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=C.col(c1) - 1, row=r1 - 1, colOff=int(dx1 * emu), rowOff=int(dy1 * emu)),
        to=AnchorMarker(col=C.col(c2) - 1, row=r2 - 1, colOff=int(dx2 * emu), rowOff=int(dy2 * emu)))


def _txpr(size=C.T_MICRO, color=C.MUTED, bold=False):
    cp = CharacterProperties(sz=int(size * 100), b=bold, solidFill=color)
    return RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _line(color, w_emu, dash=None):
    gp = GraphicalProperties()
    gp.line = LineProperties(solidFill=color, w=w_emu, prstDash=dash)
    return gp


def _label_value(ws, r, lab, val, fmt, side=LEFT, bold=False, lab_color=C.INK, val_color=C.INK):
    """Tabellenzeile einer Hälfte: Beschriftung über die ersten beiden Spalten (inkl. Rinne), Wert rechtsbündig."""
    a, b = side
    mid = C.L(C.col(a) + 1)
    C.safe_merge(ws, a, r, mid, r)
    C.set_text(ws[f"{a}{r}"], lab)
    _cell(ws, f"{a}{r}", None, C.T_BODY, bold, lab_color, "left", "center", 1)
    _cell(ws, f"{b}{r}", val, C.T_BODY, bold, val_color, "right", "center", 1, fmt=fmt)
    for c in C.iter_cells(ws, a, r, b, r):
        c.border = Border(bottom=C.side("hair", C.LINE))
    _row(ws, r, C.H_ROW)


# ================================================================================================= Aufbau
def _sheet(wb):
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    idx = wb.sheetnames.index(AFTER) + 1 if AFTER in wb.sheetnames else len(wb.sheetnames)
    ws = wb.create_sheet(SHEET, index=idx)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    ws.sheet_view.zoomScale = 100
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = C.SKY
    for k, w in COL_W.items():
        ws.column_dimensions[k].width = w
    ws.column_dimensions["J"].width = 3.71
    for h in HELP:
        ws.column_dimensions[h].width = 12
    ws.oddFooter.left.text = "MM Holding GmbH · Immobilien-Kalkulation · &A"
    ws.oddFooter.right.text = "Seite &P von &N"
    for part in (ws.oddFooter.left, ws.oddFooter.right):
        part.size = 8
        part.font = "Calibri,Regular"
        part.color = C.MUTED
    return ws


def _chrome(ws):
    """Kopfleiste wie auf allen Blättern (Navy Z. 1–2, Goldlinie Z. 3); navigation.py setzt die Reiter darüber und
    schneidet die Fläche auf die Reiterleiste zu."""
    try:
        from layouts import chrome
        chrome.masthead(ws)
        return
    except Exception as exc:  # Rückfall: dieselbe Fläche direkt
        print(f"WARNUNG expose chrome: {exc!r}")
    for r, h in ((1, 6), (2, 33), (3, 3)):
        _row(ws, r, h)
        for cc in range(1, 40):
            c = ws.cell(r, cc)
            c.fill = C.fill(C.GOLD if r == 3 else C.NAVY)
            c.border = Border()


def _header(ws):
    C.page_header(ws, C1, C2, ("Bank", "Exposé"), "Investment-Exposé",
                  "Objekt, Kennzahlen und Finanzierung auf einer Seite  ·  druckfertig für Investoren und Bank "
                  "(A4 hoch)", canonical=False)
    C.safe_merge(ws, "B", 6, "F", 6)
    C.safe_merge(ws, "B", 7, "H", 7)
    _row(ws, 8, C.H_GAP)


def _band(ws):
    """Titelband: Nachtblau über die volle Druckbreite (A–I), Gold-Linienzeichnung rechts, Monogramm oben rechts."""
    r0 = R_BAND
    heights = {r0: 12, r0 + 1: 16, r0 + 2: 34, r0 + 3: 18, r0 + 4: 10, r0 + 5: 45.75, r0 + 6: 18, r0 + 7: 12}
    for r, h in heights.items():
        _row(ws, r, h)
    last = r0 + 7
    _fill(ws, "A", r0, "I", last, C.NAVY)
    for r in range(r0, last + 1):
        for c in C.iter_cells(ws, "A", r, "I", r):
            c.border = Border(bottom=C.side("medium", C.GOLD)) if r == last else Border()
    # Texte links (B:F), das Motiv läuft rechts aus
    C.safe_merge(ws, "B", r0 + 1, "F", r0 + 1)
    _cell(ws, f"B{r0 + 1}", '="KAPITALANLAGE  ·  "&UPPER(Objektart)&"  ·  "&UPPER(Bundesland)', C.T_LABEL, True,
          C.SKY, "left", "bottom")
    C.safe_merge(ws, "B", r0 + 2, "F", r0 + 2)
    _cell(ws, f"B{r0 + 2}", "=Obj_Name", C.T_H1, True, C.WHITE, "left", "center")
    C.safe_merge(ws, "B", r0 + 3, "F", r0 + 3)
    _cell(ws, f"B{r0 + 3}", "=Obj_Adresse", C.T_BODY, False, C.WHITE, "left", "top")
    # kurze Goldlinie als redaktioneller Trenner (nur unter Spalte B)
    ws.cell(r0 + 4, 2).border = Border(bottom=C.side("thin", C.GOLD))
    lead = C.minus_text(
        '=Objektart&" mit "&FIXED(Wohnflaeche,0)&" m² Wohnfläche, Baujahr "&Baujahr&". Kaufpreis "&FIXED(Kaufpreis,0)'
        '&" € (Faktor "&FIXED(Kaufpreisfaktor,1)&"×), Nettokaltmiete "&FIXED(Miete_Monat,0)&" € im Monat, Kauf am "&'
        + DATE_BUY + '&", geplante Haltedauer "&Haltedauer&" Jahre."')
    C.safe_merge(ws, "B", r0 + 5, "D", r0 + 5)
    _cell(ws, f"B{r0 + 5}", lead, C.T_SMALL, False, C.SKY, "left", "center", wrap=True)
    C.safe_merge(ws, "B", r0 + 6, "F", r0 + 6)
    _cell(ws, f"B{r0 + 6}", f'=UPPER({MADE_FOR}&"Stand "&{DATE_TODAY}&"  ·  Kauf als "&{RF_WORD})',
          C.T_LABEL, True, C.WHITE, "left", "center")
    # Gold-Linienzeichnung (transparent) über dem Band, Monogramm oben rechts
    w = C.span_px(ws, "A", "I")
    h = _rows_px(ws, r0, last) - 2
    if I is not None:
        _image(ws, f"A{r0}", lambda: I.cover_art_path(w, h, background=False, focus="right", intensity=0.9), w, h)
        _image(ws, f"H{r0 + 1}", lambda: I.monogram_path(52), 52, 52, align="right", dx=-4, dy=-2)


def _photo(ws):
    """Fotoplatzhalter links (B:D) – so hoch wie die drei Kachelreihen rechts."""
    r1, r2 = R_TILES, R_TILES + 10
    C.safe_merge(ws, "B", r1, "D", r2)
    top = ws.cell(r1, 2)
    top.fill = C.fill(C.TINT_XL)
    for r in range(r1, r2 + 1):
        for c in C.iter_cells(ws, "B", r, "D", r):
            c.fill = C.fill(C.TINT_XL)
    w, h = C.span_px(ws, "B", "D"), _rows_px(ws, r1, r2)
    if I is not None:
        _image(ws, f"B{r1}", lambda: I.photo_placeholder_path(w, h), w, h)
    # Bildunterschrift / Hinweis (8 pt, im Druck dezent)
    C.safe_merge(ws, "B", R_CAPTION, "D", R_CAPTION)
    _cell(ws, f"B{R_CAPTION}", "Objektfoto: Platzhalter anklicken, Entf drücken, dann Einfügen › Bilder",
          C.T_MICRO, False, C.MUTED, "left", "center", 0, italic=True)


def _tiles(ws):
    """Sechs Kennzahlen (C.KPI_ORDER) als 2 × 3 Kacheln rechts neben dem Foto – dieselbe Anatomie wie überall."""
    subs = {
        "GI": '="Kaufpreis "&FIXED(Kaufpreis,0)&" €"',
        "EK": '="Quote "&FIXED(EK_Quote*100,1)&" %"',
        "CF": f'="Jahr 2: "&FIXED({C.CF_YEAR2},0)&" €"',
        "BMR": '="Faktor "&FIXED(Kaufpreisfaktor,1)&"×"',
        "DSCR": '="Rate "&FIXED(Kapitaldienst_J1/12,0)&" € / Monat"',
        "IRR": '="Multiple "&FIXED(EK_Multiple,2)&"×"',
    }
    values = {"GI": "=Gesamtinvestition", "EK": "=EK_Bedarf_gesamt", "CF": "=CF_nSt_Monat_J1",
              "BMR": "=Bruttomietrendite", "DSCR": "=DSCR_J1", "IRR": "=EK_IRR"}
    icon = {"GI": "haus", "EK": "muenzen", "CF": "euro", "BMR": "prozent", "DSCR": "schild", "IRR": "trend"}
    for i, key in enumerate(C.KPI_ORDER):
        r = R_TILES + (i // 2) * 4
        c = "F" if i % 2 == 0 else "H"
        label = C.kpi_label(key, caps=True, formula=True) if key == "IRR" else C.kpi_label(key, caps=True, tile=True)
        ref = values[key][1:] if C.KPI[key].get("rule") else None
        C.tile(ws, c, c, r, r + 1, r + 2, kpi=key, label=label, value=values[key], value_ref=ref, sub=subs[key],
               gap_right=False, gap=False)
        lab = ws[f"{c}{r}"]
        if C.is_formula(label):
            lab.value = label
            lab.data_type = "f"
        _icon(ws, f"{c}{r}", icon[key], C.GOLD, 14, align="right", valign="middle", dx=-8)
        if i // 2 < 2:
            _row(ws, r + 3, 10.5)       # Rinne 14 px = Mittelrinne der Kacheln (G)
    # Legende unter den Kacheln
    C.safe_merge(ws, "F", R_CAPTION, "H", R_CAPTION)
    leg = ws[f"F{R_CAPTION}"]
    leg.value = C.status_legend()
    leg.alignment = C.align("right", "center")
    _row(ws, R_CAPTION, 16)
    _row(ws, R_CAPTION + 1, C.H_GAP)


def _objekt(ws):
    r = R_OBJ
    _section(ws, r, *LEFT, "Objekt", icon="haus")
    _section(ws, r, *RIGHT, "Finanzierung", icon="bank")
    left = [
        ("Objektart", "=Objektart", "General"),
        ("Wohn-/Nutzfläche", "=Wohnflaeche", C.NUMFMT["m2"]),
        ("Baujahr", "=Baujahr", C.NUMFMT["year"]),
        ("Kaufdatum (geplant)", "=Kaufdatum", C.NUMFMT["date"]),
        ("Kaufpreis", "=Kaufpreis", C.EUR),
        ("Kaufpreis je m²", "=IF(Wohnflaeche>0,KP_Immobilie/Wohnflaeche,0)", EUR_QM0),
        ("Nettokaltmiete / Monat", "=Miete_Monat", C.EUR),
    ]
    right = [
        ("Darlehen I", "=Darlehen_I", C.EUR),
        ("Sollzins / anfängliche Tilgung", '=FIXED(Zins_I*100,2)&" % / "&FIXED(Tilgung_I*100,2)&" %"', "General"),
        ("Zinsbindung", "=Zinsbindung_I", C.NUMFMT["years_calc"]),
        ("Darlehen II", "=Darlehen_II", C.EUR),
        ("Eigenkapital inkl. Reserve", "=EK_Bedarf_gesamt", C.EUR),
        ("Beleihungsauslauf", "=Beleihung", C.PCT1),
        ("Kapitaldienst / Monat", "=Kapitaldienst_Monat_J1", C.EUR),
    ]
    for i, (lab, val, fmt) in enumerate(left):
        _label_value(ws, r + 1 + i, lab, val, fmt, LEFT)
    for i, (lab, val, fmt) in enumerate(right):
        _label_value(ws, r + 1 + i, lab, val, fmt, RIGHT)
    # Kapitaldienst als Kennzahlzeile (letzte Zeile der Finanzierung)
    C.sum_row(ws, r + N_OBJ, "F", "H", "kpi", value_from="H")
    _row(ws, r + N_OBJ + 1, C.H_GAP)


def _cashflow(ws):
    """Cashflow Jahr 1 als kleine Herleitung (€ / Monat, Werte der Projektion Jahr 1) + Sparkline über 40 Jahre."""
    r = R_CF
    _section(ws, r, *LEFT, "Cashflow Jahr 1", icon="muenzen", meta="€ / Monat")
    lines = [
        ("Nettokaltmiete (nach Ausfall)", "=" + Y1.format(r=15), C.EUR, None),
        ("–  Bewirtschaftungskosten", "=" + Y1.format(r=26), C.EUR, None),
        ("–  Zinsen", "=" + Y1.format(r=30), C.EUR, None),
        ("–  Tilgung", "=" + Y1.format(r=31), C.EUR, None),
        ("=  Cashflow vor Steuern", "=" + Y1.format(r=33), C.EUR, "sub"),
        ("±  Steuerwirkung", "=" + Y1.format(r=36) + "-" + Y1.format(r=33), C.NUMFMT["eur_signed"], None),
        ("=  Cashflow nach Steuern", "=" + Y1.format(r=36), C.EUR, "final"),
    ]
    for i, (lab, val, fmt, stage) in enumerate(lines):
        rr = r + 1 + i
        _label_value(ws, rr, lab, val, fmt, LEFT)
        if stage:
            C.sum_row(ws, rr, "B", "D", stage, value_from="D")
    # Sparkline: Cashflow nach Steuern je Jahr (Säulen, positiv Aqua / negativ Rot)
    rs = r + 1 + len(lines)
    C.safe_merge(ws, "B", rs, "C", rs)
    _cell(ws, f"B{rs}", "Verlauf n. St. · Jahr 1–40", C.T_SMALL, False, C.MUTED, "left", "center", 1)
    _row(ws, rs, C.H_ROW)
    if SP is not None:
        try:
            SP.register(ws, f"D{rs}", "Projektion!D36:AQ36", style="cashflow")
        except Exception as exc:
            print(f"WARNUNG expose Sparkline: {exc!r}")
    _structure(ws)
    _row(ws, r + N_CF + 1, C.H_GAP)


def _structure(ws):
    """Finanzierungsstruktur: flacher Donut (F) + Legende mit Beträgen und Anteilen (H), Fremdkapitalquote in der Mitte."""
    r = R_CF
    _section(ws, r, *RIGHT, "Finanzierungsstruktur", icon="prozent")
    names = ("Darlehen I", "Darlehen II", "Eigenkapital")
    refs = ("Darlehen_I", "Darlehen_II", "Eigenkapital")
    k, v = HELP[0], HELP[1]
    C.set_text(ws[f"{k}{r}"], "Donut Finanzierungsstruktur")
    for i, (n, ref) in enumerate(zip(names, refs)):
        ws[f"{k}{r + 1 + i}"] = n
        ws[f"{v}{r + 1 + i}"] = f"=MAX(0,{ref})"
    tot = f"SUM(${v}${r + 1}:${v}${r + 3})"
    colors = C.chart_palette(list(names))
    # Legende (H): je Position Name mit Farbpunkt, darunter Betrag · Anteil
    for i, (n, colr) in enumerate(zip(names, colors)):
        a, b = r + 1 + 2 * i, r + 2 + 2 * i
        ws[f"H{a}"].value = C.rich([("●  ", C.T_BODY, True, colr), (n, C.T_BODY, True, C.NAVY)])
        ws[f"H{a}"].alignment = C.align("left", "bottom", 1)
        c = ws[f"H{b}"]
        c.value = C.minus_text(f'=IF({v}{r + 1 + i}>0,FIXED({v}{r + 1 + i},0)&" €  ·  "&'
                               f'FIXED(IF({tot}>0,{v}{r + 1 + i}/{tot},0)*100,0)&" %","–")')
        _cell(ws, f"H{b}", None, C.T_SMALL, False, C.MUTED, "left", "top", 3)
    rt = r + 7
    ws[f"H{rt}"].value = C.minus_text('="Gesamt "&FIXED(' + tot + ',0)&" €"')
    _cell(ws, f"H{rt}", None, C.T_SMALL, True, C.NAVY, "left", "center", 1)
    ws[f"H{rt}"].border = Border(top=C.side("thin", C.LINE2))
    # Mitte des Donuts: Fremdkapitalquote in Zellen HINTER dem transparenten Diagramm (über E:G verbunden, damit die
    # Zahl exakt unter dem Loch sitzt; die Diagrammfläche ist ohne Füllung – finish_pro)
    C.safe_merge(ws, "E", r + 3, "G", r + 4)
    mid = ws[f"E{r + 3}"]
    mid.value = f'=IF({tot}>0,({v}{r + 1}+{v}{r + 2})/{tot},0)'
    _cell(ws, f"E{r + 3}", None, C.T_H2, True, C.NAVY, "center", "bottom", fmt=C.NUMFMT["pct0"])
    C.safe_merge(ws, "E", r + 5, "G", r + 5)
    _cell(ws, f"E{r + 5}", "Fremdkapital", C.T_MICRO, False, C.MUTED, "center", "top")
    ch = DoughnutChart(holeSize=62)
    ch.add_data(Reference(ws, min_col=C.col(v), min_row=r + 1, max_row=r + 3), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=C.col(k), min_row=r + 1, max_row=r + 3))
    s = ch.series[0]
    s.tx = SeriesLabel(v="Finanzierungsstruktur")
    s.dPt = []
    for i, colr in enumerate(colors):
        pt = DataPoint(idx=i)
        pt.graphicalProperties = GraphicalProperties(solidFill=colr)
        pt.graphicalProperties.line = LineProperties(solidFill=C.WHITE, w=19050)
        s.dPt.append(pt)
    ch.legend = None
    ch.title = None
    ch.firstSliceAng = 0
    ch.varyColors = True
    ch.plot_area.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                                           x=0.06, y=0.04, w=0.88, h=0.92))
    _anchor(ch, "F", r + 1, "G", r + N_CF + 1)
    ws.add_chart(ch)


def _wealth(ws):
    """Vermögensentwicklung über 40 Jahre (flach): Immobilienwert, Restschuld (gestrichelt), Nettovermögen."""
    r = R_WEALTH
    _section(ws, r, C1, C2, "Vermögensentwicklung", icon="trend")
    meta = ws.cell(r, C.col(C2))
    meta.value = C.minus_text('="Nettovermögen nach "&Haltedauer&" Jahren: "&FIXED(INDEX(Projektion!$D$43:$AQ$43,'
                              'Haltedauer),0)&" €  ·  40 Jahre, jeweils Jahresende"')
    meta.font = C.font(C.T_MICRO, False, C.BLUE)
    meta.alignment = C.align("right", "center", 1)
    for i in range(N_WEALTH):
        _row(ws, r + 1 + i, 20)
    wb = ws.parent
    P = wb["Projektion"]
    line = LineChart()
    spec = [(41, "Immobilienwert", C.chart_color("Immobilienwert"), 22225, None),
            (42, "Restschuld", C.chart_color("Restschuld"), 15875, "dash"),
            (43, "Nettovermögen", C.chart_color("Nettovermögen"), 28575, None)]
    for row, title, colr, w, dash in spec:
        line.add_data(Reference(P, min_col=4, max_col=43, min_row=row), from_rows=True, titles_from_data=False)
        s = line.series[-1]
        s.tx = SeriesLabel(v=title)
        s.smooth = False
        s.marker = Marker(symbol="none")
        s.graphicalProperties = _line(colr, w, dash)
    line.set_categories(Reference(P, min_col=4, max_col=43, min_row=10))
    line.title = None
    line.legend.position = "b"
    line.y_axis.delete = False
    line.y_axis.scaling.min = 0
    line.y_axis.numFmt = C.NUMFMT["teur"]
    line.y_axis.majorTickMark = "none"
    line.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=C.LINE, w=6350)))
    line.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    line.y_axis.txPr = _txpr()
    line.x_axis.delete = False
    line.x_axis.tickLblSkip = 5
    line.x_axis.tickMarkSkip = 5
    line.x_axis.majorTickMark = "none"
    line.x_axis.numFmt = "0"
    line.x_axis.txPr = _txpr()
    line.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=C.LINE2, w=9525))
    line.plot_area.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                                             x=0.07, y=0.05, w=0.91, h=0.72))
    _anchor(line, C1, r + 1, "I", r + N_WEALTH + 1, dy1=4)
    ws.add_chart(line)
    _row(ws, r + N_WEALTH + 1, C.H_GAP)


def _verdict(ws):
    """Gesamturteil als Status-Banner (core.status_banner): gleiche Logik wie das Dashboard (sechs Kennzahlen,
    Liquidität vor Rendite), eigene Hilfszellen in der ausgeblendeten Spalte – nur Anzeigeformeln."""
    r1, r2 = R_VERDICT, R_VERDICT + 1
    k, v = HELP[0], HELP[2]
    base = R_VERDICT
    C.set_text(ws[f"{k}{base - 1}"], "Status je Kennzahl (Gesamturteil)")
    checks = [("CF", "CF_nSt_Monat_J1", None, None),
              ("BMR", "Bruttomietrendite", "Ampel_BMR_gruen", "Ampel_BMR_gelb"),
              ("NMR", "Nettomietrendite", "Ampel_NMR_gruen", "Ampel_NMR_gelb"),
              ("DSCR", "DSCR_J1", "Ampel_DSCR_gruen", "Ampel_DSCR_gelb"),
              ("EKR", "EKR_Tile", "Ampel_EKR_gruen", "Ampel_EKR_gelb"),
              ("IRR", "EK_IRR", "Ampel_IRR_gruen", "Ampel_IRR_gelb")]
    rows = {}
    for i, (key, ref, g, y) in enumerate(checks):
        rr = base + i
        rows[key] = rr
        ws[f"{k}{rr}"] = key
        if key == "CF":
            ws[f"{v}{rr}"] = f'=IF({ref}<0,"kritisch",IF({C.CF_YEAR2}<0,"prüfen","erfüllt"))'
        else:
            ws[f"{v}{rr}"] = f'=IF({ref}>={g},"erfüllt",IF({ref}>={y},"prüfen","kritisch"))'
    st = f"${v}${base}:${v}${base + 5}"
    _row(ws, r1, C.H_ROW)
    _row(ws, r2, C.H_ROW3)
    b = f"$B${r2}"
    conds = [(f'ISNUMBER(SEARCH("kritisch",{b}))', "red"),
             (f'ISNUMBER(SEARCH("Prüfpunkten",{b}))', "amber"),
             (f'{b}="Solide"', "green")]
    for cond, lvl in conds:
        fg = C.STATUS_COLORS[lvl][0]
        C.cf_rule(ws, f"B{r2}", cond, font_=Font(color=fg, bold=True), border=Border(left=C.side("thick", fg)),
                  stop=False)
    C.status_banner(ws, "B", r1, "H", r2, conditions=conds)
    C.safe_merge(ws, "B", r1, "C", r1)
    lab = ws[f"B{r1}"]
    C.set_text(lab, "GESAMTURTEIL")
    lab.font = C.font(C.T_LABEL, True, C.BLUE)
    lab.alignment = C.align("left", "bottom", 1)
    pill = ws[f"D{r1}"]
    C.status_pill(ws, pill, conditions=conds, style="chip",
                  suffix=f'COUNTIF({st},"erfüllt")&" von 6 erfüllt"')
    pill.font = C.font(C.T_SMALL, True, C.MUTED)
    pill.alignment = C.align("right", "bottom", 1)
    C.safe_merge(ws, "B", r2, "D", r2)
    verdict = ws[f"B{r2}"]
    verdict.value = (f'=IF(OR({v}{rows["DSCR"]}="kritisch",{v}{rows["CF"]}="kritisch"),"Liquidität kritisch",'
                     f'IF(COUNTIF({st},"kritisch")>0,"Rendite kritisch",'
                     f'IF(COUNTIF({st},"prüfen")>0,"Solide mit Prüfpunkten","Solide")))')
    verdict.number_format = "General"
    verdict.font = C.font(C.T_KPI, True, C.NAVY)
    verdict.alignment = C.align("left", "center", 1)
    C.safe_merge(ws, "F", r1, "H", r2)
    ex = ws[f"F{r1}"]
    ex.value = C.minus_text(
        '="–  DSCR (Jahr 1) "&FIXED(DSCR_J1,2)&"×"&IF(DSCR_J1<Ampel_DSCR_gruen," statt mindestens "," bei Ziel ")'
        '&FIXED(Ampel_DSCR_gruen,2)&"×"'
        '&CHAR(10)&"–  IRR n. St. "&FIXED(EK_IRR*100,1)&" % über "&Haltedauer&" Jahre (Ziel "'
        '&FIXED(Ampel_IRR_gruen*100,1)&" %)"'
        '&CHAR(10)&"–  Cashflow n. St. "&FIXED(CF_nSt_Monat_J1,0)&" € / Monat, ab Jahr 2 "'
        f'&FIXED({C.CF_YEAR2},0)&" €"')
    ex.font = C.font(C.T_SMALL, False, C.INK2)
    ex.alignment = C.align("left", "center", 1, wrap=True)
    for r in (r1, r2):
        ws.cell(r, C.col("F")).border = Border(left=C.side("thin", C.LINE2))
    _row(ws, r2 + 1, C.H_GAP)


def _notes(ws):
    r = R_NOTE
    C.safe_merge(ws, C1, r, C2, r)
    c = ws[f"{C1}{r}"]
    c.value = C.minus_text(
        '="Annahmen: Miete +"&FIXED(Mietsteigerung*100,1)&" % p. a.  ·  Kosten +"&FIXED(Kostensteigerung*100,1)'
        '&" % p. a.  ·  Wert +"&FIXED(Wertsteigerung*100,1)&" % p. a.  ·  Alle Werte aus der Kalkulation (Stand "&'
        + DATE_TODAY + '&"). Prognosen sind Modellrechnungen und keine Zusicherung."')
    c.font = C.font(C.T_MICRO, False, C.MUTED)
    c.alignment = C.align("left", "center", 0, wrap=True)
    _row(ws, r, C.H_ROW2)
    _row(ws, r + 1, C.H_GAP)
    C.footer(ws, R_FOOT, C1, C2)


def apply(wb):
    ws = _sheet(wb)
    _chrome(ws)
    _header(ws)
    _band(ws)
    # Kachelzeilen vor dem Foto festlegen (Fotohöhe = drei Kachelreihen)
    for i in range(3):
        r = R_TILES + 4 * i
        _row(ws, r, C.H_TILE_LABEL)
        _row(ws, r + 1, C.H_TILE_VALUE)
        _row(ws, r + 2, C.H_TILE_SUB)
        if i < 2:
            _row(ws, r + 3, 10.5)
    _row(ws, R_TILES - 1, C.H_GAP)
    _tiles(ws)
    _photo(ws)
    _objekt(ws)
    _cashflow(ws)
    _wealth(ws)
    _verdict(ws)
    _notes(ws)
    C.hide_cols(ws, HELP[0], HELP[-1])
    C.cf_close(ws)
    return ws
