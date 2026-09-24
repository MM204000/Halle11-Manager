"""Bankgespräch (druckfertiges A4-Dokument), Haushaltsrechnung, Vermögensaufstellung.

Stil: ruhiges „Private Banking“ – viel Weißraum, feine Linien, eine Text- und eine Wertkante je Block.
Komponentensprache Runde 2 ausschließlich aus core.py: page_header, section (Ebene 1/2), sum_row (Summenstufen),
status_cf / neg_red / status_legend (Statusfarben-Disziplin), tile (Kachel), callout_box (Rail-Hinweis), text_link.
Nur Darstellung: Formeln der Vorlage bleiben unverändert; geändert werden ausschließlich reine Anzeigeformeln
(Kopfzeilen, Fußnoten), statische Beschriftungen, Stile, Zahlenformate, Höhen/Breiten, Gliederung und Diagrammanker.
"""
import datetime
import re

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Border, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import Outline

import core as C

BANK, HH, VA = "Bankgespräch", "Haushaltsrechnung", "Vermögensaufstellung"


# =============================================================================== lokale Helfer
def _unmerge(ws, r1, r2, c1=1, c2=200):
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row <= r2 and mr.max_row >= r1 and mr.min_col <= C.col(c2) and mr.max_col >= C.col(c1):
            ws.unmerge_cells(str(mr))


def _clear(ws, c1, r1, c2, r2, values=True):
    for c in C.iter_cells(ws, c1, r1, c2, r2):
        if values and not C.is_formula(c.value):
            c.value = None
        c.fill = C.NOFILL
        c.border = Border()
        c.hyperlink = None


def _drop_cf(ws, *refs):
    """Bedingte Formate auf genau diesen Bereichen entfernen (z. B. Vorgaben aus global_rules.early)."""
    cf = ws.conditional_formatting
    for key in list(cf._cf_rules):
        if str(key.sqref) in refs:
            del cf._cf_rules[key]


def _input_box(ws, c1, row, c2, h="left"):
    """Eingabefeld über mehrere Spalten: gelb, Rahmen nur außen, Einzug 1."""
    ln = C.side("thin", C.INPUT_LINE)
    for cc in range(C.col(c1), C.col(c2) + 1):
        cell = ws.cell(row, cc)
        C.input_style(cell, "required")
        cell.border = Border(left=ln if cc == C.col(c1) else None, right=ln if cc == C.col(c2) else None, top=ln,
                             bottom=ln)
    first = ws.cell(row, C.col(c1))
    first.font = C.font(C.T_BODY, True, C.INPUT_FG)
    first.alignment = C.align(h, "center", 1)


def _anchor(chart, c1, r1, c2, r2):
    """Diagramm exakt auf die Zellgrenzen c1/r1 … Ende c2/r2 (ohne Versatz)."""
    a = chart.anchor
    a._from.col, a._from.colOff, a._from.row, a._from.rowOff = C.col(c1) - 1, 0, r1 - 1, 0
    a.to.col, a.to.colOff, a.to.row, a.to.rowOff = C.col(c2), 0, r2, 0


def _label(cell, indent=1, size=C.T_BODY, bold=False, color=C.INK, wrap=False):
    cell.font = C.font(size, bold, color)
    cell.alignment = C.align("left", "center", indent, wrap=wrap)


def _value(cell, fmt=None, size=C.T_BODY, bold=False, color=C.INK, h="right", indent=1, wrap=False):
    cell.font = C.font(size, bold, color)
    cell.alignment = C.align(h, "center", indent, wrap=wrap)
    if fmt:
        cell.number_format = fmt


def _hair(ws, row, c1, c2):
    for c in C.iter_cells(ws, c1, row, c2, row):
        c.fill = C.NOFILL
        c.border = Border(bottom=C.side("hair", C.LINE))


def _fmt_sub(formula, old, new):
    return formula.replace(old, new) if isinstance(formula, str) else formula


def _text_formula_general(ws):
    """Formelzellen mit Textformat „@“ auf Standard (sonst zeigt Excel nach F2 + Enter die Formel als Text)."""
    for row in ws.iter_rows():
        for c in row:
            if c.number_format == "@" and C.is_formula(c.value):
                c.number_format = "General"


def _move_footer(ws, row, c1, c2):
    """Seitenfuß (global_rules.early) an eine neue Zeile verlegen – Platz für die Buttonzeile darüber."""
    old = None
    for c in ws._cells.values():
        if isinstance(c.value, str) and c.value == C.FOOTER_1:
            old = c.row
            break
    if old is not None and old != row:
        for r in (old, old + 1):
            for c in C.iter_cells(ws, "A", r, "Z", r):
                if not C.is_formula(c.value):
                    c.value = None
                c.border = Border()
                c.fill = C.NOFILL
            ws.row_dimensions[r].height = None
    C.footer(ws, row, c1, c2)


def _nav_row(ws, row, back, nxt, footer_cols):
    """P22: Zurück/Weiter als Buttonzeile am Seitenende (dieselbe Komponente wie S01–S12): Leerzeile · Buttons ·
    Leerzeile · Seitenfuß. back/nxt: (c1, c2, Text, Zielblatt)."""
    C.set_height(ws, row - 1, C.H_GAP)
    items = [dict(c1=back[0], c2=back[1], text=back[2], target=back[3], kind="secondary",
                  tooltip=f"Zurück zu {back[3]}"),
             dict(c1=nxt[0], c2=nxt[1], text=nxt[2], target=nxt[3], kind="primary", tooltip=f"Weiter zu {nxt[3]}")]
    _move_footer(ws, row + 2, *footer_cols)
    for it in items:
        _unmerge(ws, row, row, it["c1"], it["c2"])
    out = C.btn_row(ws, row, items)
    C.set_height(ws, row + 1, C.H_GAP)
    return out


def _made_for(cell, size=C.T_BODY):
    cell.value = '=IFERROR(IF(Erstellt_fuer="","","Erstellt für "&Erstellt_fuer),"")'
    cell.font = C.font(size, True, C.NAVY)
    cell.alignment = C.align("right", "center")


def _heights(ws, rows):
    return {r: ws.row_dimensions[r].height for r in rows}


def _restore(ws, saved):
    for r, h in saved.items():
        ws.row_dimensions[r].height = h


DATE_TXT = 'TEXT(DAY({d}),"00")&"."&TEXT(MONTH({d}),"00")&"."&YEAR({d})'   # gebietsschema-unabhängig
RECHTSFORM_LONG = 'CHOOSE(Rechtsform_Idx,"Privatperson","Vermögensverwaltende GmbH","GmbH")'


# =============================================================================== Bankgespräch
BANK_PANELS = (("B", "C"), ("E", "F"), ("H", "I"))
PHOTO_ROWS = (9, 15)                        # eingeklappte Foto-Flächen (Gruppe 9–16, Z. 16 = Fuge)
PHOTO_HINTS = ("Foto Außenansicht einfügen", "Foto Wohnraum einfügen", "Lageplan einfügen")
BANK_BULLETS = [
    "• Drucken Sie diese Übersicht (A4, Hochformat) und legen Sie sie der Finanzierungsanfrage bei. Alle Zahlen "
    "kommen aus der Kalkulation – Änderungen im Leitfaden oder auf dem Blatt „Eingaben“ aktualisieren die "
    "Übersicht automatisch.",
    "• Für die Bonitätsprüfung erwarten Banken eine Selbstauskunft. Die Vorlagen „Haushaltsrechnung“ und "
    "„Vermögensaufstellung“ sind vorbereitet und fließen in den Block „Sicherheiten & Bonität“ ein.",
    "• Argumentieren Sie mit Kapitaldienstdeckung (DSCR), Eigenkapitalquote, Beleihungsauslauf und dem Cashflow "
    "nach Steuern – das sind die Kennzahlen, auf die Kreditentscheider zuerst schauen.",
    "• Objektfotos (Außenansicht, Wohnräume, Lageplan): Zeilen 9–16 über das [+] am linken Rand einblenden und "
    "die Bilder über Einfügen › Bilder in die drei Flächen ziehen.",
]
BANK_LABELS = {
    "H20": "Nettokaltmiete Soll pro Monat",
    "H25": "Nettokaltmiete Ist (nach Ausfall)",
    "H30": "= Cashflow nach Steuern ¹",
    "E36": "Zu versteuerndes Einkommen",
    "E37": "Haushaltsüberschuss p. a.",
    "E38": "Nettovermögen",
    "E41": "Beleihungsobjekt (Ort)",
}
BANK_CHECKLIST = (
    "• Selbstauskunft: Haushaltsrechnung und Vermögensaufstellung (Unterreiter oben) ausfüllen und unterschreiben.\n\n"
    "• Einkommen: die letzten drei Gehaltsabrechnungen bzw. zwei Einkommensteuerbescheide.\n\n"
    "• Eigenkapital: aktueller Konto- oder Depotauszug.\n\n"
    "• Objekt: Exposé, Grundbuchauszug, Teilungserklärung, Energieausweis, Wohnflächenberechnung, Fotos.\n\n"
    "• Vermietung: Mietvertrag bzw. Mietaufstellung.")
BANK_NOTE = ("¹ Einschließlich Steuereffekt des Objekts (Steuererstattung bzw. -zahlung).  ·  Alle Werte stammen aus "
             "der Kalkulation (Blätter Eingaben, Finanzierung, Projektion, Steuern). Prognosewerte sind Annahmen und "
             "keine Zusicherung. Haushaltsrechnung und Vermögensaufstellung liegen als separate Vorlagen bei.  ·  "
             "Objektfotos: Zeilen 9–16 über [+] am linken Rand einblenden.")
PCT1 = ("I22", "C36", "C38", "C39", "C41", "F29", "F30")
BANK_STATUS = (("C35", "DSCR"), ("C36", "NMR"), ("C39", "IRR"), ("C41", "EKR"), ("I22", "BMR"))
CASHFLOW_CELLS = ("I29", "I30", "I39")       # Cashflow-Beträge: < 0 rot, ≥ 0 neutral (P1-10)


def _bank_header(ws):
    """Seitenkopf-Standard (P1-18): Z. 5 Überzeile · Z. 6 Titel 22 pt + „Erstellt für“ · Z. 7 Objekt + Meta."""
    _unmerge(ws, 5, 8, "B", "I")
    for c in C.iter_cells(ws, "B", 5, "I", 8):
        c.value = None
        c.fill = C.NOFILL
        c.border = Border()
    subtitle = ('=Obj_Name&"  –  "&Obj_Adresse&"  ·  Kauf geplant "&' + DATE_TXT.format(d="Kaufdatum"))
    C.page_header(ws, "B", "I", "Bank  ›  Bankgespräch", "Investitionsübersicht für das Bankgespräch", subtitle)
    C.safe_merge(ws, "B", 7, "G", 7)
    ws["B7"].alignment = C.align("left", "top")

    made_for = ws["H6"]
    made_for.value = '=IFERROR(IF(Erstellt_fuer="","","Erstellt für "&Erstellt_fuer),"")'
    C.safe_merge(ws, "H", 6, "I", 6)
    made_for.font = C.font(C.T_BODY, True, C.NAVY)
    made_for.alignment = C.align("right", "bottom")
    meta = ws["H7"]
    meta.value = ('="Kauf als: "&' + RECHTSFORM_LONG + '&"  ·  Stand "&' + DATE_TXT.format(d="TODAY()"))
    C.safe_merge(ws, "H", 7, "I", 7)
    meta.font = C.font(C.T_SMALL, False, C.MUTED)
    meta.alignment = C.align("right", "top")
    C.set_height(ws, 4, 12)
    C.set_height(ws, 8, C.H_GAP)


def _bank_photos(ws):
    """P1-01: Foto-Flächen ohne Strichrahmen, mit Hinweis, eingeklappt (Gliederung Z. 9–16)."""
    r1, r2 = PHOTO_ROWS
    _unmerge(ws, r1, r2 + 1, "B", "I")
    _clear(ws, "B", r1, "I", r2 + 1)
    for (c1, c2), hint in zip(BANK_PANELS, PHOTO_HINTS):
        for c in C.iter_cells(ws, c1, r1, c2, r2):
            c.fill = C.fill(C.TINT_XL)
        C.safe_merge(ws, c1, r1, c2, r2)
        cell = ws[f"{c1}{r1}"]
        cell.value = hint
        cell.font = C.font(C.T_MICRO, False, C.MUTED, italic=True)
        cell.alignment = C.align("center", "center")
    for r in range(r1, r2 + 1):
        C.set_height(ws, r, 12)
    C.set_height(ws, r2 + 1, C.H_GAP)
    for r in range(r1, r2 + 2):
        ws.row_dimensions[r].outline_level = 1
        ws.row_dimensions[r].hidden = True
    if ws.sheet_properties.outlinePr is None:
        ws.sheet_properties.outlinePr = Outline(summaryBelow=True, summaryRight=True)
    else:
        ws.sheet_properties.outlinePr.summaryBelow = True


def bank(ws):
    # ---- Raster: drei Spalten Beschriftung + Wert, schmale Stege (Druck A4 hoch)
    widths = {"A": 4.5, "B": 27, "C": 18, "D": 1.5, "E": 27, "F": 18, "G": 1.5, "H": 28, "I": 18, "J": 2.5,
              "K": 48}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    ws.sheet_view.showGridLines = False

    _bank_header(ws)
    _bank_photos(ws)

    # ---- Abschnitte Ebene 1 (Z. 17 / 32) und Unterabschnitte Ebene 2 (Linie, ohne Fläche)
    for row in (17, 32):
        for c1, c2 in BANK_PANELS:
            C.section(ws, row, c1, c2)
        C.set_height(ws, row + 1, 6)
    for row, panels in ((19, BANK_PANELS), (24, BANK_PANELS), (27, [("E", "F")])):
        for c1, c2 in panels:
            C.section(ws, row, c1, c2, level=2, variant="line", height=C.H_ROW)

    for coord, text in BANK_LABELS.items():
        C.set_text(ws[coord], text)
    ws["F36"].value = _fmt_sub(ws["F36"].value, '"GmbH – Jahresabschluss"', '"lt. Jahresabschluss"')
    ws["C40"].value = _fmt_sub(ws["C40"].value, '" J. / "', '" Jahre / "')
    if not C.is_formula(ws["B39"].value):                  # P20: kanonische KPI-Beschriftung
        ws["B39"].value = C.kpi_label("IRR", formula=True)
    # P3-09: Beleihungsobjekt als Label-Wert-Paar (F41 = reine Anzeigeformel, rechtsbündig)
    _unmerge(ws, 41, 41, "E", "F")
    ws["F41"].value = '=IFERROR(TRIM(MID(Obj_Adresse,FIND(",",Obj_Adresse)+1,200)),Obj_Adresse)'   # nur Ort

    # ---- Datenzeilen: eine Text- und eine Wertkante je Panel, Haarlinie, alles 10 pt 1A1D21
    heads = {(19, "B"), (19, "E"), (19, "H"), (24, "B"), (24, "E"), (24, "H"), (27, "E")}
    for r in list(range(20, 31)) + list(range(34, 42)):
        C.set_height(ws, r, C.H_ROW)
        for c1, c2 in BANK_PANELS:
            if (r, c1) in heads:
                continue
            lab, val = ws[f"{c1}{r}"], ws[f"{c2}{r}"]
            if lab.value is None and val.value is None:
                continue
            _hair(ws, r, c1, c2)
            _label(lab)
            _value(val)
    # P27: alle Datenzeilen 18 pt, keine Umbrüche in Wertzellen (F41 zeigt nur den Ort, die Adresse steht im Kopf)
    ws["F41"].alignment = C.align("right", "center", 1)
    C.set_height(ws, 41, C.H_ROW)

    for coord in PCT1:
        ws[coord].number_format = C.NUMFMT["pct1"]
    ws["C35"].number_format = C.NUMFMT["dscr"]
    ws["I23"].number_format = C.NUMFMT["mult1"]
    ws["I21"].number_format = C.NUMFMT["eur2"]

    # ---- Summenstufen (P1-14): Blockergebnisse Gesamtinvestition / Eigenkapital / CF n. St., Zwischensumme CF v. St.
    C.sum_row(ws, 29, "B", "C", "result")
    C.sum_row(ws, 28, "E", "F", "result")
    C.sum_row(ws, 29, "H", "I", "sub")
    C.sum_row(ws, 30, "H", "I", "result")
    for coord in ("B29", "E28", "H29", "H30"):
        ws[coord].alignment = C.align("left", "center", 1)
    for coord in ("C29", "F28", "I29", "I30"):
        ws[coord].alignment = C.align("right", "center", 1)
    C.set_height(ws, 31, C.H_GAP)

    # ---- Statusfarben-Disziplin (P1-10): Kennzahlen mit Zielkorridor per Schriftfarbe (Legende in B64),
    #      Cashflow-Beträge nur rot, wenn negativ – sonst neutral
    for coord, kpi in BANK_STATUS:
        C.status_cf(ws, coord, C.status_conditions(kpi, coord))
    for coord in CASHFLOW_CELLS:
        C.neg_red(ws, coord, bold=True)

    # ---- Diagrammzeile: Abschnitt, Kreis B:C, Säulen E:I (Zeilen 44–57)
    C.section(ws, 43, "B", "I")
    C.set_height(ws, 42, C.H_GAP)
    for r in range(44, 58):
        C.set_height(ws, r, 15)
    charts = sorted(ws._charts, key=lambda ch: ch.anchor._from.col)
    if len(charts) >= 2:
        _anchor(charts[0], "B", 44, "C", 57)
        _anchor(charts[1], "E", 44, "I", 57)
    for r, h in {58: 10, 59: 4, 60: 4}.items():
        C.set_height(ws, r, h)

    # ---- Annahmen, Fußnote ¹ (+ Foto-Hinweis) und Farblegende (P1-01, P1-10)
    _unmerge(ws, 61, 64, "B", "I")
    b61 = ws["B61"]
    if isinstance(b61.value, str):
        b61.value = b61.value.replace(' | ', '  ·  ')
    C.set_text(ws["B63"], BANK_NOTE)
    for coord in ("B61", "B63"):
        r = ws[coord].row
        C.safe_merge(ws, "B", r, "I", r)
        cell = ws[coord]
        cell.font = C.font(C.T_SMALL, False, C.MUTED)
        cell.alignment = C.align("left", "top", 0, wrap=True)
        n = C.lines_needed_metric(C.display_text(cell), C.span_px(ws, "B", "I"), C.T_SMALL)
        need = n * C.line_pt(C.T_SMALL) + 5
        C.set_height(ws, r, min(h for h in C.ROW_RASTER if h >= need))
    C.set_height(ws, 62, 4)
    legend = ws["B64"]
    legend.value = C.rich([("Farbkennzeichnung der Kennzahlen (Schwellen lt. Konfiguration):   ", C.T_MICRO, False,
                            C.MUTED)] + list(_legend_parts()) +
                          [("   ·   negative Cashflow-Beträge rot", C.T_MICRO, False, C.MUTED)])
    legend.font = C.font(C.T_MICRO, False, C.MUTED)
    legend.alignment = C.align("left", "center")
    C.set_height(ws, 64, 16)

    # ---- Rail rechts (nur Bildschirm, außerhalb des Druckbereichs): zwei Hinweis-Boxen, bündig mit den beiden
    #      Abschnittsreihen (Z. 17–30 und 32–41); Navigation nur noch Unterreiter oben + Buttonzeile unten (P22)
    _unmerge(ws, 5, 45, "K", "K")
    _clear(ws, "K", 5, "K", 45)
    C.callout_box(ws, "K", 17, "K", 18, 30, title="Bessere Konditionen im Bankgespräch",
                  text="\n\n".join(BANK_BULLETS), pill=False, fit=None)
    C.set_height(ws, 17, C.H_BAND)          # Kopf bündig mit den Abschnittsköpfen der Zeile 17
    C.set_height(ws, 18, 6)
    C.callout_box(ws, "K", 32, "K", 33, 41, title="Unterlagen für die Finanzierungsanfrage", text=BANK_CHECKLIST,
                  pill=False, fit=None)
    C.set_height(ws, 32, C.H_BAND)
    C.set_height(ws, 33, 6)

    # ---- Buttonzeile am Seitenende (P22) – gleiche Breiten: B:C ≈ H:I
    _nav_row(ws, 66, ("B", "C", "‹  Zurück: Sensitivität", "Sensitivität"),
             ("H", "I", "Weiter: Haushaltsrechnung  ›", HH), ("B", "I"))
    _text_formula_general(ws)
    C.cf_close(ws)


def _legend_parts():
    """„● erfüllt · ● prüfen · ● kritisch“ als Rich-Text-Teile (8 pt, Punkte in Statusfarbe)."""
    for i, (lvl, lab) in enumerate((("green", "erfüllt"), ("amber", "prüfen"), ("red", "kritisch"))):
        if i:
            yield ("   ·   ", C.T_MICRO, False, C.MUTED)
        yield (C.STATUS_DOT + " ", C.T_MICRO, True, C.STATUS_COLORS[lvl][0])
        yield (lab, C.T_MICRO, False, C.MUTED)


# =============================================================================== Haushaltsrechnung / Vermögensaufstellung
FORM_WIDTHS = {"A": 4.5, "B": 56, "C": 14, "D": 14, "E": 36, "F": 3,
               "G": 12, "H": 12, "I": 12, "J": 2, "K": 12, "L": 12, "M": 12}
HH_LABELS = {
    "B42": "Bewirtschaftung des kalkulierten Objekts (Jahr 1, inkl. Rücklage)",
    "B27": "= Summe Einnahmen",
    "B43": "= Summe Ausgaben",
}
VA_LABELS = {"B19": "= Summe Vermögenswerte", "B29": "= Summe Verbindlichkeiten"}
HH_LINKED = ("C26", "C41", "C42")            # aus der Kalkulation übernommen (überschreibbar)
VA_LINKED = ("C33",)


def _form_common(ws, first, last, comment_rows):
    """Spaltenraster, eine Textkante (Einzug 1), Werte rechts mit Einzug 1, Kommentare 9 pt grau."""
    for k, w in FORM_WIDTHS.items():
        ws.column_dimensions[k].width = w
    ws.sheet_view.showGridLines = False
    bands = set()
    for r in range(first, last + 1):
        b = ws.cell(r, 2)
        fill = b.fill.fgColor.rgb[-6:] if b.fill is not None and b.fill.fill_type == "solid" else None
        is_head = isinstance(b.value, str) and (b.value.upper() == b.value or fill in (C.TINT, C.HEAD))
        if fill in (C.TINT, C.HEAD) and is_head:
            bands.add(r)
            continue
        if isinstance(b.value, str) and not C.is_formula(b.value):
            b.alignment = C.align("left", "center", 1)
        for cc in (3, 4):
            v = ws.cell(r, cc)
            if v.value is None:
                continue
            if v.alignment.horizontal == "right" or isinstance(v.value, (int, float)) or C.is_formula(v.value):
                v.alignment = C.align("right", "center", 1)
    for r in comment_rows:
        e = ws.cell(r, 5)
        # Spaltenköpfe („KOMMENTAR“ …) und Bänder behalten ihren Stil
        if r in bands or (isinstance(e.value, str) and e.value.isupper()):
            continue
        if e.value is not None:
            e.font = C.font(C.T_SMALL, False, C.MUTED)
            e.alignment = C.align("left", "center", 1)
    return bands


def _linked(ws, coords):
    """P31: überschreibbare Verknüpfungen als eigene Eingabe-Variante – 1D4F8A normal auf FFF5D6, Rahmen gestrichelt
    E6CB77 (Legende im Seitenkopf rechts: „gestrichelt = aus der Kalkulation, überschreibbar“)."""
    dash = C.side("dashed", C.INPUT_LINE)
    for coord in coords:
        cell = ws[coord]
        cell.font = C.font(C.T_BODY, False, C.BLUE)
        cell.fill = C.fill(C.INPUT_BG)
        cell.border = Border(left=dash, right=dash, top=dash, bottom=dash)
        cell.alignment = C.align("right", "center", 1)


def _comment_inputs(ws, rows, bands=()):
    """P31: entsperrte Kommentar-/Nachweiszellen (Spalte E) als dezente Eingabe – FFF9EA, Unterkante E6CB77,
    9 pt kursiv 5B6068. Spaltenköpfe und Bänder bleiben unberührt."""
    for r in rows:
        e = ws.cell(r, 5)
        if r in bands or e.protection.locked is not False:
            continue
        if isinstance(e.value, str) and (e.value.isupper() or C.is_formula(e.value)):
            continue
        e.fill = C.fill(C.NOTE_BG)
        e.border = Border(bottom=C.side("thin", C.INPUT_LINE))
        e.font = C.font(C.T_SMALL, False, C.MUTED, italic=True)
        e.alignment = C.align("left", "center", 1)


def _form_header(ws):
    """P04: Kopf rechts an der Inhaltskante M – Z. 6 „Erstellt für …“, Z. 7 Legende der Eingabe-Arten.
    (Z. 5 rechts: Unterreiter, navigation.py.)"""
    for coord in ("E6", "E7"):
        cell = ws[coord]
        cell.value = None
        cell.hyperlink = None
    _made_for(ws["M6"])
    lg = ws["M7"]
    lg.value = C.rich([("■ ", C.T_LABEL, True, C.INPUT_LINE), ("Eingabe", C.T_LABEL, False, C.MUTED),
                       ("   ·   ", C.T_LABEL, False, C.MUTED),
                       ("┅ ", C.T_LABEL, True, C.INPUT_LINE),
                       ("gestrichelt = aus der Kalkulation, überschreibbar", C.T_LABEL, False, C.MUTED)])
    lg.font = C.font(C.T_LABEL, False, C.MUTED)
    lg.alignment = C.align("right", "center")


def _gaps(ws, rows):
    for r in rows:
        C.set_height(ws, r, C.H_GAP)


def _side_section(ws, row, title, meta=None):
    """Abschnittskopf über der rechten Spalte G:M (bündig mit dem Tabellenabschnitt derselben Zeile)."""
    for c in C.iter_cells(ws, "G", row, "M", row):
        c.value = None if not C.is_formula(c.value) else c.value
    C.section(ws, row, "G", "M", title, meta=meta)


def _tile_pair(ws, row, left, right):
    """Zwei helle Kacheln G:I und K:M (Rinne J) über vier Tabellenzeilen: Label · Wert (2 Zeilen) · Kontext.
    Die Zeilenhöhen der Tabelle links bleiben unverändert."""
    saved = _heights(ws, range(row, row + 4))
    for (c1, c2), spec in zip((("G", "I"), ("K", "M")), (left, right)):
        C.tile(ws, c1, c2, row, row + 1, row + 3, variant="light", value_rows=2, gap_right=False, **spec)
    _restore(ws, saved)


def household(ws):
    for coord, text in HH_LABELS.items():
        C.set_text(ws[coord], text)
    bands = _form_common(ws, 8, 47, range(18, 48))

    # Persönliche Angaben: alle Felder über C:E (volle Breite), Kinder als Zahl, „Beschäftigt seit“ als echtes Datum
    _unmerge(ws, 9, 14, "C", "E")
    kids = ws["C12"]
    if isinstance(kids.value, str) and kids.value.strip().isdigit():
        kids.value = int(kids.value.strip())
    since = ws["C14"]
    if isinstance(since.value, str) and re.match(r"^\d{2}\.\d{2}\.\d{4}$", since.value.strip()):
        d, m, y = (int(x) for x in since.value.strip().split("."))
        since.value = datetime.datetime(y, m, d)
    for r in range(9, 15):
        for cc in ("D", "E"):
            ws[f"{cc}{r}"].value = None
        C.safe_merge(ws, "C", r, "E", r)
        _input_box(ws, "C", r, "E")
    kids.number_format = C.NUMFMT["int_in"]
    since.number_format = C.NUMFMT["date"]
    dv_k = DataValidation(type="whole", operator="between", formula1="0", formula2="20", allow_blank=True,
                          showInputMessage=True, showErrorMessage=True)
    dv_k.promptTitle, dv_k.prompt = "Kinder", "Anzahl unterhaltspflichtiger Kinder (0–20)."
    dv_k.errorTitle, dv_k.error = "Ungültige Anzahl", "Bitte eine ganze Zahl zwischen 0 und 20 eingeben."
    dv_d = DataValidation(type="date", operator="between", formula1="DATE(1950,1,1)", formula2="TODAY()",
                          allow_blank=True, showInputMessage=True, showErrorMessage=True)
    dv_d.promptTitle, dv_d.prompt = "Beschäftigt seit", "Datum im Format TT.MM.JJJJ, z. B. 01.01.2015."
    dv_d.errorTitle, dv_d.error = "Ungültiges Datum", "Bitte ein Datum zwischen 1950 und heute eingeben."
    for dv, ref in ((dv_k, "C12"), (dv_d, "C14")):
        ws.add_data_validation(dv)
        dv.add(ref)

    _linked(ws, HH_LINKED)
    _comment_inputs(ws, range(18, 48), bands)
    # Nullwerte je Zeile einheitlich (Befund): Eingabe links „0 €“ → Jahreswert rechts ebenfalls „0 €“
    for r in list(range(18, 27)) + list(range(31, 43)):
        ws[f"D{r}"].number_format = C.NUMFMT["eur_in"]
        ws[f"C{r}"].number_format = C.NUMFMT["eur_in"]
    for r in (27, 43):
        C.sum_row(ws, r, "B", "E", "sub")
    _gaps(ws, (15, 28, 44))
    # P29: Spaltenköpfe im Ergebnisband, „Überschuss“ statt „Überschuss pro Monat“
    for coord, text in (("C45", "PRO MONAT"), ("D45", "PRO JAHR")):
        cell = ws[coord]
        C.set_text(cell, text)
        cell.font = C.font(C.T_LABEL, True, C.BLUE)
        cell.alignment = C.align("right", "center", 1)
    C.set_text(ws["B46"], "Überschuss")

    # Ergebnis: EINE Statusregel auf Basis der Überschussquote C47 (< 0 rot · < 5 % amber · sonst grün), P1-10
    C.set_text(ws["E46"], "Banken erwarten > 0 nach neuem Kapitaldienst")
    ws["E46"].font = C.font(C.T_SMALL, False, C.MUTED)
    ws["E46"].alignment = C.align("left", "center", 1)
    ws["C47"].number_format = C.NUMFMT["pct1"]
    _drop_cf(ws, "C46:D46", "C47")
    conds = [("AND(ISNUMBER($C$47),$C$47<0)", "red"), ("AND(ISNUMBER($C$47),$C$47<0.05)", "amber"),
             ("ISNUMBER($C$47)", "green")]
    C.status_cf(ws, "C46:D46", conds)
    C.status_cf(ws, "C47", conds)
    C.status_pill(ws, ws["E47"], conditions=conds, style="chip", h="left", suffix='"Ziel ≥ 5 %"')
    C.neg_red(ws, "C46:D46")

    # Rechte Spalte: Abschnitt + Diagramm (Oberkante bündig mit Z. 8), darunter die Ergebnis-Kacheln
    _side_section(ws, 8, "Ausgabenstruktur")
    for ch in ws._charts:
        _anchor(ch, "G", 9, "M", 27)
    _side_section(ws, 29, "Ergebnis der Haushaltsrechnung")
    _tile_pair(ws, 30,
               dict(label="Überschuss pro Monat", value="=C46", fmt=C.NUMFMT["eur"], sub='="p. a. "&FIXED(D46,0)&" €"',
                    status_col="I", conditions=conds),
               dict(label="Überschussquote", value="=C47", fmt=C.NUMFMT["pct1"], sub="Ziel ≥ 5 % vom Einkommen",
                    status_col="M", conditions=conds))
    _form_header(ws)
    _nav_row(ws, 49, ("C", "D", "‹  Zurück: Bankgespräch", BANK),
             ("E", "E", "Weiter: Vermögensaufstellung  ›", VA), ("B", "E"))
    _text_formula_general(ws)
    C.cf_close(ws)


def assets(ws):
    for coord, text in VA_LABELS.items():
        C.set_text(ws[coord], text)
    bands = _form_common(ws, 8, 34, range(10, 35))
    # Verpfändet? – Auswahlliste ja/nein, „ja“ in Amber (belastete Werte fallen auf)
    dv = DataValidation(type="list", formula1='"ja,nein"', allow_blank=True, showInputMessage=True,
                        showErrorMessage=True)
    dv.promptTitle, dv.prompt = "Verpfändet?", "„ja“, wenn der Vermögenswert bereits als Sicherheit dient."
    dv.errorTitle, dv.error = "Bitte auswählen", "Nur „ja“ oder „nein“ aus der Liste."
    ws.add_data_validation(dv)
    dv.add("D10:D18")
    for r in range(10, 19):
        ws.cell(r, 4).alignment = C.align("center", "center")
    ws.conditional_formatting.add("D10:D18", FormulaRule(formula=['LOWER(TRIM(D10))="ja"'], stopIfTrue=True,
                                                         font=Font(color=C.AMBER, bold=True)))
    ws["D9"].alignment = C.align("center", "center")

    _linked(ws, VA_LINKED)
    _comment_inputs(ws, range(10, 35), bands)
    for r in (19, 29):
        C.sum_row(ws, r, "B", "E", "sub")
    # P12: Summenzeile mit EINER Schrift – „–“ in D29 im Stil von C29 (fett 0B2A4A), die Monatsrate wird nicht summiert
    d29 = ws["D29"]
    if d29.value is None:
        d29.value = "–"
    d29.font = C.font(C.T_BODY, True, C.NAVY)
    d29.alignment = C.align("right", "center", 1)
    C.neg_red(ws, "C34")
    _gaps(ws, (20, 30))

    # Rechte Spalte: Abschnitt + Diagramm, darunter Kennzahl-Kacheln (P2-17) als Anzeige-Verweise
    _side_section(ws, 8, "Vermögensstruktur")
    for ch in ws._charts:
        _anchor(ch, "G", 9, "M", 19)
    _side_section(ws, 21, "Kennzahlen für die Bank")
    liq = [("AND(ISNUMBER($C$34),$C$34<0)", "red"), ("ISNUMBER($C$34)", "green")]
    _tile_pair(ws, 22,
               dict(label="Nettovermögen", value="=C32", fmt=C.NUMFMT["eur"],
                    sub="Vermögenswerte − Verbindlichkeiten", status=None),
               dict(label="Liquide Mittel nach EK-Einsatz", value="=C34", fmt=C.NUMFMT["eur"],
                    sub='="nach Eigenkapital "&FIXED(C33,0)&" €"', conditions=liq))
    _form_header(ws)
    _nav_row(ws, 36, ("C", "D", "‹  Zurück: Haushaltsrechnung", HH),
             ("E", "E", "↺  Zurück zur Übersicht: Bankgespräch", BANK), ("B", "E"))
    _text_formula_general(ws)
    C.cf_close(ws)


def apply(wb):
    if BANK in wb.sheetnames:
        bank(wb[BANK])
    if HH in wb.sheetnames:
        household(wb[HH])
    if VA in wb.sheetnames:
        assets(wb[VA])
